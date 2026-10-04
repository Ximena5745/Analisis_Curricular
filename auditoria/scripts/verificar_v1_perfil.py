"""
V1 (propuesta): correspondencia perfil de egreso → competencias → RA, por campo del Paso 1.

Estructura del Paso 1 (fila 3 en adelante):
  C Perfil profesional  → se desagrega en E Saber, F SaberHacer, G SaberSer (una fila por componente)
  D Perfil ocupacional  → se desagrega en H Áreas profesionales, I Tareas, J Poblaciones de actuación
  K Valor agregado
Vínculos formales de la plantilla (listas desplegables):
  Paso 2: Objeto conceptual ← Saber; Finalidad ← SaberHacer; Condición ← SaberSer
  Paso 3: SaberAsociado ← saber del perfil del tipo indicado en TipoSaber
Un elemento «corresponde» si el texto normalizado es igual o su similitud es ≥ 0,80 (difflib, sin autojunk).
Unidad: celda = elemento (sin dividir; criterio de conteo de la auditoría del Paso 1).

Uso: python auditoria/scripts/verificar_v1_perfil.py
Salida: auditoria/evidencia_v1_perfil.json
"""
import glob
import json
import os
import re
import sys
import unicodedata
import warnings
from difflib import SequenceMatcher

import openpyxl

warnings.filterwarnings('ignore')
sys.stdout.reconfigure(encoding='utf-8')
TIPOS = ('Saber', 'SaberHacer', 'SaberSer')


def norm(t):
    t = unicodedata.normalize('NFKD', str(t)).encode('ascii', 'ignore').decode('ascii').lower()
    return re.sub(r'\s+', ' ', re.sub(r'[^a-z0-9 ]', ' ', t)).strip()


def ne(v):
    return v is not None and str(v).strip() not in ('', 'nan', 'None') and not str(v).strip().startswith('[')


def igual(a, b):
    return a == b or SequenceMatcher(None, a, b, autojunk=False).ratio() >= 0.8


def hoja(wb, pref):
    return next(wb[s] for s in wb.sheetnames if s.strip().startswith(pref) and 'backup' not in s.lower())


def tipo(t):
    t = norm(t).replace(' ', '')
    return 'SaberSer' if 'ser' in t else 'SaberHacer' if 'hacer' in t else 'Saber' if 'saber' in t else None


filas = []
for f in sorted(glob.glob('data/raw/FORMATOS RA CICLO UNO RC/*.xlsx')):
    matriz = os.path.basename(f)[10:-5]
    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    perfil = {'componentes': [], 'Perfil profesional': 0, 'Perfil ocupacional': 0,
              'Áreas profesionales': 0, 'Tareas profesionales': 0, 'Poblaciones actuación': 0, 'Valor agregado': 0}
    for r in hoja(wb, 'Paso1').iter_rows(min_row=3, values_only=True):
        r = list(r) + [None] * 12
        perfil['Perfil profesional'] += ne(r[2])
        perfil['Perfil ocupacional'] += ne(r[3])
        for j, k in ((7, 'Áreas profesionales'), (8, 'Tareas profesionales'), (9, 'Poblaciones actuación'), (10, 'Valor agregado')):
            perfil[k] += ne(r[j])
        if any(ne(r[j]) for j in (4, 5, 6)):
            c = {t: norm(r[j]) if ne(r[j]) else '' for t, j in zip(TIPOS, (4, 5, 6))}
            c['generica'] = 'fenomenos contemporaneos' in c['Saber']
            perfil['componentes'].append(c)
    comp = []
    for r in hoja(wb, 'Paso 2').iter_rows(min_row=3, values_only=True):
        if len(r) > 5 and ne(r[5]) and not norm(r[5]).startswith('redaccion'):
            comp.append({'Saber': norm(r[2] or ''), 'SaberHacer': norm(r[3] or ''), 'SaberSer': norm(r[4] or ''),
                         'generica': 'fenomenos contemporaneos' in norm(r[5]), 'texto': str(r[5]).strip()})
    ra = {}
    for r in hoja(wb, 'Paso 3').iter_rows(min_row=3, values_only=True):
        if len(r) > 8 and ne(r[8]) and tipo(r[2] or ''):
            d = ra.setdefault(str(r[8]).strip().lower(), {'asociados': []})
            d['asociados'].append((tipo(r[2]), norm(r[3] or '')))
    wb.close()

    elementos = [(t, c[t]) for c in perfil['componentes'] for t in TIPOS if c[t]]
    gen = [c['generica'] for c in perfil['componentes'] for t in TIPOS if c[t]]
    # Perfil profesional → competencia: componente con competencia cuyo campo equivalente coincide
    comp_ok = [any(c[t] and igual(c[t], k[t]) for k in comp) for c in perfil['componentes'] for t in TIPOS if c[t]]
    # Competencia completa: las tres partes provienen de una misma fila del perfil
    comp_trazable = [any(all(k[t] and c[t] and igual(c[t], k[t]) for t in TIPOS) for c in perfil['componentes']) for k in comp]
    # Perfil → RA: elemento del perfil asociado a ≥ 1 RA (SaberAsociado del mismo tipo)
    asoc = [a for d in ra.values() for a in d['asociados']]
    elem_ra = [any(t == ta and igual(e, a) for ta, a in asoc) for t, e in elementos]
    # RA → perfil: RA único cuyo SaberAsociado corresponde a un elemento del perfil
    ra_ok = [any(any(t == te and igual(a, e) for te, e in elementos) for t, a in d['asociados']) for d in ra.values()]
    filas.append({
        'matriz': matriz, 'sede': matriz.rsplit('_', 1)[1],
        'perfil_profesional': perfil['Perfil profesional'], 'perfil_ocupacional': perfil['Perfil ocupacional'],
        'componentes_perfil': len(elementos),
        'componentes_con_competencia': sum(comp_ok),
        'competencias': len(comp), 'competencias_trazables': sum(comp_trazable),
        'componentes_con_ra': sum(elem_ra),
        'componentes_programa': sum(not g for g in gen), 'componentes_programa_con_ra': sum(o for o, g in zip(elem_ra, gen) if not g),
        'componentes_genericos': sum(gen), 'componentes_genericos_con_ra': sum(o for o, g in zip(elem_ra, gen) if g),
        'componentes_programa_con_competencia': sum(o for o, g in zip(comp_ok, gen) if not g),
        'componentes_programa_sin_ra': [f'{t}: {e}' for (t, e), o, g in zip(elementos, elem_ra, gen) if not o and not g],
        'por_tipo': {t: [sum(o for (te, _), o in zip(elementos, elem_ra) if te == t), sum(te == t for te, _ in elementos)] for t in TIPOS},
        'ra_unicos': len(ra), 'ra_trazables': sum(ra_ok),
        'ocupacional': {k: perfil[k] for k in ('Áreas profesionales', 'Tareas profesionales', 'Poblaciones actuación', 'Valor agregado')},
        'competencias_no_trazables': [k['texto'] for k, o in zip(comp, comp_trazable) if not o],
        'componentes_sin_ra': [f'{t}: {e}' for (t, e), o in zip(elementos, elem_ra) if not o],
    })

S = lambda k: sum(m[k] for m in filas)
pct = lambda a, b: round(100 * a / b, 1) if b else None
cob = json.load(open('auditoria/evidencia_cobertura_perfil.json', encoding='utf-8'))['por_campo']['D']
EV = {
    'matrices': len(filas),
    'con_perfil_profesional': sum(m['perfil_profesional'] > 0 for m in filas),
    'con_perfil_ocupacional': sum(m['perfil_ocupacional'] > 0 for m in filas),
    'A_profesional_a_competencia': {'componentes': S('componentes_perfil'), 'con_competencia': S('componentes_con_competencia'),
                                    'pct': pct(S('componentes_con_competencia'), S('componentes_perfil'))},
    'A2_competencias_trazables': {'competencias': S('competencias'), 'trazables': S('competencias_trazables'),
                                  'pct': pct(S('competencias_trazables'), S('competencias'))},
    'B_profesional_a_RA': {'componentes': S('componentes_perfil'), 'con_ra': S('componentes_con_ra'),
                           'pct': pct(S('componentes_con_ra'), S('componentes_perfil')),
                           'por_tipo': {t: [sum(m['por_tipo'][t][0] for m in filas), sum(m['por_tipo'][t][1] for m in filas)] for t in TIPOS}},
    'B_diferenciada': {'programa': [S('componentes_programa_con_ra'), S('componentes_programa'), pct(S('componentes_programa_con_ra'), S('componentes_programa'))],
                       'generico': [S('componentes_genericos_con_ra'), S('componentes_genericos'), pct(S('componentes_genericos_con_ra'), S('componentes_genericos'))]},
    'A_programa': [S('componentes_programa_con_competencia'), S('componentes_programa'), pct(S('componentes_programa_con_competencia'), S('componentes_programa'))],
    'programa_sin_ra': [(m['matriz'], e) for m in filas for e in m['componentes_programa_sin_ra']],
    'B2_RA_a_perfil': {'ra_unicos': S('ra_unicos'), 'trazables': S('ra_trazables'), 'pct': pct(S('ra_trazables'), S('ra_unicos'))},
    'C_ocupacional_a_asignaturas_D9': {k: cob[k] for k in ('Áreas profesionales', 'Tareas profesionales', 'Poblaciones actuación', 'Valor agregado')},
    'C_profesional_a_asignaturas_D9': {k: cob[k] for k in ('Perfil profesional', 'Saber', 'SaberHacer', 'SaberSer')},
    'matrices_100_A': sum(m['componentes_con_competencia'] == m['componentes_perfil'] for m in filas),
    'matrices_100_B': sum(m['componentes_con_ra'] == m['componentes_perfil'] for m in filas),
    'por_matriz': filas,
}
json.dump(EV, open('auditoria/evidencia_v1_perfil.json', 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
print(json.dumps({k: v for k, v in EV.items() if k != 'por_matriz'}, ensure_ascii=False, indent=1))
for m in filas:
    if m['competencias_no_trazables']:
        print(m['matriz'], '| comp no trazables:', [c[:90] for c in m['competencias_no_trazables']])
