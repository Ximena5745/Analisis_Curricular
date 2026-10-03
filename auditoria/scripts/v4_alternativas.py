"""
V4 (trazabilidad): comparación de metodologías sobre RA únicos por matriz (D4).

  A. Actual: RA vinculado a una estrategia meso (Paso 4) con instrumento, por similitud ≥ 0,80.
  B. Sensibilidad del emparejamiento: exacto vs similitud 0,70 / 0,80 / 0,90.
  C. Embudo por eslabones: RA → estrategia meso → indicador + instrumento → asignatura (Paso 5)
     → actividad de evaluación en la asignatura.
  D. Ruta completa meso + micro: RA con estrategia e instrumento Y asignatura con actividad de evaluación.
  E. Profundidad: estrategias y asignaturas por RA.

Uso: python auditoria/scripts/v4_alternativas.py
Salida: auditoria/evidencia_v4_alternativas.json
"""
import glob
import json
import logging
import os
import re
import statistics as st
import sys
import unicodedata
import warnings
from difflib import SequenceMatcher

import openpyxl
import pandas as pd

warnings.filterwarnings('ignore')
logging.disable(logging.CRITICAL)
sys.path.insert(0, '.')
from src.extractor import ExcelExtractor  # noqa: E402

sys.stdout.reconfigure(encoding='utf-8')


def norm(t):
    t = unicodedata.normalize('NFKD', str(t)).encode('ascii', 'ignore').decode('ascii').lower()
    return re.sub(r'\s+', ' ', re.sub(r'[^a-z0-9 ]', ' ', t)).strip()


def ne(v):
    return v is not None and str(v).strip() not in ('', 'nan', 'None')


def parecido(a, b, u):
    return a == b if u >= 1 else (a == b or SequenceMatcher(None, a, b, autojunk=False).ratio() >= u)


filas = []
for f in sorted(glob.glob('data/raw/FORMATOS RA CICLO UNO RC/*.xlsx')):
    nombre = os.path.basename(f)[10:-5]
    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    p3 = wb[[s for s in wb.sheetnames if s.startswith('Paso 3 Redacción RA') and 'Backup' not in s][0]]
    ra = {}
    for r in p3.iter_rows(min_row=3, values_only=True):
        if len(r) > 8 and ne(r[8]):
            ra.setdefault(str(r[8]).strip().lower(), norm(r[8]))   # clave D4 -> texto normalizado
    p4 = next(wb[s] for s in wb.sheetnames if s.strip().startswith('Paso 4'))
    meso, est = [], None
    for r in p4.iter_rows(min_row=3, values_only=True):
        if not r or len(r) < 6:
            continue
        # Bloque de estrategia: la columna B (estrategia), C y E son celdas combinadas que
        # abarcan varias filas de RA; indicadores (D) e instrumentos (F) ocupan filas propias
        # que no coinciden con las de los RA. El bloque tiene indicador/instrumento si
        # cualquiera de sus filas lo declara (el dict es compartido por todos sus RA).
        if ne(r[1]):
            est = {'ind': False, 'ins': False}
        if est is not None:
            est['ind'] |= ne(r[3])
            est['ins'] |= ne(r[5])
        if ne(r[0]) and est is not None:
            meso.append((norm(r[0]), est))
    wb.close()
    micro = ExcelExtractor(f).extract_estrategias_micro()
    col_ev = next((c for c in micro.columns if norm(c).startswith('actividades de evaluacion')), None)
    col_ra = next((c for c in micro.columns if norm(c) == 'resultado de aprendizaje'), None)
    col_as = next((c for c in micro.columns if norm(c).startswith('nombre asignatura')), None)
    m = micro[[col_ra, col_as, col_ev]].copy() if col_ra and col_ev else pd.DataFrame(columns=['ra', 'as', 'ev'])
    m.columns = ['ra', 'as', 'ev']
    m['as'] = m['as'].ffill()
    m['ev'] = m.groupby(m['as'].fillna('').astype(str))['ev'].transform(lambda s: s.ffill().bfill())
    micro_ra = [(norm(x), ne(e)) for x, e in zip(m['ra'], m['ev']) if ne(x)]
    # Emparejamiento por mejor coincidencia mutua (≥ 0,50): recupera RA reescritos en el
    # Paso 3 cuyo texto anterior sigue en el Paso 4 (el texto del Paso 4 y el del Paso 3
    # son, cada uno, el más parecido del otro).
    sim = lambda a, b: SequenceMatcher(None, a, b, autojunk=False).ratio()
    p4_textos = sorted({x for x, _ in meso})
    p3_textos = list(ra.values())
    mejor_p3 = {x: max(p3_textos, key=lambda t: sim(t, x)) for x in p4_textos} if p3_textos else {}
    mutuo = set()
    for t in p3_textos:
        if p4_textos:
            x = max(p4_textos, key=lambda x: sim(t, x))
            if mejor_p3.get(x) == t and sim(t, x) >= 0.5:
                mutuo.add((t, x))
    for clave, t in ra.items():
        fila = {'matriz': nombre, 'sede': nombre.rsplit('_', 1)[1], 'programa': nombre.rsplit('_', 1)[0],
                'ra': clave, 'generica': 'fenomenos contemporaneos' in t}
        for u, et in ((1, 'exacto'), (0.7, 's70'), (0.8, 's80'), (0.9, 's90')):
            vinc = [e for x, e in meso if parecido(t, x, u)]
            fila[f'meso_{et}'] = bool(vinc)
            fila[f'meso_ins_{et}'] = any(e['ins'] for e in vinc)
            fila[f'meso_ind_ins_{et}'] = any(e['ins'] and e['ind'] for e in vinc)
        vinc = [e for x, e in meso if parecido(t, x, 0.8) or (t, x) in mutuo]
        fila['meso_mutuo'] = bool(vinc)
        fila['meso_ins_mutuo'] = any(e['ins'] for e in vinc)
        mic = [ev for x, ev in micro_ra if parecido(t, x, 0.8)]
        fila['micro'] = bool(mic)
        fila['micro_eval'] = any(mic)
        fila['n_estrategias'] = sum(parecido(t, x, 0.8) for x, _ in meso)
        fila['n_asignaturas'] = len(mic)
        filas.append(fila)

d = pd.DataFrame(filas)
d['completa'] = d['meso_ind_ins_s80'] & d['micro_eval']
pct = lambda c: round(100 * d[c].mean(), 1)
media_matriz = lambda c: round(d.groupby('matriz')[c].mean().mean() * 100, 1)
EV = {
    'ra_unicos': len(d),
    'A_actual_s80_instrumento': {'global': pct('meso_ins_s80'), 'media_matriz': media_matriz('meso_ins_s80')},
    'B_sensibilidad': {et: pct(f'meso_ins_{et}') for et in ('exacto', 's70', 's80', 's90')},
    'C_embudo_global': {'1_estrategia_meso': pct('meso_s80'), '2_con_indicador_e_instrumento': pct('meso_ind_ins_s80'),
                        '3_en_asignatura_paso5': pct('micro'), '4_asignatura_con_actividad_evaluacion': pct('micro_eval')},
    'D_ruta_completa_meso_y_micro': {'global': pct('completa'), 'media_matriz': media_matriz('completa'),
                                     'min_matriz': round(d.groupby('matriz')['completa'].mean().min() * 100, 1)},
    'E_profundidad': {'estrategias_por_ra_mediana': float(d.n_estrategias.median()),
                      'asignaturas_por_ra_mediana': float(d.n_asignaturas.median()),
                      'ra_sin_asignatura': int((d.n_asignaturas == 0).sum())},
    'por_sede_completa': d.groupby('sede')['completa'].mean().mul(100).round(1).to_dict(),
    'mutuo': {'global': pct('meso_ins_mutuo'), 'media_matriz': media_matriz('meso_ins_mutuo'),
              'sin_estrategia': int((~d['meso_ins_mutuo']).sum()),
              'sin_estrategia_programa': int((~d['meso_ins_mutuo'] & ~d['generica']).sum()),
              'V4_programa': round(100 * d.loc[~d['generica'], 'meso_ins_mutuo'].mean(), 1),
              'recuperados': d.loc[d['meso_ins_mutuo'] & ~d['meso_ins_s80'], ['matriz', 'ra']].to_dict('records'),
              'sin_estrategia_programa_lista': d.loc[~d['meso_ins_mutuo'] & ~d['generica'], ['matriz', 'ra']].to_dict('records')},
    'ra_sin_estrategia': int((~d['meso_ins_s80']).sum()),
    'ra_sin_estrategia_genericos': int((~d['meso_ins_s80'] & d['generica']).sum()),
    'ra_genericos': int(d['generica'].sum()),
    'V4_sin_ra_generico': round(100 * d.loc[~d['generica'], 'meso_ins_s80'].mean(), 1),
    'lista_sin_estrategia': d.loc[~d['meso_ins_s80'], ['matriz', 'generica', 'ra']].to_dict('records'),
    # V4 oficial (D16): todos los RA únicos (D4, incluido el RA genérico institucional) con
    # estrategia meso e instrumento, emparejados por similitud ≥ 0,80 o mejor coincidencia mutua ≥ 0,50
    'V4_oficial': {'ra_unicos': len(d), 'con_estrategia': int(d['meso_ins_mutuo'].sum()),
                   'V4': round(100 * d['meso_ins_mutuo'].mean(), 1),
                   'media_matriz': round(d.groupby('matriz')['meso_ins_mutuo'].mean().mean() * 100, 1),
                   'min_matriz': round(d.groupby('matriz')['meso_ins_mutuo'].mean().min() * 100, 1),
                   'ra_programa': f"{int(d.loc[~d['generica'], 'meso_ins_mutuo'].sum())}/{int((~d['generica']).sum())}",
                   'generico_con_estrategia': f"{int(d.loc[d['generica'], 'meso_ins_mutuo'].sum())}/{int(d['generica'].sum())}"},
    # Clasificación diferenciada: RA de programa vs RA genérico institucional, por sede y total
    'diferenciada': {sede: {tipo: {'ra': int(len(g)), 'con_estrategia': int(g['meso_ins_mutuo'].sum()),
                                   'V4': round(100 * g['meso_ins_mutuo'].mean(), 1) if len(g) else None}
                            for tipo, g in (('programa', gs[~gs['generica']]), ('generico', gs[gs['generica']]), ('total', gs))}
                     for sede, gs in list(d.groupby('sede')) + [('Total', d)]},
    'por_matriz': d.groupby(['matriz', 'sede', 'programa']).agg(ra_unicos=('meso_ins_mutuo', 'size'),
                                                               con_estrategia=('meso_ins_mutuo', 'sum'))
                   .assign(V4=lambda x: (100 * x.con_estrategia / x.ra_unicos).round(1)).reset_index().to_dict('records'),
}
json.dump(EV, open('auditoria/evidencia_v4_alternativas.json', 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
print(json.dumps(EV, ensure_ascii=False, indent=1))
