"""
Evaluación de las clasificaciones automáticas del aplicativo (dashboard_tematico.py y src/).

  A. Taxonomía de RA (página «Bloom & Integración», detectar_taxonomia): dominio y subcategoría
     inferidos del texto del RA frente a lo DECLARADO en el Paso 3 (Dominio Asociado, Nivel Dominio).
  B. Nivel de formación (_detectar_nivel): Pregrado/Posgrado inferido de las columnas del Paso 5
     frente al nivel real (denominación del programa).
  C. Listas de palabras clave: config.TEMATICAS (src/thematic_detector) frente a
     config_tendencias.json (página Tendencias): coherencia entre listas.

Las funciones del dashboard se extraen con ast (sin ejecutar Streamlit) para evaluar el código real.

Uso: python auditoria/scripts/evaluar_clasificaciones_app.py
Salida: auditoria/evidencia_clasificaciones_app.json
"""
import ast
import glob
import json
import logging
import os
import re
import sys
import unicodedata
import warnings
from collections import Counter

import openpyxl
import pandas as pd

warnings.filterwarnings('ignore')
logging.disable(logging.CRITICAL)
sys.path.insert(0, '.')
sys.stdout.reconfigure(encoding='utf-8')

# --- extraer del dashboard las piezas necesarias, sin ejecutar la app
src = open('dashboard_tematico.py', encoding='utf-8').read()
tree = ast.parse(src)
NOMBRES = {'_SUBCAT_TO_DOMAIN', '_SUBCAT_ORDER', '_TAXONOMIAS_MATRIZ_PATH', '_stem_es', 'leer_taxonomias_bloom',
           '_normalize_value', '_normalize_column_name', '_find_column', '_detectar_nivel'}
piezas = [ast.get_source_segment(src, n) for n in tree.body
          if (isinstance(n, ast.FunctionDef) and n.name in NOMBRES)
          or (isinstance(n, (ast.Assign, ast.AnnAssign)) and any(
              getattr(t, 'id', None) in NOMBRES for t in (n.targets if isinstance(n, ast.Assign) else [n.target])))]
ns = {'pd': pd, 'unicodedata': unicodedata, 're': re, 'Dict': dict, 'Optional': None}
exec('from typing import Dict, Optional\n' + '\n\n'.join(piezas), ns)

# --- reconstrucción literal de verb_to_tax y detectar_taxonomia (pagina_bloom_integracion)
tax = ns['leer_taxonomias_bloom']()
verb_to_tax = {}
for subcat_norm, verbos in tax.items():
    dominio = ns['_SUBCAT_TO_DOMAIN'].get(subcat_norm, 'Cognitivo')
    orden = ns['_SUBCAT_ORDER'].get(subcat_norm, 0)
    for v in verbos:
        v_n = unicodedata.normalize('NFKD', v).encode('ascii', 'ignore').decode('ascii')
        if v_n not in verb_to_tax or orden > verb_to_tax[v_n][2]:
            verb_to_tax[v_n] = (dominio, subcat_norm, orden)
_stem_es = ns['_stem_es']


def detectar_taxonomia(texto):
    if not texto or str(texto).strip() in ('nan', ''):
        return ('No identificado', 'No identificado')
    t = unicodedata.normalize('NFKD', str(texto).lower()).encode('ascii', 'ignore').decode('ascii')
    words_t = re.findall(r'\b[a-z]{3,}\b', t)
    text_stems = [(w, _stem_es(w)) for w in words_t]
    best_orden, best_dom, best_sub = -1, 'No identificado', 'No identificado'
    for v_n, (dom, sub, orden) in verb_to_tax.items():
        if orden <= best_orden:
            continue
        if re.search(r'\b' + re.escape(v_n) + r'\b', t):
            best_orden, best_dom, best_sub = orden, dom, sub
            continue
        v_stem = _stem_es(v_n)
        if len(v_stem) < 4:
            continue
        for word, word_stem in text_stems:
            if word_stem == v_stem or word.startswith(v_stem) or v_n.startswith(word_stem):
                best_orden, best_dom, best_sub = orden, dom, sub
                break
    return (best_dom, best_sub)


def norm(t):
    return unicodedata.normalize('NFKD', str(t)).encode('ascii', 'ignore').decode('ascii').lower()


def dominio_declarado(d):
    d = norm(d)
    return 'Cognitivo' if 'cogn' in d else 'Procedimental' if 'proced' in d else 'Actitudinal' if 'actit' in d else None


# --- A. taxonomía de RA frente a lo declarado
filas = []
for f in sorted(glob.glob('data/raw/FORMATOS RA CICLO UNO RC/*.xlsx')):
    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    ws = wb[[s for s in wb.sheetnames if s.startswith('Paso 3 Redacción RA') and 'Backup' not in s][0]]
    vistos = set()
    for r in ws.iter_rows(min_row=3, values_only=True):
        if len(r) > 8 and r[8] and str(r[8]).strip():
            k = str(r[8]).strip().lower()
            if k in vistos:
                continue
            vistos.add(k)
            dom_d = dominio_declarado(r[5])
            dom_a, sub_a = detectar_taxonomia(r[8])
            filas.append({'matriz': os.path.basename(f)[10:-5], 'tipo_saber': r[2], 'dominio_declarado': dom_d,
                          'subcat_declarada': norm(re.sub(r'(B|BAK)$', '', str(r[6] or ''))).strip(),
                          'verbo_declarado': norm(r[7] or ''), 'dominio_app': dom_a, 'subcat_app': norm(sub_a), 'ra': str(r[8])[:160]})
    wb.close()
A = pd.DataFrame(filas)
val = A[A['dominio_declarado'].notna()]
conf = pd.crosstab(val['dominio_declarado'], val['dominio_app'])
EV = {'A_taxonomia': {
    'ra_unicos': len(A), 'con_dominio_declarado': len(val),
    'acuerdo_dominio': round(100 * (val['dominio_declarado'] == val['dominio_app']).mean(), 1),
    'acuerdo_subcategoria': round(100 * (val['subcat_declarada'] == val['subcat_app']).mean(), 1),
    'distribucion_declarada': val['dominio_declarado'].value_counts().to_dict(),
    'distribucion_app': val['dominio_app'].value_counts().to_dict(),
    'matriz_confusion': {i: conf.loc[i].to_dict() for i in conf.index},
    'ejemplos_desacuerdo': val[val['dominio_declarado'] != val['dominio_app']]
        [['verbo_declarado', 'dominio_declarado', 'dominio_app', 'subcat_app', 'ra']].head(10).to_dict('records'),
    'verbos_bd': len(verb_to_tax)}}

# --- B. nivel de formación
from src.extractor import ExcelExtractor  # noqa: E402
REAL = lambda c: 'Posgrado' if re.match(r'^(Esp|M[A-Z])', c) else 'Pregrado'
nivel = []
for f in sorted(glob.glob('data/raw/FORMATOS RA CICLO UNO RC/*.xlsx')):
    m = ExcelExtractor(f).extract_estrategias_micro()
    cod = os.path.basename(f)[10:-5].rsplit('_', 1)[0]
    try:
        app = ns['_detectar_nivel'](m)
    except Exception as e:  # noqa: BLE001
        app = f'error: {e}'
    nivel.append({'matriz': os.path.basename(f)[10:-5], 'real': REAL(cod), 'app': app})
B = pd.DataFrame(nivel)
EV['B_nivel'] = {'acuerdo': round(100 * (B['real'] == B['app']).mean(), 1), 'tabla': pd.crosstab(B['real'], B['app']).to_dict(),
                 'desacuerdos': B[B['real'] != B['app']].to_dict('records')}

# --- C. listas de palabras clave
from config import TEMATICAS  # noqa: E402
TEND = json.load(open('config_tendencias.json', encoding='utf-8'))['TENDENCIAS_GLOBALES']
EV['C_listas'] = {'config_TEMATICAS': sorted(TEMATICAS.keys()), 'config_tendencias_json': sorted(TEND.keys()),
                  'en_ambas': sorted(set(TEMATICAS) & set(TEND)),
                  'solo_TEMATICAS': sorted(set(TEMATICAS) - set(TEND)), 'solo_tendencias': sorted(set(TEND) - set(TEMATICAS))}
json.dump(EV, open('auditoria/evidencia_clasificaciones_app.json', 'w', encoding='utf-8'), ensure_ascii=False, indent=1, default=str)
print(json.dumps(EV, ensure_ascii=False, indent=1, default=str)[:6000])
