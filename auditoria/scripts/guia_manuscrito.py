"""
Guía para actualizar el manuscrito de la autora (Articulo Politécnico Grancolombiano.docx, OneDrive institucional)
con los textos auditados. El manuscrito no se modifica: la guía indica, por sección, qué reemplazar y con qué texto.

Uso:    python auditoria/scripts/guia_manuscrito.py
Salida: auditoria/Guia_actualizacion_manuscrito.docx
"""
import sys

from docx import Document
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, 'auditoria/scripts')
import textos_articulo as TXT  # noqa: E402

NAVY = RGBColor(0x0F, 0x38, 0x5A)
doc = Document()
st = doc.styles['Normal']
st.font.name, st.font.size = 'Times New Roman', Pt(11)
for s in doc.sections:
    s.left_margin = s.right_margin = Cm(2.2)


def titulo(t, n=1):
    p = doc.add_paragraph()
    r = p.add_run(t)
    r.bold, r.font.size, r.font.color.rgb = True, Pt(15 if n == 1 else 12), NAVY


def accion(t):
    p = doc.add_paragraph()
    r = p.add_run('Acción: ')
    r.bold = True
    p.add_run(t)


def bloque(texto, etiqueta='Texto nuevo'):
    p = doc.add_paragraph()
    r = p.add_run(f'{etiqueta}: ')
    r.bold, r.font.color.rgb = True, RGBColor(0x1F, 0xB2, 0xDE)
    p.add_run(texto)


def tabla(cab, filas):
    t = doc.add_table(rows=1, cols=len(cab))
    t.style = 'Table Grid'
    for i, c in enumerate(cab):
        t.rows[0].cells[i].paragraphs[0].add_run(c).bold = True
    for f in filas:
        cs = t.add_row().cells
        for i, v in enumerate(f):
            cs[i].text = v


titulo('Guía de actualización del manuscrito')
doc.add_paragraph('Manuscrito: «Articulo Politécnico Grancolombiano.docx» (OneDrive institucional). Fuente de los textos '
                  'auditados: auditoria/Resultados_consolidado.docx y auditoria/scripts/textos_articulo.py. Se recomienda '
                  'aplicar los cambios con control de cambios activado. El orden sigue el del manuscrito.')

titulo('1. Resumen', 2)
accion('Reemplazar las cifras de «Método» y «Resultados». Son las del borrador previo a la auditoría.')
tabla(['En el manuscrito', 'Valor auditado'],
      [['341 resultados de aprendizaje únicos', '338 (D23: deduplicación sin puntuación)'],
       ['La evaluabilidad alcanzó 98,3 %', '100 %, condición impuesta por la plantilla (R3)'],
       ['solo el 52,3 % ... ruta documentada completa', '87,0 % con estrategia e instrumento; 98,0 % sin el RA genérico (R4)'],
       ['El 10,6 % de los elementos del perfil generó alertas', '96,2 % de los 2.574 atributos con alineación en alguna capa; '
        '3,8 % sin ninguna (R1)'],
       ['el 28,9 % de las asignaturas homónimas', '9,4 % (15 de 159) divergentes (R8.1)'],
       ['solo en el 3,3 % de los registros microcurriculares', '3,8 % de las asignaturas (62 de 1.616, lista oficial D21)']])
bloque(TXT.RESUMEN[0][1], 'Texto auditado de «Método» y «Resultados»')

titulo('2. Materiales y Métodos — Variables', 2)
accion('B2: retirar «umbrales de cobertura por campo del perfil» y «detector de valores atípicos» (retirados: D24 y '
       'medida complementaria) y sustituir la Tabla 2 por las definiciones vigentes.')
bloque('No describen el currículo sino el instrumento, y se reportan como auditoría del pipeline (P9): tasa de rechazo del '
       'filtrado de núcleos, control de calidad de núcleos con reglas explícitas y puntaje orientativo de calidad temática.',
       'B2, primera frase')
p = doc.add_paragraph()
p.add_run('Tabla 2. Definición operativa de las variables de análisis').italic = True
tabla(['Variable', 'Definición operativa', 'Categoría de origen'],
      [['V1. Correspondencia del perfil', 'Proporción de atributos del perfil de egreso con alineación en al menos una capa '
        'de la matriz (competencias, RA o asignaturas); se informa además la alineación con las tres capas a la vez (D28).',
        'A. Cobertura de las competencias'],
       ['V2. Coherencia horizontal', 'Proporción de matrices sin verbo repetido entre competencias específicas; la '
        'competencia genérica institucional no cuenta (D14).', 'A. Redundancia informacional'],
       ['V3. Evaluabilidad', 'Proporción de RA únicos con verbo observable (en SaberSer se admite el verbo afectivo) y '
        'finalidad de desempeño o producto declarado (D15).', 'A. Observabilidad del resultado de aprendizaje'],
       ['V4. Trazabilidad', 'Proporción de RA únicos vinculados a una estrategia mesocurricular con indicador e instrumento '
        '(Paso 4) (D16).', 'B1. Cobertura de perfil'],
       ['V5. Nivel de evidencia de los indicadores', 'Proporción de estrategias con al menos un indicador de aprendizaje '
        'demostrado (N3) o de transferencia (N4); descriptiva, no exigida por la plantilla (D17, D26).',
        'A. Consistencia de los indicadores de logro']])

titulo('3. Procedimiento, Resultados, Discusión y Observaciones', 2)
accion('Reemplazar completas estas secciones por las de auditoria/Resultados_consolidado.docx: desde «Procedimiento» '
       '(Etapas 1–7, Tablas 3–6) hasta el final de «4. Discusión», y añadir «Observaciones» antes de las Referencias. '
       'Las del manuscrito corresponden a la versión previa a la auditoría (p. ej., V1 como trazabilidad RA–competencia, '
       'Tabla 14 con n = 1.466, Figura 10 «ausencia de sesgo disciplinar», puntaje de calidad 0–100 en la Etapa 3 y umbrales '
       'en la Tabla 5).')
tabla(['Sección del manuscrito', 'Qué cambia en la versión auditada'],
      [['Etapa 2', 'Sin detector de valores atípicos; control de núcleos con reglas (Rahm & Do, 2000; Tukey, 1977); 6.684 núcleos.'],
       ['Etapa 3', 'Valoración por criterios Cumple/Parcial/No cumple (Glaser, 1963) en lugar del puntaje 0–100.'],
       ['Etapa 4 y Tabla 5', 'Método de asociación del perfil (V1); la cobertura TF-IDF/BM25 queda como medida complementaria '
        'de saberes y valor agregado; Tabla 5 = atributos por componente.'],
       ['Etapa 7 y Tabla 6', 'Informes por matriz programa-sede; Python 3.10 (despliegue); sin Isolation Forest.'],
       ['3. Resultados', 'Tabla 7 por sede; Figura 1 = ruta de alineación curricular (Biggs, 1996); R1–R8 con las cifras '
        'auditadas y frase de contexto en cada apartado.'],
       ['4. Discusión', 'Versión auditada 4.1–4.5 (D22).'],
       ['Observaciones', 'Textos retirados por no estar respaldados en el proyecto (P17, P20, P21, P22).']])

titulo('4. Conclusiones', 2)
accion('Actualizar las cifras del primer y tercer párrafo.')
tabla(['En el manuscrito', 'Texto auditado'],
      [['examinar 50 matrices, 1.466 elementos de perfil y 5.780 núcleos temáticos',
        'examinar 50 matrices, 2.574 atributos del perfil y 6.684 núcleos temáticos'],
       ['El 10,6 % de los elementos examinados generó alertas de posible falta de respaldo, concentradas principalmente en '
        'las áreas y poblaciones de actuación.',
        'El 96,2 % de los atributos del perfil tiene alineación con al menos una capa curricular; los 98 sin alineación '
        '(3,8 %) se concentran en las poblaciones de actuación (10,9 %).']])

titulo('5. Referencias', 2)
accion('Aplicar los cambios de la lista de referencias.')
for b in TXT.REFERENCIAS_CAMBIOS:
    doc.add_paragraph(b[1], style='List Bullet')

doc.save('auditoria/Guia_actualizacion_manuscrito.docx')
print('OK auditoria/Guia_actualizacion_manuscrito.docx')
