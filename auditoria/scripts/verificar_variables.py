"""
Auditoría independiente - Sección "Variables" (Tabla 2: V1-V5).

Recalcula cada variable según su definición operativa literal en el artículo y
la compara con la salida del proyecto (data/output/recalculo_2026.xlsx), para
la que no existe código en el repositorio.

  V1  % de RA únicos cuya competencia (Paso 3, col. A) corresponde a una
      competencia del Paso 2. El vínculo competencia→perfil no existe en la matriz.
  V2  % de competencias cuyo verbo rector (Paso 2, col. B) se repite en la matriz
      (definición del artículo) y su complemento (lo que reporta el proyecto);
      indicador binario de vacíos en los tres tipos de saber (Paso 3).
  V4  % de RA únicos con ruta hasta instrumento: RA presente en Paso 4 (similitud
      >= 0,80) y bloque de estrategia con "Instrumentos de medición" (col. F).
  V5  % de estrategias declaradas (D5) con indicador de impacto (col. D) en su
      bloque, y % de filas RA–estrategia con indicador (lo que reporta el proyecto).

Uso: python auditoria/scripts/verificar_variables.py
"""
import glob
import json
import os
import re
import statistics
import sys
import unicodedata
import warnings
from collections import Counter
from difflib import SequenceMatcher

import openpyxl
import pandas as pd

warnings.filterwarnings('ignore')
RAW = 'data/raw/FORMATOS RA CICLO UNO RC'
OUT = 'auditoria/evidencia_variables.json'


def ne(v):
    return v is not None and str(v).strip() != '' and not str(v).strip().startswith('[')


def norm(t):
    t = unicodedata.normalize('NFKD', str(t)).encode('ascii', 'ignore').decode().lower()
    return re.sub(r'[\s\.;,:]+$', '', re.sub(r'\s+', ' ', t).strip())


def hoja(wb, p):
    return wb[[s for s in wb.sheetnames if s.startswith(p) and 'Backup' not in s][0]]


def parecido(a, b):
    return a == b or a.startswith(b) or b.startswith(a) or SequenceMatcher(None, a, b).ratio() >= 0.80


def main():
    filas = []
    for f in sorted(glob.glob(RAW + '/*.xlsx')):
        nombre = os.path.basename(f)
        prog, sede = re.match(r'FormatoRA_(.+?)_([A-Z]{4})\.xlsx', nombre).groups()
        wb = openpyxl.load_workbook(f, data_only=True)

        # Paso 2: competencias (redacción col F) y verbo (col B)
        comp = [(norm(r[1]) if ne(r[1]) else '', norm(r[5]))
                for r in hoja(wb, 'Paso 2').iter_rows(min_row=3, values_only=True)
                if len(r) > 5 and ne(r[5]) and not norm(r[5]).replace(' ', '').startswith('redaccion')]
        verbos = Counter(v for v, _ in comp if v)
        rep = sum(1 for v, _ in comp if v and verbos[v] > 1)

        # Paso 3: RA únicos con su competencia (col A) y tipo de saber (col C)
        p3 = [r for r in hoja(wb, 'Paso 3 Redacción RA').iter_rows(min_row=3, values_only=True) if len(r) > 8 and ne(r[8])]
        ra = {}
        for r in p3:
            ra.setdefault(norm(r[8]), set()).add(norm(r[0]) if ne(r[0]) else '')
        tipos = {re.sub(r'[^a-z]', '', norm(r[2])) for r in p3 if ne(r[2])}
        vacio_tipo = int(not {'saber', 'saberhacer', 'saberser'} <= tipos)
        comp_txt = [c for _, c in comp]
        v1_ok = sum(any(c and any(parecido(c, x) or (len(c) > 15 and c[:25] in x) for x in comp_txt) for c in cs)
                    for cs in ra.values())

        # Paso 4: bloques de estrategia (nombre en col B); indicador (col D) e instrumento (col F) por bloque
        bloques, actual = [], None
        for r in hoja(wb, 'Paso 4').iter_rows(min_row=3, values_only=True):
            if not ne(r[0]):
                continue
            if ne(r[1]) or actual is None:
                actual = {'ra': [], 'ind': False, 'ins': False, 'filas': 0, 'filas_ind': 0}
                bloques.append(actual)
            actual['ra'].append(norm(r[0]))
            actual['filas'] += 1
            tiene_ind = len(r) > 3 and ne(r[3])
            actual['filas_ind'] += tiene_ind
            actual['ind'] |= tiene_ind
            actual['ins'] |= len(r) > 5 and ne(r[5])
        estr = [b for b in bloques if b['filas']]
        v5_estr = sum(b['ind'] for b in estr)
        filas4 = sum(b['filas'] for b in estr)
        filas4_ind = sum(b['filas_ind'] for b in estr)
        ra_con_estr = ra_con_ins = 0
        for t in ra:
            bs = [b for b in estr if any(parecido(t, x) for x in b['ra'])]
            ra_con_estr += bool(bs)
            ra_con_ins += any(b['ins'] for b in bs)

        n = len(ra)
        filas.append({'archivo': nombre, 'prog': prog, 'sede': sede, 'competencias': len(comp), 'ra_unicos': n,
                      'V1_aud': round(100 * v1_ok / n, 2),
                      'verbos_repetidos': rep, 'V2_def_articulo_pct_repetidos': round(100 * rep / len(comp), 2) if comp else None,
                      'V2_complemento': round(100 * (1 - rep / len(comp)), 2) if comp else None, 'vacio_tipo_saber': vacio_tipo,
                      'V4_sin_instrumento': round(100 * ra_con_estr / n, 2), 'V4_con_instrumento': round(100 * ra_con_ins / n, 2),
                      'estrategias': len(estr), 'V5_por_estrategia': round(100 * v5_estr / len(estr), 2) if estr else None,
                      'filas_meso': filas4, 'V5_por_fila': round(100 * filas4_ind / filas4, 2) if filas4 else None})

    df = pd.DataFrame(filas)
    rc = pd.read_excel('data/output/recalculo_2026.xlsx', 'Variables_V1_V5')
    df = df.merge(rc[['prog', 'sede', 'V1', 'V2', 'V3', 'V4', 'V5', 'verbos_rep', 'n_meso']], on=['prog', 'sede'], how='left')
    m = lambda c: round(float(df[c].mean()), 1)
    res = {
        'V1_aud': m('V1_aud'), 'V1_proyecto': m('V1'), 'V1_min_aud': float(df.V1_aud.min()),
        'V2_def_articulo': m('V2_def_articulo_pct_repetidos'), 'V2_complemento': m('V2_complemento'), 'V2_proyecto': m('V2'),
        'V2_coincide_complemento': int((df.V2_complemento.round(1) == df.V2.round(1)).sum()),
        'verbos_rep_coincide': int((df.verbos_repetidos == df.verbos_rep).sum()),
        'matrices_con_vacio_tipo_saber': int(df.vacio_tipo_saber.sum()),
        'V3_proyecto': m('V3'), 'V3_matrices_bajo_100': int((df.V3 < 100).sum()),
        'V4_sin_instrumento': m('V4_sin_instrumento'), 'V4_con_instrumento': m('V4_con_instrumento'), 'V4_proyecto': m('V4'),
        'V5_por_estrategia': m('V5_por_estrategia'), 'V5_por_fila': m('V5_por_fila'), 'V5_proyecto': m('V5'),
        'V5_por_fila_coincide': int((df.V5_por_fila.round(1) == df.V5.round(1)).sum()),
        'n_meso_coincide_filas': int((df.filas_meso == df.n_meso).sum()),
    }
    json.dump({'resumen': res, 'por_matriz': json.loads(df.to_json(orient='records', force_ascii=False))},
              open(OUT, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    sys.stdout.reconfigure(encoding='utf-8')
    print(json.dumps(res, ensure_ascii=False, indent=1))
    print(df[['prog', 'sede', 'V2_complemento', 'V2', 'V5_por_fila', 'V5_por_estrategia', 'V5', 'V4_con_instrumento', 'V4']].head(8).to_string())


if __name__ == '__main__':
    main()
