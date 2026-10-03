"""
Verificación de Resultados R2 (V2): verbos de competencia, programas con verbos
repetidos, vacíos de tipo de saber, distribución de RA por tipo de saber
(Figura 2) y rangos por matriz.

Uso: python auditoria/scripts/verificar_r2.py
Salida: auditoria/evidencia_r2.json
"""
import glob
import json
import re
import sys
import unicodedata
import warnings
from collections import Counter, defaultdict

import openpyxl

warnings.filterwarnings('ignore')
sys.stdout.reconfigure(encoding='utf-8')
ARCH = sorted(glob.glob('data/raw/FORMATOS RA CICLO UNO RC/*.xlsx'))


def norm(t):
    t = unicodedata.normalize('NFKD', str(t)).encode('ascii', 'ignore').decode('ascii').lower()
    return re.sub(r'\s+', ' ', re.sub(r'[^a-z0-9 ]', ' ', t)).strip()


def ne(v):
    return v is not None and str(v).strip() != ''


def hoja(wb, pref):
    return next(wb[n] for n in wb.sheetnames if n.strip().startswith(pref) and 'backup' not in n.lower())


def tipo(t):
    t = norm(t).replace(' ', '')
    return 'SaberSer' if 'ser' in t else 'SaberHacer' if 'hacer' in t else 'Saber' if 'saber' in t else None


verbos_total, verbo_matrices, verbo_programas = Counter(), defaultdict(set), defaultdict(set)
por_matriz, tipo_filas, tipo_unicos = [], Counter(), Counter()
for f in ARCH:
    nombre = f.split('\\')[-1].split('/')[-1]
    prog = re.sub(r'^FormatoRA_|_(PBOG|PMED|HBOG|HMED|VNAL)\.xlsx$', '', nombre)
    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    comp = []
    for r in hoja(wb, 'Paso 2').iter_rows(min_row=3, values_only=True):
        if len(r) > 5 and ne(r[5]) and not str(r[5]).strip().startswith('['):
            if norm(r[5]).startswith('redaccion') or norm(r[1] or '') == 'verbo':
                continue
            comp.append(norm(r[1]) if ne(r[1]) else '')
    cnt = Counter(v for v in comp if v)
    for v in cnt:
        verbo_matrices[v].add(nombre)
        verbo_programas[v].add(prog)
    verbos_total.update(cnt)
    rep_comp = sum(1 for v in comp if v and cnt[v] > 1)

    filas, vistos, tipos_ra = Counter(), {}, defaultdict(set)
    for r in hoja(wb, 'Paso 3').iter_rows(min_row=3, values_only=True):
        if len(r) > 8 and ne(r[8]) and ne(r[2]):
            tp = tipo(r[2])
            if tp:
                filas[tp] += 1
                tipos_ra[str(r[8]).strip().lower()].add(tp)  # clave D4
    unic = Counter()
    for ra, tps in tipos_ra.items():
        unic[sorted(tps)[0] if len(tps) == 1 else 'Varios'] += 1
    tipo_filas.update(filas)
    n = sum(filas.values())
    por_matriz.append({'archivo': nombre, 'programa': prog, 'competencias': len(comp),
                       'verbos_repetidos': sum(1 for v in cnt if cnt[v] > 1), 'competencias_con_verbo_repetido': rep_comp,
                       'tiene_repeticion': rep_comp > 0, 'tipos_presentes': len(filas),
                       'pct_saber': round(100 * filas['Saber'] / n, 1), 'pct_hacer': round(100 * filas['SaberHacer'] / n, 1),
                       'pct_ser': round(100 * filas['SaberSer'] / n, 1),
                       'ra_con_3_tipos': sum(1 for t in tipos_ra.values() if len(t) == 3), 'ra_unicos': len(tipos_ra)})
    wb.close()

N = len(por_matriz)
progs = {m['programa'] for m in por_matriz}
prog_rep = {m['programa'] for m in por_matriz if m['tiene_repeticion']}
EV = {
    'competencias': sum(m['competencias'] for m in por_matriz),
    'top_verbos': [{'verbo': v, 'n': c, 'matrices': len(verbo_matrices[v]), 'programas': len(verbo_programas[v])}
                   for v, c in verbos_total.most_common(15)],
    'matrices_con_repeticion': sum(m['tiene_repeticion'] for m in por_matriz),
    'programas_con_repeticion': f'{len(prog_rep)}/{len(progs)}',
    'matrices_con_vacio_tipo': sum(m['tipos_presentes'] < 3 for m in por_matriz),
    'tipo_filas': dict(tipo_filas),
    'rango_pct_saber': [min(m['pct_saber'] for m in por_matriz), max(m['pct_saber'] for m in por_matriz)],
    'rango_pct_hacer': [min(m['pct_hacer'] for m in por_matriz), max(m['pct_hacer'] for m in por_matriz)],
    'rango_pct_ser': [min(m['pct_ser'] for m in por_matriz), max(m['pct_ser'] for m in por_matriz)],
    'matrices_hacer_ge_35': sum(m['pct_hacer'] >= 35 for m in por_matriz),
    'ra_unicos': sum(m['ra_unicos'] for m in por_matriz),
    'ra_unicos_con_3_tipos': sum(m['ra_con_3_tipos'] for m in por_matriz),
    'por_matriz': por_matriz,
}
json.dump(EV, open('auditoria/evidencia_r2.json', 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
print(json.dumps({k: v for k, v in EV.items() if k != 'por_matriz'}, ensure_ascii=False, indent=1))
