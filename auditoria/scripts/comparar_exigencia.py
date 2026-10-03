"""
Comparativo E3-2: nivel de exigencia de los RA con el código actual
(src/analyzer.py:_get_nivel_taxonomico) frente a la escala propuesta, que
traduce cada nivel a 1-6 según su posición dentro de la progresión de su
propio dominio (descripción del artículo, Procedimiento, Paso 3).

Unidad: RA único por matriz (D4). Índice de exigencia = (nivel medio − 1) / 5 × 100.

Uso: python auditoria/scripts/comparar_exigencia.py
"""
import glob
import json
import logging
import os
import re
import statistics
import sys
import unicodedata
import warnings
from collections import Counter

import openpyxl

warnings.filterwarnings('ignore')
logging.disable(logging.CRITICAL)
sys.path.insert(0, '.')
from src.analyzer import CurricularAnalyzer  # noqa: E402

# Progresión de cada dominio, de menor a mayor exigencia
PROGRESIONES = {
    ('bloom', None): ['conocimiento', 'comprension', 'aplicacion', 'analisis', 'sintesis', 'evaluacion'],
    ('bak', 'cognitivo'): ['conocimiento', 'comprension', 'aplicacion', 'analisis', 'sintesis', 'evaluacion'],
    ('bak', 'procedimental'): ['imitacion', 'manipulacion', 'precision', 'control'],
    ('bak', 'actitudinal'): ['percepcion', 'responder', 'valorar', 'organizar', 'caracterizar'],
}


def k(t):
    return re.sub(r'[^a-z]', '', unicodedata.normalize('NFKD', str(t)).encode('ascii', 'ignore').decode().lower())


def nivel_propuesto(taxonomia, dominio, nivel):
    tax = 'bak' if 'bak' in k(taxonomia) else 'bloom'
    dom = next((d for d in ('cognitivo', 'procedimental', 'actitudinal') if d in k(dominio)), None)
    prog = PROGRESIONES[(tax, None if tax == 'bloom' else dom)]
    base = re.sub(r'(bak|b)$', '', k(nivel))
    base = 'conocimiento' if base.startswith('conocimiento') else base
    if base not in prog:
        return None
    pos = prog.index(base) + 1
    return round(1 + (pos - 1) * 5 / (len(prog) - 1), 2)


def ne(v):
    return v is not None and str(v).strip() != ''


def main():
    actual = CurricularAnalyzer.__new__(CurricularAnalyzer)
    tabla, filas = {}, []
    for f in sorted(glob.glob('data/raw/FORMATOS RA CICLO UNO RC/*.xlsx')):
        wb = openpyxl.load_workbook(f, data_only=True)
        ws = wb[[s for s in wb.sheetnames if s.startswith('Paso 3 Redacción RA') and 'Backup' not in s][0]]
        vistos = {}
        for r in ws.iter_rows(min_row=3, values_only=True):
            if len(r) > 8 and ne(r[8]):
                vistos.setdefault(k(r[8]), (str(r[4]).strip(), str(r[5]).strip(), str(r[6]).strip(), r[7]))
        n_act, n_pro = [], []
        for tax, dom, niv, verbo in vistos.values():
            a = actual._get_nivel_taxonomico(verbo, niv)
            p = nivel_propuesto(tax, dom, niv)
            n_act.append(a)
            n_pro.append(p)
            clave = (tax, dom, niv)
            tabla.setdefault(clave, {'ra': 0, 'actual': Counter(), 'propuesto': p})
            tabla[clave]['ra'] += 1
            tabla[clave]['actual'][a] += 1
        idx = lambda xs: round((statistics.mean(xs) - 1) / 5 * 100, 1)
        filas.append({'archivo': os.path.basename(f)[10:-5], 'ra_unicos': len(vistos),
                      'indice_actual': idx(n_act), 'indice_propuesto': idx(n_pro),
                      'por_defecto_actual': sum(1 for (t, d, nv, v), a in zip(vistos.values(), n_act)
                                                if a == 2 and not re.search(r'comprend|entiend', k(nv)))})
    for i, x in enumerate(sorted(filas, key=lambda x: -x['indice_actual']), 1):
        x['rank_actual'] = i
    for i, x in enumerate(sorted(filas, key=lambda x: -x['indice_propuesto']), 1):
        x['rank_propuesto'] = i
    out = {'tabla_niveles': [{'taxonomia': t, 'dominio': d, 'nivel': n, 'ra_unicos': v['ra'],
                              'nivel_actual': dict(v['actual']), 'nivel_propuesto': v['propuesto']}
                             for (t, d, n), v in sorted(tabla.items())],
           'por_matriz': filas,
           'resumen': {'indice_actual_media': round(statistics.mean(x['indice_actual'] for x in filas), 1),
                       'indice_propuesto_media': round(statistics.mean(x['indice_propuesto'] for x in filas), 1),
                       'indice_actual_rango': [min(x['indice_actual'] for x in filas), max(x['indice_actual'] for x in filas)],
                       'indice_propuesto_rango': [min(x['indice_propuesto'] for x in filas), max(x['indice_propuesto'] for x in filas)],
                       'ra_por_defecto_actual': sum(x['por_defecto_actual'] for x in filas),
                       'ra_sin_mapeo_propuesto': sum(v['ra'] for v in tabla.values() if v['propuesto'] is None)}}
    json.dump(out, open('auditoria/evidencia_exigencia.json', 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    sys.stdout.reconfigure(encoding='utf-8')
    print(json.dumps(out['resumen'], ensure_ascii=False))
    for t in out['tabla_niveles']:
        print(f"{t['taxonomia']:6}|{t['dominio']:17}|{t['nivel']:15}|{t['ra_unicos']:4}| actual {t['nivel_actual']} | propuesto {t['nivel_propuesto']}")
    print('--- matrices (ordenadas por cambio de posición)')
    for x in sorted(filas, key=lambda x: -abs(x['rank_actual'] - x['rank_propuesto'])):
        print(x)


if __name__ == '__main__':
    main()
