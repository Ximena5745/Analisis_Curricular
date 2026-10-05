"""
Etapa 2: control de calidad de los núcleos temáticos con reglas explícitas (P21; sustituye a Isolation Forest).

Aplica src.nucleos_cleaner.control_nucleos_reglas a los 6.671 núcleos (Paso 5, separación D7) y guarda
el conteo por regla y la lista de núcleos señalados.

Uso:    python auditoria/scripts/control_reglas_nucleos.py
Salida: auditoria/evidencia_control_nucleos.json, auditoria/Control_nucleos_reglas.xlsx
"""
import glob
import json
import logging
import re
import sys
import warnings

import openpyxl
import pandas as pd

warnings.filterwarnings('ignore')
logging.disable(logging.CRITICAL)
sys.path.insert(0, '.')
from src.nucleos_cleaner import (REGLAS_CONTROL_NUCLEOS, control_nucleos_reglas, limpiar_nucleo,  # noqa: E402
                                 tokenizar_nucleo_celda)

ne = lambda v: v is not None and str(v).strip() != ''
filas = []
for f in sorted(glob.glob('data/raw/FORMATOS RA CICLO UNO RC/*.xlsx')):
    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    ws = wb[[s for s in wb.sheetnames if s.startswith('Paso 5')][0]]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    j = next(i for i, h in enumerate(str(c or '').lower() for c in rows[0]) if 'cleo' in h)
    for r in rows[1:]:
        if ne(r[3]) and not re.fullmatch(r'[\d\.\s]+', str(r[3]).strip()) and j < len(r) and ne(r[j]):
            for t in tokenizar_nucleo_celda(r[j]):
                filas.append({'matriz': f.replace('\\', '/').split('/')[-1][10:-5], 'asignatura': str(r[3]).strip(),
                              'nucleo': limpiar_nucleo(t)})
    wb.close()

d = control_nucleos_reglas(pd.DataFrame(filas))
n = len(d)
q1, q3 = d['palabras'].quantile([0.25, 0.75])
res = {'nucleos': n, 'umbral_extension': q3 + 1.5 * (q3 - q1),
       'por_regla': {k: int(d[k].sum()) for k in REGLAS_CONTROL_NUCLEOS},
       'revisar': int(d['revisar'].sum()), 'pct_revisar': round(100 * d['revisar'].mean(), 1),
       'asignaturas_varios_temas': sorted({f"{m} | {a}" for m, a in d.loc[d.varios_temas, ['matriz', 'asignatura']].values}),
       'igual_al_nombre': sorted({f"{m} | {a}" for m, a in d.loc[d.igual_al_nombre, ['matriz', 'asignatura']].values})}
json.dump(res, open('auditoria/evidencia_control_nucleos.json', 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
d[d.revisar].rename(columns=REGLAS_CONTROL_NUCLEOS).to_excel('auditoria/Control_nucleos_reglas.xlsx', index=False)
sys.stdout.reconfigure(encoding='utf-8')
print(json.dumps(res, ensure_ascii=False, indent=1))
