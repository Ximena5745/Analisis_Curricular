"""
Auditoría independiente - Etapa 2: EDA de resultados de aprendizaje (RA).

Por cada matriz describe la unidad de análisis (filas vs RA únicos), la
consistencia de la numeración, los nulos en columnas clave y los vínculos
de cada RA único con Paso 4 (meso) y Paso 5 (micro). Solo lectura.

Regla de RA único: texto de la columna I normalizado (minúsculas, sin
tildes, espacios comprimidos, sin puntuación final), deduplicado DENTRO
de cada matriz.

Uso: python auditoria/scripts/eda_resultados_aprendizaje.py
"""
import glob
import json
import os
import re
import statistics
import sys
import unicodedata
import warnings
from collections import Counter, defaultdict

import openpyxl

warnings.filterwarnings('ignore')
RAW = 'data/raw/FORMATOS RA CICLO UNO RC'
OUT = 'auditoria/evidencia_eda_ra.json'
COLS_P3 = ['Competencia', 'Número', 'TipoSaber', 'SaberAsociado', 'Taxonomía',
           'Dominio', 'Nivel', 'Verbo', 'Texto RA']


def ne(v):
    return v is not None and str(v).strip() != ''


def norm(t):
    t = unicodedata.normalize('NFKD', str(t)).encode('ascii', 'ignore').decode().lower()
    return re.sub(r'[\s\.;,:]+$', '', re.sub(r'\s+', ' ', t).strip())


def hoja(wb, p, backup=False):
    c = [s for s in wb.sheetnames if s.startswith(p) and (('Backup' in s) == backup)]
    return wb[c[0]] if c else None


def tipo_norm(v):
    k = re.sub(r'[^a-z]', '', norm(v)) if ne(v) else ''
    return {'saber': 'Saber', 'saberhacer': 'SaberHacer', 'saberser': 'SaberSer'}.get(k, v if k else '(vacío)')


def main():
    filas_out, inconsist, sin_ruta, backups, detalle_ra = [], [], [], [], []
    for f in sorted(glob.glob(RAW + '/*.xlsx')):
        nombre = os.path.basename(f)
        wb = openpyxl.load_workbook(f, data_only=True)
        p3, filas_excel = [], defaultdict(list)
        for i, r in enumerate(hoja(wb, 'Paso 3 Redacción RA').iter_rows(min_row=3, values_only=True), start=3):
            if len(r) > 8 and ne(r[8]):
                p3.append(r)
                filas_excel[norm(r[8])].append(i)
        textos = defaultdict(list)
        for r in p3:
            textos[norm(r[8])].append(r)
        for j, (t, rs) in enumerate(textos.items(), 1):
            detalle_ra.append({'archivo': nombre, 'n': j, 'ra': str(rs[0][8]).strip(),
                               'numeros': ', '.join(sorted({str(r[1]).strip() for r in rs if ne(r[1])})),
                               'tipos': ', '.join(sorted({tipo_norm(r[2]) for r in rs})),
                               'filas_p3': len(rs), 'filas_excel': ', '.join(map(str, filas_excel[t])),
                               'competencias': len({norm(r[0]) for r in rs if ne(r[0])})})
        nums = {str(r[1]).strip() for r in p3 if ne(r[1])}
        nulos = {c: sum(not ne(r[j]) for r in p3) for j, c in enumerate(COLS_P3[:8])}

        # Inconsistencias de numeración: un texto con varios números, o un número con varios textos
        num_a_txt = defaultdict(set)
        for t, rs in textos.items():
            ns = {str(r[1]).strip() for r in rs if ne(r[1])}
            for x in ns:
                num_a_txt[x].add(t)
            if len(ns) > 1:
                inconsist.append({'archivo': nombre, 'tipo': 'Mismo texto, varios números',
                                  'detalle': f'Números {sorted(ns)}', 'ra': str(rs[0][8])[:150]})
        for x, ts in num_a_txt.items():
            if len(ts) > 1:
                inconsist.append({'archivo': nombre, 'tipo': 'Mismo número, varios textos',
                                  'detalle': f'Número {x}: {len(ts)} textos', 'ra': ' || '.join(t[:70] for t in ts)})

        tipos_por_ra = {t: {tipo_norm(r[2]) for r in rs} for t, rs in textos.items()}
        filas_tipo = Counter(tipo_norm(r[2]) for r in p3)
        unicos_tipo = Counter(min(ts) if len(ts) == 1 else 'Varios tipos' for ts in tipos_por_ra.values())

        p4 = Counter(norm(r[0]) for r in hoja(wb, 'Paso 4').iter_rows(min_row=3, values_only=True) if ne(r[0]))
        p5 = Counter(norm(r[1]) for r in hoja(wb, 'Paso 5').iter_rows(min_row=3, values_only=True)
                     if ne(r[1]) and not re.fullmatch(r'[\d\.\s]+', str(r[3] or '').strip()))
        con4 = {t for t in textos if p4.get(t)}
        con5 = {t for t in textos if p5.get(t)}
        for t in textos:
            if t not in con4 or t not in con5:
                sin_ruta.append({'archivo': nombre, 'ra': str(textos[t][0][8])[:150],
                                 'en_paso4': p4.get(t, 0), 'en_paso5': p5.get(t, 0),
                                 'falta': ' y '.join(x for x, ok in [('Paso 4 (meso)', t in con4), ('Paso 5 (micro)', t in con5)] if not ok)})

        bk = hoja(wb, 'Paso 3 Redacción RA', backup=True)
        if bk:
            bt = {norm(r[8]) for r in bk.iter_rows(min_row=3, values_only=True) if len(r) > 8 and ne(r[8])}
            comp = len(bt & set(textos))
            backups.append({'archivo': nombre, 'ra_en_backup': len(bt), 'compartidos_con_hoja_principal': comp,
                            'diagnostico': 'Vacía' if not bt else (
                                'Contenido ajeno al programa (posible plantilla copiada)' if comp / len(bt) < 0.5
                                else 'Versión anterior del mismo programa')})

        filas_out.append({
            'archivo': nombre, 'programa': re.match(r'FormatoRA_(.+?)_[A-Z]{4}', nombre).group(1),
            'sede': nombre[-9:-5],
            'p3_filas': len(p3), 'ra_unicos_texto': len(textos), 'ra_unicos_numero': len(nums),
            'factor_repeticion': round(len(p3) / len(textos), 2) if textos else None,
            'filas_Saber': filas_tipo['Saber'], 'filas_SaberHacer': filas_tipo['SaberHacer'],
            'filas_SaberSer': filas_tipo['SaberSer'],
            'unicos_Saber': unicos_tipo['Saber'], 'unicos_SaberHacer': unicos_tipo['SaberHacer'],
            'unicos_SaberSer': unicos_tipo['SaberSer'], 'unicos_varios_tipos': unicos_tipo['Varios tipos'],
            'nulos_columnas_clave': sum(nulos.values()),
            'nulos_detalle': '; '.join(f'{k}: {v}' for k, v in nulos.items() if v) or '—',
            'p4_filas_que_citan_ra': sum(p4.values()), 'p5_filas_que_citan_ra': sum(p5.values()),
            'p4_textos_huerfanos': sum(1 for t in p4 if t not in textos),
            'p5_textos_huerfanos': sum(1 for t in p5 if t not in textos),
            'ra_con_meso': len(con4), 'ra_con_micro': len(con5), 'ra_con_ambos': len(con4 & con5),
            'ra_sin_meso': len(textos) - len(con4), 'ra_sin_micro': len(textos) - len(con5),
        })

    tot = lambda k: sum(r[k] for r in filas_out)
    unicos = [r['ra_unicos_texto'] for r in filas_out]
    resumen = {
        'matrices': len(filas_out), 'p3_filas': tot('p3_filas'), 'ra_unicos_texto': tot('ra_unicos_texto'),
        'ra_unicos_numero': tot('ra_unicos_numero'),
        'unicos_min': min(unicos), 'unicos_mediana': statistics.median(unicos), 'unicos_max': max(unicos),
        'factor_min': min(r['factor_repeticion'] for r in filas_out),
        'factor_max': max(r['factor_repeticion'] for r in filas_out),
        'filas_por_tipo': {k: tot('filas_' + k) for k in ['Saber', 'SaberHacer', 'SaberSer']},
        'unicos_por_tipo': {k: tot('unicos_' + k) for k in ['Saber', 'SaberHacer', 'SaberSer', 'varios_tipos']},
        'nulos': tot('nulos_columnas_clave'),
        'ra_con_meso': tot('ra_con_meso'), 'ra_con_micro': tot('ra_con_micro'), 'ra_con_ambos': tot('ra_con_ambos'),
        'p4_huerfanos': tot('p4_textos_huerfanos'), 'p5_huerfanos': tot('p5_textos_huerfanos'),
        'matrices_con_inconsistencia_numeracion': len({i['archivo'] for i in inconsist}),
        'inconsistencias': len(inconsist), 'backups_ajenos': sum(b['diagnostico'].startswith('Contenido ajeno') for b in backups),
        'backups': len(backups),
    }
    json.dump({'resumen': resumen, 'por_programa': filas_out, 'inconsistencias': inconsist,
               'ra_sin_ruta': sin_ruta, 'backups': backups, 'detalle_ra': detalle_ra},
              open(OUT, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    sys.stdout.reconfigure(encoding='utf-8')
    print(json.dumps(resumen, ensure_ascii=False, indent=1))
    a = next(r for r in filas_out if r['archivo'] == 'FormatoRA_AdmonEmpresas_HMED.xlsx')
    print('control AdmonEmpresas_HMED:', a['p3_filas'], a['ra_unicos_texto'], a['p5_filas_que_citan_ra'], a['ra_sin_meso'])


if __name__ == '__main__':
    main()
