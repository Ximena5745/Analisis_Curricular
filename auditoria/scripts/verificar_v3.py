"""
V3 (evaluabilidad) con regla operativa versionada, sobre RA únicos por matriz (D4).

Estructura de redacción del RA en las matrices: VERBO + objeto + "para" + FINALIDAD + condición.
Un RA es evaluable cuando:
  1. Verbo observable: el verbo del RA (Paso 3, col. H) no es de estado mental; en RA de
     SaberSer (dominio actitudinal) se acepta el verbo de la taxonomía afectiva (D15).
  2. Producto o evidencia: el RA declara una finalidad de desempeño, es decir, la
     construcción "para + infinitivo" (o "mediante/a través de/con el fin de") cuyo
     infinitivo no es de estado mental, o nombra un producto concreto
     ("un/una estudio, propuesta, proyecto, informe, plan, prototipo…").

Uso: python auditoria/scripts/verificar_v3.py
Salida: auditoria/evidencia_v3.json
"""
import glob
import json
import os
import re
import statistics as st
import sys
import unicodedata
import warnings
from collections import Counter

import openpyxl

warnings.filterwarnings('ignore')
sys.stdout.reconfigure(encoding='utf-8')

# Verbos de estado mental (no observables)
NO_OBSERVABLES = {'conocer', 'comprender', 'entender', 'reflexionar', 'valorar', 'apreciar', 'interiorizar',
                  'sensibilizar', 'concientizar', 'percibir', 'aceptar', 'respetar', 'saber', 'creer',
                  'pensar', 'sentir', 'asumir', 'apropiar', 'apropiarse', 'familiarizarse', 'considerar'}
FINALIDAD = re.compile(r'\b(?:para|mediante|a traves de|con el fin de|con el proposito de)\s+(?:(?:la|el|los|las|su|sus)\s+)?([a-z]+(?:ar|er|ir)(?:se|lo|la|los|las)?)\b')
PRODUCTO = re.compile(r'\b(?:un|una|unos|unas)\s+(?:estudio|propuesta|proyecto|informe|plan|prototipo|modelo|diagnostico|'
                      r'producto|documento|ensayo|articulo|programa|estrategia|solucion|analisis|disen[oa]|reporte|portafolio)\b')
SEDES = r'_(PBOG|PMED|HBOG|HMED|VNAL)$'


def norm(t):
    t = unicodedata.normalize('NFKD', str(t)).encode('ascii', 'ignore').decode('ascii').lower()
    return re.sub(r'\s+', ' ', re.sub(r'[^a-z0-9 ]', ' ', t)).strip()


def infinitivo(v):
    return re.sub(r'(se|lo|la|los|las)$', '', v) if len(v) > 6 else v


def clasificar(verbo, texto, actitudinal=False):
    v = norm(verbo)
    t = norm(texto)
    # D15: en RA de SaberSer o dominio actitudinal se acepta el verbo de la taxonomía
    # afectiva (valorar, apreciar, respetar…) cuando hay finalidad observable.
    obs = v not in NO_OBSERVABLES or actitudinal
    m = FINALIDAD.search(t)
    fin = infinitivo(m.group(1)) if m else None
    prod = (fin is not None and fin not in NO_OBSERVABLES) or PRODUCTO.search(t) is not None
    motivo = [] if obs else [f'verbo no observable ({v})']
    if prod:
        pass
    elif not m:
        motivo.append('sin finalidad o producto declarado')
    elif not prod:
        motivo.append(f'finalidad no observable (para {fin})')
    return obs and prod, fin, '; '.join(motivo)


por_matriz, detalle = [], []
for f in sorted(glob.glob('data/raw/FORMATOS RA CICLO UNO RC/*.xlsx')):
    nombre = os.path.basename(f)[10:-5]
    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    ws = next(wb[x] for x in wb.sheetnames if x.strip().startswith('Paso 3') and 'backup' not in x.lower())
    unicos = {}
    for r in ws.iter_rows(min_row=3, values_only=True):
        if len(r) > 8 and r[8] not in (None, ''):
            unicos.setdefault(str(r[8]).strip().lower(), (r[7] or '', str(r[8]).strip(),  # clave D4
                              'ser' in norm(r[2] or '') or 'actitudinal' in norm(r[5] or '')))
    wb.close()
    ev = 0
    for verbo, texto, actitudinal in unicos.values():
        ok, fin, motivo = clasificar(verbo, texto, actitudinal)
        ev += ok
        detalle.append({'matriz': nombre, 'verbo': norm(verbo), 'finalidad': fin, 'evaluable': ok, 'motivo': motivo, 'ra': texto})
    sede = re.search(SEDES, nombre).group(1)
    por_matriz.append({'matriz': nombre, 'programa': re.sub(SEDES, '', nombre), 'sede': sede,
                       'ra_unicos': len(unicos), 'evaluables': ev, 'V3': round(100 * ev / len(unicos), 1)})


def resumen(filas):
    v = [m['V3'] for m in filas]
    return {'n': len(v), 'media': round(st.mean(v), 1), 'mediana': round(st.median(v), 1),
            'desv': round(st.stdev(v), 1) if len(v) > 1 else None, 'min': min(v), 'max': max(v)}


por_sede = {s: resumen([m for m in por_matriz if m['sede'] == s]) for s in ('VNAL', 'PBOG', 'PMED', 'HMED', 'HBOG')}
# Una matriz por programa (la primera por nombre de archivo) para no duplicar los multisede
una = {}
for m in por_matriz:
    una.setdefault(m['programa'], m)
n_ra = sum(m['ra_unicos'] for m in por_matriz)
n_ev = sum(m['evaluables'] for m in por_matriz)
EV = {'regla': {'no_observables': sorted(NO_OBSERVABLES), 'finalidad': FINALIDAD.pattern},
      'ra_unicos': n_ra, 'evaluables': n_ev, 'V3_global': round(100 * n_ev / n_ra, 1),
      'por_matriz_50': resumen(por_matriz), 'por_programa_39': resumen(list(una.values())),
      'por_sede': por_sede,
      'motivos': Counter(d['motivo'] for d in detalle if not d['evaluable']).most_common(),
      'verbos_principales_no_observables': Counter(d['verbo'] for d in detalle if d['verbo'] in NO_OBSERVABLES).most_common(),
      'finalidades_frecuentes': Counter(d['finalidad'] for d in detalle).most_common(15),
      'no_evaluables': [d for d in detalle if not d['evaluable']],
      'por_matriz': por_matriz}
json.dump(EV, open('auditoria/evidencia_v3.json', 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
print(json.dumps({k: v for k, v in EV.items() if k not in ('no_evaluables', 'por_matriz', 'regla')}, ensure_ascii=False, indent=1))
for d in EV['no_evaluables']:
    print(f"- [{d['matriz']}] {d['motivo']} :: {d['ra'][:200]}")
