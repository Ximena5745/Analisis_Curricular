"""
Control antes/después de la corrección P5 (src/nucleos_cleaner.py).

Aplica las funciones del proyecto (tokenizar_nucleo_celda, limpiar_nucleo,
es_nucleo_valido) a cada celda de núcleos del Paso 5 (fila de inicio de cada
asignatura, sin filas de totales) y guarda ítems, válidos, únicos y rechazos.

Uso:
    python auditoria/scripts/control_nucleos.py antes
    python auditoria/scripts/control_nucleos.py despues
"""
import glob
import json
import logging
import re
import sys
import unicodedata
import warnings
from collections import Counter

import openpyxl

warnings.filterwarnings('ignore')
logging.disable(logging.CRITICAL)
sys.path.insert(0, '.')
from src.nucleos_cleaner import es_nucleo_valido, limpiar_nucleo, tokenizar_nucleo_celda  # noqa: E402

etiqueta = sys.argv[1] if len(sys.argv) > 1 else 'antes'
ne = lambda v: v is not None and str(v).strip() != ''
clave = lambda t: re.sub(r'[^a-z0-9 ]', '', unicodedata.normalize('NFKD', re.sub(r'\s+', ' ', t).strip().lower())
                         .encode('ascii', 'ignore').decode())
items = validos = celdas = 0
unicos, motivos = Counter(), Counter()
for f in sorted(glob.glob('data/raw/FORMATOS RA CICLO UNO RC/*.xlsx')):
    wb = openpyxl.load_workbook(f, data_only=True)
    ws = wb[[s for s in wb.sheetnames if s.startswith('Paso 5')][0]]
    enc = [str(c.value or '').lower() for c in ws[2]]
    j = next(i for i, h in enumerate(enc) if 'cleo' in h)
    for r in ws.iter_rows(min_row=3, values_only=True):
        if not (ne(r[3]) and not re.fullmatch(r'[\d\.\s]+', str(r[3]).strip()) and j < len(r) and ne(r[j])):
            continue
        celdas += 1
        for t in tokenizar_nucleo_celda(r[j]):
            items += 1
            lt = limpiar_nucleo(t)
            ok, motivo = es_nucleo_valido(lt)
            if ok:
                validos += 1
                unicos[clave(lt)] += 1
            else:
                motivos[re.sub(r'\s*\(.*', '', motivo)] += 1
res = {'etiqueta': etiqueta, 'celdas': celdas, 'items': items, 'validos': validos, 'unicos': len(unicos),
       'rechazados': items - validos, 'motivos': dict(motivos.most_common())}
json.dump(res, open(f'auditoria/control_nucleos_{etiqueta}.json', 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
sys.stdout.reconfigure(encoding='utf-8')
print(json.dumps(res, ensure_ascii=False, indent=1))
