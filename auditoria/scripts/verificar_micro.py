"""
Auditoría independiente - Componente microcurricular (asignaturas y núcleos temáticos).

Recalcula desde los archivos crudos (Paso 5, solo lectura). La unidad de lectura
es el BLOQUE de asignatura: la fila con nombre en la columna D y su celda de
núcleos (celda combinada, se cuenta una sola vez). Se excluye la fila de totales.

Reproduce las dos reglas del proyecto:
  R1 analisis_tematico_avanzado.py: split [,;\\n]+, longitud > 3 (sin pasar a minúsculas)
     → produce "menciones" (7.906) y "únicos" (3.432) del artículo.
  R2 src/nucleos_cleaner.py: tokenizar_nucleo_celda (; \\n | y numeración, NO coma)
     + limpiar_nucleo + es_nucleo_valido → produce "válidos" (5.780) del artículo.

  R3 (auditoría): un núcleo = un ítem numerado ("1. …"), aunque ocupe varias líneas.

Y evalúa: efecto de cortar por coma, calidad de los rechazos, duplicación entre
sedes de un mismo programa, sensibilidad a mayúsculas y densidad por asignatura.

Uso: python auditoria/scripts/verificar_micro.py
"""
import glob
import json
import logging
import os
import re
import statistics
import sys
import unicodedata
import warnings
from collections import Counter, defaultdict

import openpyxl

warnings.filterwarnings('ignore')
logging.disable(logging.CRITICAL)
sys.path.insert(0, '.')
from src.nucleos_cleaner import es_nucleo_valido, limpiar_nucleo, tokenizar_nucleo_celda  # noqa: E402

RAW = 'data/raw/FORMATOS RA CICLO UNO RC'
OUT = 'auditoria/evidencia_micro.json'


def ne(v):
    return v is not None and str(v).strip() != ''


def sin_tildes(t):
    return unicodedata.normalize('NFKD', t).encode('ascii', 'ignore').decode()


NUM = re.compile(r'(?m)^\s*(\d+(?:\.\d+)*)[\.\)]?\s*(?=[^\d\s])')


def nucleos_por_numeracion(txt):
    """R3 (auditoría): cada núcleo empieza con su número al inicio de línea ("1. …", "2.…").
    Un núcleo puede ocupar varias líneas. Si la celda no está numerada, cada línea es un núcleo."""
    t = str(txt).strip()
    ms = list(NUM.finditer(t))
    if not ms or ms[0].start() > 3:
        return [l.strip() for l in t.split('\n') if l.strip()], False
    out = []
    for a, b in zip(ms, ms[1:] + [None]):
        it = re.sub(r'\s+', ' ', t[a.end():b.start() if b else len(t)]).strip()
        if it:
            out.append(it)
    return out, True


def clave(t):
    return re.sub(r'[^a-z0-9 ]', '', sin_tildes(re.sub(r'\s+', ' ', t).strip().lower()))


def main():
    asig_raw, por_matriz = [], []
    r1 = Counter()
    r2_val, r2_motivos = Counter(), Counter()
    ej_rech = defaultdict(list)
    ej_coma = []
    una_palabra_tematica = 0
    val_por_matriz = {}
    densidad_r1, densidad_r2, densidad_r3 = [], [], []
    r3_unicos, r3_rech, ej_r3 = Counter(), Counter(), defaultdict(list)
    sin_num = 0
    for f in sorted(glob.glob(RAW + '/*.xlsx')):
        nombre = os.path.basename(f)
        prog = re.match(r'FormatoRA_(.+?)_[A-Z]{4}', nombre).group(1)
        wb = openpyxl.load_workbook(f, data_only=True)
        ws = wb[[s for s in wb.sheetnames if s.startswith('Paso 5')][0]]
        enc = [sin_tildes(str(c.value)).lower().strip() if ne(c.value) else '' for c in ws[2]]
        j = next(i for i, h in enumerate(enc) if h.startswith('nucleo'))
        n_asig = celdas = m1 = t2 = v2 = n3 = 0
        val_set = set()
        for r in ws.iter_rows(min_row=3, values_only=True):
            d = r[3]
            es_total = ne(d) and re.fullmatch(r'[\d\.\s]+', str(d).strip())
            if not ne(d) or es_total:
                continue  # los núcleos se declaran en la fila de inicio del bloque (celda combinada)
            asig_raw.append(str(d))
            n_asig += 1
            if not (j < len(r) and ne(r[j])):
                continue
            celdas += 1
            txt = str(r[j])
            p1 = [x.strip() for x in re.split(r'[,;\n]+', txt) if x.strip() and len(x.strip()) > 3]
            r1.update(p1)
            m1 += len(p1)
            densidad_r1.append(len(p1))
            r3, numerada = nucleos_por_numeracion(txt)
            n3 += len(r3)
            sin_num += not numerada
            r3_unicos.update(clave(x) for x in r3)
            for x in r3:
                ok3, mot3 = es_nucleo_valido(x)
                if not ok3:
                    m3 = re.sub(r'\s*\(.*', '', mot3)
                    r3_rech[m3] += 1
                    if len(ej_r3[m3]) < 8:
                        ej_r3[m3].append({'archivo': nombre, 'asignatura': str(d).strip(), 'nucleo': x})
            densidad_r3.append(len(r3))
            toks = tokenizar_nucleo_celda(txt)
            t2 += len(toks)
            nv = 0
            for t in toks:
                lt = limpiar_nucleo(t)
                ok, motivo = es_nucleo_valido(lt)
                if ok:
                    nv += 1
                    r2_val[lt] += 1
                    val_set.add(clave(lt))
                    if ',' in lt and len(ej_coma) < 8:
                        ej_coma.append({'archivo': nombre, 'nucleo': lt,
                                        'como_lo_corta_R1': ' | '.join(x.strip() for x in lt.split(',') if len(x.strip()) > 3)})
                else:
                    m = re.sub(r'\s*\(.*', '', motivo)
                    r2_motivos[m] += 1
                    if len(ej_rech[m]) < 6:
                        ej_rech[m].append({'archivo': nombre, 'texto': lt})
                    if m == 'Menos de 2 palabras' and re.fullmatch(r'[A-Za-zÁÉÍÓÚáéíóúÑñ]{5,}', lt.strip()):
                        una_palabra_tematica += 1
            v2 += nv
            densidad_r2.append(nv)
        val_por_matriz[nombre] = (prog, val_set)
        por_matriz.append({'archivo': nombre, 'asignaturas': n_asig, 'celdas_nucleos': celdas,
                           'asig_sin_nucleos': n_asig - celdas, 'menciones_R1': m1, 'tokens_R2': t2,
                           'validos_R2': v2, 'validos_unicos_R2': len(val_set), 'nucleos_R3': n3})

    # Duplicación entre sedes: núcleos válidos (por clave) de matrices hermanas
    por_prog = defaultdict(list)
    for nombre, (prog, st) in val_por_matriz.items():
        por_prog[prog].append(st)
    multisede = []
    for prog, sets in por_prog.items():
        if len(sets) > 1:
            inter = set.intersection(*sets)
            multisede.append({'programa': prog, 'matrices': len(sets),
                              'pct_compartido': round(100 * len(inter) / max(1, len(set.union(*sets))), 1)})
    nuc_matriz = sum(len(s) for _, s in val_por_matriz.values())
    nuc_programa = sum(len(set.union(*s)) for s in por_prog.values())

    norm = {
        'texto exacto (strip)': {a.strip() for a in asig_raw},
        'minúsculas': {re.sub(r'\s+', ' ', a).strip().lower() for a in asig_raw},
        'minúsculas sin tildes': {sin_tildes(re.sub(r'\s+', ' ', a).strip().lower()) for a in asig_raw},
        'minúsculas sin tildes ni puntuación': {clave(a) for a in asig_raw},
    }
    res = {
        'registros_asignatura': len(asig_raw),
        'denominaciones': {k: len(v) for k, v in norm.items()},
        'celdas_nucleos': sum(p['celdas_nucleos'] for p in por_matriz),
        'asig_sin_nucleos': sum(p['asig_sin_nucleos'] for p in por_matriz),
        'R1_menciones': sum(r1.values()), 'R1_unicos': len(r1),
        'R1_unicos_sin_mayusculas': len({clave(k) for k in r1}),
        'R2_tokens': sum(p['tokens_R2'] for p in por_matriz), 'R2_validos': sum(r2_val.values()),
        'R2_validos_unicos_texto': len(r2_val), 'R2_validos_unicos_normalizados': len({clave(k) for k in r2_val}),
        'R2_rechazados': sum(r2_motivos.values()), 'R2_rechazos_por_motivo': dict(r2_motivos.most_common()),
        'rechazos_una_palabra_posible_tema': una_palabra_tematica,
        'validos_con_coma': sum(v for k, v in r2_val.items() if ',' in k),
        'densidad_R1_media': round(statistics.mean(densidad_r1), 1),
        'densidad_R2_media': round(statistics.mean(densidad_r2), 1),
        'densidad_R2_mediana': statistics.median(densidad_r2),
        'densidad_R2_max': max(densidad_r2),
        'R3_nucleos': sum(densidad_r3), 'R3_unicos': len(r3_unicos), 'R3_celdas_sin_numeracion': sin_num,
        'R3_densidad_media': round(statistics.mean(densidad_r3), 1), 'R3_densidad_mediana': statistics.median(densidad_r3),
        'R3_densidad_max': max(densidad_r3),
        'R3_rechazados_por_filtro_R2': sum(r3_rech.values()), 'R3_rechazos_por_motivo': dict(r3_rech.most_common()),
        'nucleos_unicos_sumando_matrices': nuc_matriz,
        'nucleos_unicos_sumando_programas': nuc_programa,
        'multisede': multisede,
    }
    json.dump({'resumen': res, 'por_matriz': por_matriz, 'ejemplos_rechazo': ej_rech, 'ejemplos_rechazo_R3': ej_r3, 'ejemplos_coma': ej_coma},
              open(OUT, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    sys.stdout.reconfigure(encoding='utf-8')
    print(json.dumps(res, ensure_ascii=False, indent=1))
    for k, v in ej_r3.items():
        print('R3', k, [e['nucleo'] for e in v[:6]])
    for k, v in ej_rech.items():
        print(k, [e['texto'] for e in v[:5]])
    print([e['nucleo'] for e in ej_coma[:4]])


if __name__ == '__main__':
    main()
