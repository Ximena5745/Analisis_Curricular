"""
Asociación del perfil ocupacional con el contenido del programa (componentes esenciales).

Componentes (Paso 1): Áreas profesionales (H), Tareas profesionales (I), Poblaciones de actuación (J).
Ítem = cada entrada de la celda, separada por salto de línea, «/», «;» o viñeta (se quitan
numeraciones y conteos entre paréntesis). La competencia genérica no aplica (no hay perfil ocupacional genérico).

Capas de contenido de la misma matriz, en unidades finas:
  Competencias  Paso 2: redacción de cada competencia y su condición de contexto
  RA            Paso 3: texto de cada RA único (D4)
  Asignaturas   Paso 5: nombre de cada asignatura, cada núcleo temático (separado por su numeración, D7)
                e indicadores de logro; se conserva la asignatura de origen
Similitud: coseno entre embeddings multilingües (sentence-transformers,
paraphrase-multilingual-MiniLM-L12-v2); puntaje del ítem en una capa = máxima similitud con sus unidades.

Umbral por capa (índice de Youden): positivos = componentes del perfil profesional frente a su propia matriz
(vínculo formal por lista desplegable); negativos = los mismos frente a un programa de OTRO campo amplio CINE-F.
Se informa además la tasa de falsos positivos de los ítems ocupacionales con ese umbral.

Uso: python auditoria/scripts/asociacion_perfil_ocupacional.py
Salida: auditoria/evidencia_perfil_ocupacional.json y auditoria/Asociacion_perfil_ocupacional.xlsx
"""
import glob
import json
import logging
import os
import re
import sys
import unicodedata
import warnings

import numpy as np
import openpyxl
import pandas as pd

warnings.filterwarnings('ignore')
logging.disable(logging.CRITICAL)
sys.path.insert(0, '.')
sys.stdout.reconfigure(encoding='utf-8')
from sentence_transformers import SentenceTransformer  # noqa: E402

from src.extractor import ExcelExtractor  # noqa: E402
from src.nucleos_cleaner import es_nucleo_valido, tokenizar_nucleo_celda  # noqa: E402

MODELO = 'paraphrase-multilingual-MiniLM-L12-v2'
CAMPOS = {7: 'Áreas profesionales', 8: 'Tareas profesionales', 9: 'Poblaciones actuación'}
CAPAS = ('Competencias', 'RA', 'Asignaturas')


def norm(t):
    t = unicodedata.normalize('NFKD', str(t)).encode('ascii', 'ignore').decode('ascii').lower()
    return re.sub(r'\s+', ' ', re.sub(r'[^a-z0-9 ]', ' ', t)).strip()


def ne(v):
    return v is not None and str(v).strip() not in ('', 'nan', 'None') and not str(v).strip().startswith('[')


def items(celda):
    out = []
    for linea in re.split(r'\n+|;|•', str(celda)):
        partes = re.split(r'/', linea)
        # «/» separa ítems solo si cada parte es una expresión (≥ 2 palabras); si no, es una alternativa
        # dentro del mismo ítem («Diseño/elaboración de guías»)
        if len(partes) > 1 and not all(len(x.split()) >= 2 for x in partes):
            partes = [linea]
        for p in partes:
            p = re.sub(r'\(\s*\d+\s*\)', '', p)
            p = re.sub(r'^\s*(\d+[.)]|[-–—*])\s*', '', p).strip(' .,:-\t')
            if len(p) >= 4 and re.search(r'[A-Za-zÁÉÍÓÚáéíóúñÑ]{3}', p):
                out.append(p)
    return out


def hoja(wb, pref):
    return next(wb[s] for s in wb.sheetnames if s.strip().startswith(pref) and 'backup' not in s.lower())


def col(df, prefijo):
    return next((c for c in df.columns if norm(c).startswith(prefijo)), None)


modelo = SentenceTransformer(MODELO)
emb = lambda textos: modelo.encode(list(textos), batch_size=64, normalize_embeddings=True, show_progress_bar=False)

MATRICES = []
for f in sorted(glob.glob('data/raw/FORMATOS RA CICLO UNO RC/*.xlsx')):
    matriz = os.path.basename(f)[10:-5]
    programa, sede = matriz.rsplit('_', 1)
    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    ocup, prof = [], []
    for r in hoja(wb, 'Paso1').iter_rows(min_row=3, values_only=True):
        r = list(r) + [None] * 12
        for j, campo in CAMPOS.items():
            if ne(r[j]):
                ocup += [(campo, i) for i in items(r[j])]
        if ne(r[4]) and 'fenomenos contemporaneos' not in norm(r[4]):
            prof += [(c, str(r[j]).strip()) for j, c in ((4, 'Saber'), (5, 'SaberHacer'), (6, 'SaberSer')) if ne(r[j])]
    unidades = {c: [] for c in CAPAS}
    for r in hoja(wb, 'Paso 2').iter_rows(min_row=3, values_only=True):
        if len(r) > 5 and ne(r[5]) and not norm(r[5]).startswith('redaccion') and 'fenomenos contemporaneos' not in norm(r[5]):
            unidades['Competencias'].append((str(r[5]).strip(), 'Competencia'))
            if ne(r[4]):
                unidades['Competencias'].append((str(r[4]).strip(), 'Condición de contexto'))
    vistos = set()
    for r in hoja(wb, 'Paso 3').iter_rows(min_row=3, values_only=True):
        if len(r) > 8 and ne(r[8]) and norm(r[8]) != 'resultados aprendizaje' and str(r[8]).strip().lower() not in vistos:
            vistos.add(str(r[8]).strip().lower())
            unidades['RA'].append((str(r[8]).strip(), 'RA'))
    wb.close()
    micro = ExcelExtractor(f).extract_estrategias_micro()
    c_as, c_nu, c_in = col(micro, 'nombre asignatura'), col(micro, 'nucleos tematicos'), col(micro, 'indicadores de logro')
    asig = micro[c_as].ffill() if c_as else pd.Series('', index=micro.index)
    for i, row in micro.iterrows():
        a = str(asig[i]).strip()
        if not ne(a) or norm(a).startswith('electiva'):
            continue
        if ne(row.get(c_as)):
            unidades['Asignaturas'].append((a, a))
        if c_nu and ne(row.get(c_nu)):
            unidades['Asignaturas'] += [(n, a) for n in tokenizar_nucleo_celda(row[c_nu]) if es_nucleo_valido(n)[0]]
        if c_in and ne(row.get(c_in)):
            unidades['Asignaturas'].append((str(row[c_in]).strip(), a))
    unidades = {c: list(dict.fromkeys(u)) for c, u in unidades.items()}
    MATRICES.append({'matriz': matriz, 'programa': programa, 'sede': sede, 'ocup': ocup, 'prof': prof,
                     'unidades': unidades, 'E': {c: emb(t for t, _ in u) if u else None for c, u in unidades.items()}})
    print('ok', matriz, len(ocup), {c: len(u) for c, u in unidades.items()})

# Control negativo: matriz siguiente cuyo programa pertenece a otro campo amplio CINE-F 2013
ws = openpyxl.load_workbook('auditoria/Clasificacion_estandar_oferta.xlsx', read_only=True)['CINE-F_programas']
CINE = {str(r[0]).strip(): r[2] for r in list(ws.iter_rows(values_only=True))[1:] if r[0]}
for i, m in enumerate(MATRICES):
    m['otra'] = next(MATRICES[(i + k) % len(MATRICES)] for k in range(1, len(MATRICES))
                     if CINE.get(MATRICES[(i + k) % len(MATRICES)]['programa']) != CINE.get(m['programa']))


def mejor(e_item, E):
    if E is None or not len(E):
        return None, -1
    s = E @ e_item
    j = int(np.argmax(s))
    return float(s[j]), j


filas, negativos, POS, NEG = [], {}, {c: [] for c in CAPAS}, {c: [] for c in CAPAS}
for m in MATRICES:
    e_oc = emb(t for _, t in m['ocup']) if m['ocup'] else []
    for (campo, texto), e in zip(m['ocup'], e_oc):
        fila = {'Matriz': m['matriz'], 'Programa': m['programa'], 'Sede': m['sede'], 'Componente': campo, 'Ítem': texto}
        for capa in CAPAS:
            s, j = mejor(e, m['E'][capa])
            fila[f'{capa} similitud'] = None if s is None else round(s, 3)
            if s is not None:
                u_txt, u_src = m['unidades'][capa][j]
                fila[f'{capa} unidad más afín'] = u_txt[:200] if capa != 'Asignaturas' else f'{u_src}: {u_txt[:160]}'
            sn, _ = mejor(e, m['otra']['E'][capa])
            if sn is not None:
                negativos.setdefault((campo, capa), []).append(sn)
        filas.append(fila)
    if m['prof']:
        for e in emb(t for _, t in m['prof']):
            for capa in CAPAS:
                sp, sn = mejor(e, m['E'][capa])[0], mejor(e, m['otra']['E'][capa])[0]
                if sp is not None and sn is not None:
                    POS[capa].append(sp)
                    NEG[capa].append(sn)

# Umbral por capa = máximo índice de Youden entre positivos (componentes del perfil profesional frente a
# su propia matriz, vínculo formal) y negativos (los mismos frente a un programa de otro campo CINE-F)
REFERENCIA, UMB_CAPA = {}, {}
for capa in CAPAS:
    pos, neg = np.array(POS[capa]), np.array(NEG[capa])
    j, u = max(((pos >= u).mean() - (neg >= u).mean(), u) for u in np.round(np.arange(0.30, 0.95, 0.005), 3))
    UMB_CAPA[capa] = float(u)
    REFERENCIA[capa] = {'umbral': float(u), 'youden': round(float(j), 3), 'sensibilidad': round(float((pos >= u).mean()), 3),
                        'falsos_positivos': round(float((neg >= u).mean()), 3), 'n': int(len(pos)),
                        'mediana_positivos': round(float(np.median(pos)), 3), 'mediana_negativos': round(float(np.median(neg)), 3)}
# Falsos positivos esperados en el perfil ocupacional (mismo ítem frente a otro campo) con ese umbral
UMBRALES = {f'{campo} | {capa}': UMB_CAPA[capa] for (campo, capa) in negativos}
FP_OCUP = {f'{campo} | {capa}': round(float((np.array(v) >= UMB_CAPA[capa]).mean()), 3) for (campo, capa), v in negativos.items()}
d = pd.DataFrame(filas)
for capa in CAPAS:
    d[f'{capa} asociado'] = [None if pd.isna(s) else bool(s >= UMBRALES[f'{c} | {capa}'])
                             for c, s in zip(d['Componente'], d[f'{capa} similitud'])]
asoc = [f'{c} asociado' for c in CAPAS]
d['Capas asociadas'] = d[asoc].fillna(False).astype(int).sum(axis=1)
d['Estado'] = np.select([d['Capas asociadas'] == 0, d['Competencias asociado'].fillna(False).astype(bool)],
                        ['Sin asociación', 'Asociado a competencia'], 'Solo en RA o asignaturas')


def resumen(g):
    out = {'items': int(len(g))}
    for c in CAPAS:
        v = g[f'{c} asociado'].dropna().astype(float)
        out[c] = round(100 * v.mean(), 1) if len(v) else None
    out['sin_asociacion'] = int((g['Capas asociadas'] == 0).sum())
    out['pct_sin_asociacion'] = round(100 * (g['Capas asociadas'] == 0).mean(), 1)
    return out


EV = {'modelo': MODELO, 'umbrales_por_capa': UMB_CAPA, 'calibracion_perfil_profesional': REFERENCIA, 'falsos_positivos_ocupacional': FP_OCUP,
      'items': len(d), 'total': resumen(d), 'por_componente': {c: resumen(g) for c, g in d.groupby('Componente')},
      'por_matriz': {m: resumen(g) for m, g in d.groupby('Matriz')}}
json.dump(EV, open('auditoria/evidencia_perfil_ocupacional.json', 'w', encoding='utf-8'), ensure_ascii=False, indent=1)

pm = pd.DataFrame([{'Matriz': m, **v} for m, v in EV['por_matriz'].items()]).sort_values('pct_sin_asociacion', ascending=False)
with pd.ExcelWriter('auditoria/Asociacion_perfil_ocupacional.xlsx') as xw:
    pd.DataFrame([{'Componente': 'Total', **EV['total']}] + [{'Componente': c, **v} for c, v in EV['por_componente'].items()]
                 ).to_excel(xw, sheet_name='Resumen', index=False)
    pm.to_excel(xw, sheet_name='Por_matriz', index=False)
    d.sort_values(['Matriz', 'Componente', 'Capas asociadas']).to_excel(xw, sheet_name='Items', index=False)
    pd.DataFrame([{'Capa': c, **v, 'FP ocupacional': ', '.join(f"{k.split(' | ')[0]} {x:.0%}" for k, x in FP_OCUP.items() if k.endswith(c))}
                  for c, v in REFERENCIA.items()]).to_excel(xw, sheet_name='Calibracion', index=False)
    pd.DataFrame({'Elemento': ['Ítem', 'Capas', 'Similitud', 'Umbral', 'Estado', 'Referencia', 'Limitación'],
                  'Definición': [
                      'Entrada de Áreas profesionales, Tareas profesionales o Poblaciones de actuación (Paso 1), separada por línea, «/», «;» o viñeta.',
                      'Competencias (redacción y condición de contexto, Paso 2); RA únicos (Paso 3); asignaturas (nombre, núcleos e indicadores, Paso 5; sin electivas).',
                      f'Coseno entre embeddings multilingües ({MODELO}); máximo sobre las unidades de la capa.',
                      'Por capa, máximo índice de Youden: componentes del perfil profesional frente a su matriz (positivos, vínculo formal) y frente a un programa de otro campo CINE-F (negativos).',
                      'Asociado a competencia; Solo en RA o asignaturas; Sin asociación.',
                      'Hoja Calibracion: sensibilidad y falsos positivos por capa; FP ocupacional = ítems ocupacionales que superan el umbral frente a otro campo.',
                      'Mide cercanía de significado, no suficiencia formativa; la unidad más afín permite la validación del comité curricular.']}
                 ).to_excel(xw, sheet_name='Definiciones', index=False)
print(json.dumps({k: v for k, v in EV.items() if k != 'por_matriz'}, ensure_ascii=False, indent=1))
