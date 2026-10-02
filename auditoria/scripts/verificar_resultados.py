"""
Auditoría independiente - Etapa 3: Resultados (A7-A11).

Recalcula desde los archivos crudos, con las unidades adoptadas
(D1: 50 programas; D4: RA únicos por matriz; D6: 1.757 asignaturas):
  - V4 trazabilidad: % de RA únicos que aparecen en el Paso 4 (exacto y
    por similitud >= 0,80, como declara el artículo).
  - IA: presencia por programa y por registro de asignatura, con
    coincidencia de palabra completa (lista de config_tendencias.json y
    lista estricta), frente a la coincidencia por subcadena del código.
Además resume la salida existente data/output/recalculo_2026.xlsx.

Uso: python auditoria/scripts/verificar_resultados.py
"""
import glob
import json
import os
import re
import statistics
import sys
import unicodedata
import warnings
from difflib import SequenceMatcher

import openpyxl
import pandas as pd

warnings.filterwarnings('ignore')
RAW = 'data/raw/FORMATOS RA CICLO UNO RC'
OUT = 'auditoria/evidencia_resultados.json'

IA_CONFIG = json.load(open('config_tendencias.json', encoding='utf-8'))['TENDENCIAS_GLOBALES']['INTELIGENCIA_ARTIFICIAL']['keywords']
IA_ESTRICTA = ['inteligencia artificial', 'ia', 'machine learning', 'aprendizaje automatico', 'deep learning',
               'aprendizaje profundo', 'redes neuronales', 'chatbot', 'nlp', 'procesamiento de lenguaje natural',
               'vision artificial', 'ia generativa']


def ne(v):
    return v is not None and str(v).strip() != ''


def norm(t):
    t = unicodedata.normalize('NFKD', str(t)).encode('ascii', 'ignore').decode().lower()
    return re.sub(r'[\s\.;,:]+$', '', re.sub(r'\s+', ' ', t).strip())


def hoja(wb, p):
    return wb[[s for s in wb.sheetnames if s.startswith(p) and 'Backup' not in s][0]]


def patron(kws):
    return re.compile(r'\b(' + '|'.join(re.escape(norm(k)) for k in kws) + r')\b')


P_CONF, P_EST = patron(IA_CONFIG), patron(IA_ESTRICTA)
P_SUB = [norm(k) for k in IA_CONFIG]


def main():
    prog, asig_rows, ejemplos_sub = [], [], []
    for f in sorted(glob.glob(RAW + '/*.xlsx')):
        nombre = os.path.basename(f)
        wb = openpyxl.load_workbook(f, data_only=True)
        # --- V4: RA únicos del Paso 3 presentes en el Paso 4
        ra = {norm(r[8]) for r in hoja(wb, 'Paso 3 Redacción RA').iter_rows(min_row=3, values_only=True)
              if len(r) > 8 and ne(r[8])}
        p4 = {norm(r[0]) for r in hoja(wb, 'Paso 4').iter_rows(min_row=3, values_only=True) if ne(r[0])}
        exact = sum(t in p4 for t in ra)
        fuzzy = sum(t in p4 or any(SequenceMatcher(None, t, x).ratio() >= 0.80 for x in p4) for t in ra)

        # --- IA: bloques de asignatura del Paso 5 (nombre en col D; filas siguientes hasta la próxima asignatura)
        ws = hoja(wb, 'Paso 5')
        enc = [norm(c.value) if ne(c.value) else '' for c in ws[2]]
        j_nuc = next((j for j, h in enumerate(enc) if h.startswith('nucleo')), None)
        bloques, actual = [], None
        for r in ws.iter_rows(min_row=3, values_only=True):
            d = r[3]
            if ne(d) and re.fullmatch(r'[\d\.\s]+', str(d).strip()):
                continue  # fila de totales
            if ne(d):
                actual = {'asig': str(d).strip(), 'txt': []}
                bloques.append(actual)
            if actual is not None:
                for j in [1, 3, 4] + ([j_nuc] if j_nuc is not None else []):
                    if j < len(r) and ne(r[j]):
                        actual['txt'].append(str(r[j]))
        n_conf = n_est = n_sub = 0
        for b in bloques:
            t = norm(' '.join(b['txt']))
            c, e = bool(P_CONF.search(t)), bool(P_EST.search(t))
            s = any(k in t for k in P_SUB)
            n_conf += c
            n_est += e
            n_sub += s
            if s and not c and len(ejemplos_sub) < 15:
                k = next(k for k in P_SUB if k in t)
                i = t.find(k)
                ejemplos_sub.append({'archivo': nombre, 'asignatura': b['asig'], 'keyword': k,
                                     'contexto': t[max(0, i - 40):i + 40]})
            if e:
                asig_rows.append({'archivo': nombre, 'asignatura': b['asig'], 'termino': P_EST.search(t).group(0)})
        prog.append({'archivo': nombre, 'ra_unicos': len(ra), 'ra_en_p4_exacto': exact, 'ra_en_p4_similitud': fuzzy,
                     'v4_exacto': round(100 * exact / len(ra), 2), 'v4_similitud': round(100 * fuzzy / len(ra), 2),
                     'asignaturas': len(bloques), 'ia_config': n_conf, 'ia_estricta': n_est, 'ia_subcadena': n_sub})

    df = pd.DataFrame(prog)
    N = df.asignaturas.sum()
    res = {
        'programas': len(df), 'asignaturas': int(N), 'ra_unicos': int(df.ra_unicos.sum()),
        'v4_global_exacto': round(100 * df.ra_en_p4_exacto.sum() / df.ra_unicos.sum(), 1),
        'v4_global_similitud': round(100 * df.ra_en_p4_similitud.sum() / df.ra_unicos.sum(), 1),
        'v4_media_prog_exacto': round(df.v4_exacto.mean(), 1), 'v4_media_prog_similitud': round(df.v4_similitud.mean(), 1),
        'v4_min_similitud': df.v4_similitud.min(), 'v4_max_similitud': df.v4_similitud.max(),
        'ia_prog_config': int((df.ia_config > 0).sum()), 'ia_reg_config': int(df.ia_config.sum()),
        'ia_prog_estricta': int((df.ia_estricta > 0).sum()), 'ia_reg_estricta': int(df.ia_estricta.sum()),
        'ia_prog_subcadena': int((df.ia_subcadena > 0).sum()), 'ia_reg_subcadena': int(df.ia_subcadena.sum()),
    }
    rec = pd.read_excel('data/output/recalculo_2026.xlsx', sheet_name=None)
    v, c, t, b = rec['Variables_V1_V5'], rec['Consistencia_asignaturas'], rec['Tendencias'], rec['Brechas_perfil']
    ia = t[t.tendencia == 'Inteligencia artificial'].iloc[0]
    res['recalculo_2026'] = {
        'V3_media': round(v.V3.mean(), 1), 'V3_min': v.V3.min(), 'V4_media': round(v.V4.mean(), 1),
        'V4_min': v.V4.min(), 'V4_max': v.V4.max(), 'n_ra_unicos': int(v.n_ra_unicos.sum()),
        'asig_compartidas': len(c), 'divergentes': int((c.sim < 0.60).sum()),
        'pct_divergentes': round(100 * (c.sim < 0.60).mean(), 1),
        'ia_registros': int(ia.registros), 'ia_pct_reg': float(ia.pct_reg), 'ia_programas': int(ia.programas),
        'ia_pct_prog': float(ia.pct_prog), 'brechas_perfil': len(b),
    }
    json.dump({'resumen': res, 'por_programa': prog, 'ia_asignaturas': asig_rows, 'falsos_positivos_subcadena': ejemplos_sub},
              open(OUT, 'w', encoding='utf-8'), ensure_ascii=False, indent=1, default=float)
    sys.stdout.reconfigure(encoding='utf-8')
    print(json.dumps(res, ensure_ascii=False, indent=1, default=float))
    for e in ejemplos_sub[:6]:
        print(e)


if __name__ == '__main__':
    main()
