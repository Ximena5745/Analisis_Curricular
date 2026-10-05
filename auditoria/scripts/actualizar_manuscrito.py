"""
Actualiza el manuscrito de la autora con los textos auditados (Guia_actualizacion_manuscrito.docx):
Resumen, Variables (B2 y Tabla 2), Procedimiento → Discusión (+ Observaciones), Conclusiones y Referencias.

Uso:    python auditoria/scripts/actualizar_manuscrito.py <entrada.docx> <salida.docx>
Fuente: auditoria/Resultados_consolidado.docx y auditoria/scripts/textos_articulo.py.
Hacer antes una copia de respaldo del manuscrito; el script no toca las demás secciones.
"""
import copy
import io
import re
import sys
import unicodedata

import docx
from docx.oxml.ns import qn

sys.path.insert(0, 'auditoria/scripts')
import textos_articulo as TXT  # noqa: E402

ENTRADA, SALIDA = sys.argv[1], sys.argv[2]
M = docx.Document(ENTRADA)
S = docx.Document('auditoria/Resultados_consolidado.docx')
log = []


def poner_texto(p, texto):
    """Reemplaza el texto del párrafo conservando el formato de su primer run."""
    runs = p.runs
    if not runs:
        p.add_run(texto)
        return
    runs[0].text = texto
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def reemplazar_en(p, viejo, nuevo):
    if viejo in p.text:
        poner_texto(p, p.text.replace(viejo, nuevo))
        return True
    return False


# 1. Resumen ------------------------------------------------------------------------------------
P = M.paragraphs
res_nuevo = TXT.RESUMEN[0][1].split('Resultados. ', 1)[1]
for p in P[:6]:
    if reemplazar_en(p, '341 resultados de aprendizaje únicos', '338 resultados de aprendizaje únicos'):
        log.append('Resumen: 341 → 338 RA únicos')
    if p.text.startswith('Resultados.') and ' Discusión. ' in p.text:
        cola = p.text.split(' Discusión. ', 1)[1]
        poner_texto(p, f'Resultados. {res_nuevo} Discusión. {cola}')
        log.append('Resumen: párrafo de Resultados con cifras auditadas')

# 2. Variables: B2 y Tabla 2 ---------------------------------------------------------------------
for p in M.paragraphs:
    if p.text.startswith('No describen el currículo sino el instrumento'):
        poner_texto(p, 'No describen el currículo sino el instrumento, y se reportan como auditoría del pipeline (P9): tasa de '
                       'rechazo del filtrado de núcleos, control de calidad de núcleos con reglas explícitas y puntaje '
                       'orientativo de calidad temática.')
        log.append('Variables B2: sin umbrales de cobertura ni detector de valores atípicos')
TABLA2 = [
    ('V1. Correspondencia del perfil', 'Proporción de atributos del perfil de egreso con alineación en al menos una capa de la '
     'matriz (competencias, RA o asignaturas); se informa además la alineación con las tres capas a la vez.'),
    ('V2. Coherencia horizontal', 'Proporción de matrices sin verbo repetido entre competencias específicas; la competencia '
     'genérica institucional no cuenta.'),
    ('V3. Evaluabilidad', 'Proporción de RA únicos con verbo observable (en SaberSer se admite el verbo afectivo) y finalidad '
     'de desempeño o producto declarado.'),
    ('V4. Trazabilidad', 'Proporción de RA únicos vinculados a una estrategia mesocurricular con indicador e instrumento (Paso 4).'),
    ('V5. Nivel de evidencia de los indicadores', 'Proporción de estrategias con al menos un indicador de aprendizaje '
     'demostrado (N3) o de transferencia (N4); descriptiva, no exigida por la plantilla.'),
]
for t in M.tables:
    if len(t.rows) >= 6 and t.cell(1, 0).text.strip().startswith('V1.'):
        for i, (var, defi) in enumerate(TABLA2, 1):
            poner_texto(t.cell(i, 0).paragraphs[0], var)
            poner_texto(t.cell(i, 1).paragraphs[0], defi)
        log.append('Tabla 2: definiciones vigentes de V1–V5')
        break

# 3. Procedimiento → Discusión: reemplazo por la versión auditada (con Observaciones) ---------------
body = M.element.body
hijos = list(body.iterchildren())
txt = lambda e: ''.join(t.text or '' for t in e.iter(qn('w:t'))).strip()
ini = next(i for i, e in enumerate(hijos) if e.tag == qn('w:p') and txt(e) == 'Procedimiento')
fin = next(i for i, e in enumerate(hijos) if e.tag == qn('w:p') and txt(e) == 'Conclusiones')
ancla = hijos[fin]
for e in hijos[ini:fin]:
    body.remove(e)
log.append(f'Eliminados {fin - ini} elementos (Procedimiento → Discusión anteriores)')

s_hijos = list(S.element.body.iterchildren())
s_ini = next(i for i, e in enumerate(s_hijos) if e.tag == qn('w:p') and txt(e) == 'Procedimiento')
s_fin = next(i for i, e in enumerate(s_hijos) if e.tag == qn('w:p') and txt(e).startswith('Control de avance'))
while s_fin > s_ini and s_hijos[s_fin - 1].find('.//' + qn('w:br')) is not None and not txt(s_hijos[s_fin - 1]):
    s_fin -= 1
ESTILO = {}
for e in s_hijos[s_ini:s_fin]:
    t = txt(e)
    if e.tag != qn('w:p'):
        continue
    if t in ('3. Resultados', '4. Discusión', 'Observaciones'):
        ESTILO[id(e)] = 'Heading 1'
    elif t == 'Procedimiento' or re.match(r'^(R\d\. .+|4\.\d .+)$', t):
        ESTILO[id(e)] = 'Heading 2'
    elif re.match(r'^R8\.\d ', t):
        ESTILO[id(e)] = 'Heading 3'
id_estilo = {n: M.styles[n].style_id for n in ('Heading 1', 'Heading 2', 'Heading 3')}
def formato_tabla_manuscrito(tbl):
    """Formato directo de las tablas del manuscrito: bordes simples; encabezado #002060 en blanco y negrita, centrado;
    cuerpo blanco; fila «Total» #A6C9EC en negrita; Times New Roman 10 pt (el estilo TableGrid no existe allí)."""
    tblPr = tbl.find(qn('w:tblPr'))
    for x in list(tblPr):
        if x.tag in (qn('w:tblStyle'), qn('w:tblW'), qn('w:tblBorders')):
            tblPr.remove(x)
    tblPr.insert(0, tblPr.makeelement(qn('w:tblW'), {qn('w:w'): '5000', qn('w:type'): 'pct'}))
    for i, tr in enumerate(tbl.findall(qn('w:tr'))):
        tcs = tr.findall(qn('w:tc'))
        primera = ''.join(t.text or '' for t in tcs[0].iter(qn('w:t'))).strip() if tcs else ''
        tipo = 'cab' if i == 0 else ('total' if primera == 'Total' else 'cuerpo')
        fondo, color, negrita = {'cab': ('002060', 'FFFFFF', True), 'total': ('A6C9EC', '000000', True),
                                 'cuerpo': ('FFFFFF', '000000', False)}[tipo]
        for tc in tcs:
            tcPr = tc.find(qn('w:tcPr'))
            if tcPr is None:
                tcPr = tc.makeelement(qn('w:tcPr'), {})
                tc.insert(0, tcPr)
            for x in list(tcPr):
                if x.tag in (qn('w:tcBorders'), qn('w:shd'), qn('w:vAlign')):
                    tcPr.remove(x)
            bordes = tcPr.makeelement(qn('w:tcBorders'), {})
            for lado in ('top', 'left', 'bottom', 'right'):
                bordes.append(bordes.makeelement(qn(f'w:{lado}'), {qn('w:val'): 'single', qn('w:sz'): '4',
                                                                   qn('w:space'): '0', qn('w:color'): 'auto'}))
            tcPr.append(bordes)
            tcPr.append(tcPr.makeelement(qn('w:shd'), {qn('w:val'): 'clear', qn('w:color'): '000000', qn('w:fill'): fondo}))
            tcPr.append(tcPr.makeelement(qn('w:vAlign'), {qn('w:val'): 'center'}))
            for p in tc.findall(qn('w:p')):
                if tipo == 'cab':
                    pPr = p.find(qn('w:pPr'))
                    if pPr is None:
                        pPr = p.makeelement(qn('w:pPr'), {})
                        p.insert(0, pPr)
                    for x in pPr.findall(qn('w:jc')):
                        pPr.remove(x)
                    jc = pPr.makeelement(qn('w:jc'), {qn('w:val'): 'center'})
                    rpr_p = pPr.find(qn('w:rPr'))
                    (rpr_p.addprevious if rpr_p is not None else pPr.append)(jc)  # jc va antes de rPr
                for r in p.findall(qn('w:r')):
                    viejo = r.find(qn('w:rPr'))
                    cursiva = viejo is not None and viejo.find(qn('w:i')) is not None
                    if viejo is not None:
                        r.remove(viejo)
                    rPr = r.makeelement(qn('w:rPr'), {})  # orden del esquema: rFonts, b, bCs, i, iCs, color, sz, szCs
                    rPr.append(rPr.makeelement(qn('w:rFonts'), {qn('w:ascii'): 'Times New Roman', qn('w:hAnsi'): 'Times New Roman',
                                                                qn('w:eastAsia'): 'Times New Roman', qn('w:cs'): 'Times New Roman'}))
                    if negrita:
                        rPr.append(rPr.makeelement(qn('w:b'), {}))
                        rPr.append(rPr.makeelement(qn('w:bCs'), {}))
                    if cursiva:
                        rPr.append(rPr.makeelement(qn('w:i'), {}))
                        rPr.append(rPr.makeelement(qn('w:iCs'), {}))
                    rPr.append(rPr.makeelement(qn('w:color'), {qn('w:val'): color}))
                    rPr.append(rPr.makeelement(qn('w:sz'), {qn('w:val'): '20'}))
                    rPr.append(rPr.makeelement(qn('w:szCs'), {qn('w:val'): '20'}))
                    r.insert(0, rPr)


doc_pr = 5000
n_img = 0
for e in s_hijos[s_ini:s_fin]:
    nuevo = copy.deepcopy(e)
    if id(e) in ESTILO:
        pPr = nuevo.find(qn('w:pPr'))
        if pPr is None:
            pPr = nuevo.makeelement(qn('w:pPr'), {})
            nuevo.insert(0, pPr)
        pStyle = pPr.find(qn('w:pStyle'))
        if pStyle is None:
            pStyle = pPr.makeelement(qn('w:pStyle'), {})
            pPr.insert(0, pStyle)
        pStyle.set(qn('w:val'), id_estilo[ESTILO[id(e)]])
        for r in nuevo.iter(qn('w:rPr')):  # el estilo de título manda sobre el formato directo
            for x in list(r):
                if x.tag in (qn('w:rFonts'), qn('w:sz'), qn('w:b')):
                    r.remove(x)
    if nuevo.tag == qn('w:tbl'):
        formato_tabla_manuscrito(nuevo)
    for blip in nuevo.iter(qn('a:blip')):
        rid = blip.get(qn('r:embed'))
        parte = S.part.related_parts[rid]
        nuevo_rid, _ = M.part.get_or_add_image(io.BytesIO(parte.blob))
        blip.set(qn('r:embed'), nuevo_rid)
        n_img += 1
    for dp in nuevo.iter(qn('wp:docPr')):
        doc_pr += 1
        dp.set('id', str(doc_pr))
    ancla.addprevious(nuevo)
log.append(f'Insertados {s_fin - s_ini} elementos auditados ({n_img} figuras)')

# 4. Conclusiones ---------------------------------------------------------------------------------
CONC = [('examinar 50 matrices, 1.466 elementos de perfil y 5.780 núcleos temáticos',
         'examinar 50 matrices, 2.574 atributos del perfil y 6.684 núcleos temáticos'),
        ('El 10,6 % de los elementos examinados generó alertas de posible falta de respaldo, concentradas principalmente en las '
         'áreas y poblaciones de actuación.',
         'El 96,2 % de los atributos del perfil tiene alineación con al menos una capa curricular; los 98 sin alineación '
         '(3,8 %) se concentran en las poblaciones de actuación (10,9 %).'),
        ('Estas alertas no prueban la ausencia de formación asociada', 'Estos vacíos no prueban la ausencia de formación asociada')]
for p in M.paragraphs:
    for a, b in CONC:
        if reemplazar_en(p, a, b):
            log.append(f'Conclusiones: «{a[:45]}…»')

# 5. Referencias -----------------------------------------------------------------------------------
def clave(t):
    t = unicodedata.normalize('NFKD', t).encode('ascii', 'ignore').decode().lower()
    return re.sub(r'[^a-z0-9 ]', '', t)


P = M.paragraphs
i_ref = next(i for i, p in enumerate(P) if p.text.strip() == 'Referencias')
refs = [p for p in P[i_ref + 1:] if p.text.strip()]
for b in TXT.REFERENCIAS_CAMBIOS:
    t = b[1]
    if t.startswith('Se retira: '):
        ref = t[len('Se retira: '):].split(' — ')[0]
        autor = clave(ref.split('(')[0])[:25]
        anio = re.search(r'\((\d{4})\)', ref).group(1)
        for p in refs:
            if clave(p.text).startswith(autor) and f'({anio})' in p.text:
                p._p.getparent().remove(p._p)
                refs.remove(p)
                log.append(f'Referencia retirada: {ref[:50]}')
                break
    elif t.startswith('Se agrega: '):
        ref = t[len('Se agrega: '):]
        autor = clave(ref.split('(')[0])[:20]
        anio = re.search(r'\((\d{4})\)', ref).group(1)
        if any(clave(p.text).startswith(autor) and f'({anio})' in p.text for p in refs):
            continue
        sig = next((p for p in refs if clave(p.text) > clave(ref)), None)
        modelo = refs[0]
        nuevo = copy.deepcopy(modelo._p)
        (sig._p.addprevious if sig is not None else refs[-1]._p.addnext)(nuevo)
        par = docx.text.paragraph.Paragraph(nuevo, modelo._parent)
        poner_texto(par, ref)
        refs.insert(refs.index(sig) if sig is not None else len(refs), par)
        log.append(f'Referencia agregada: {ref[:50]}')

M.save(SALIDA)
print('\n'.join(log))
print('OK', SALIDA)
