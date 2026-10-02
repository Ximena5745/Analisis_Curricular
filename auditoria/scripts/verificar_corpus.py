"""
Auditoría independiente - Sección "Corpus" y Tabla 1 del artículo.

Recalcula por matriz y por nivel de formación: competencias (Paso 2), filas y
RA únicos (Paso 3), filas meso y estrategias declaradas (Paso 4), asignaturas,
denominaciones y núcleos temáticos (Paso 5). También mide cuántas matrices
multisede son idénticas entre sedes (independencia de las unidades).

Nivel de formación: se infiere del prefijo del nombre del archivo
(Esp → Especialización; M* de la lista → Maestría; Tecnol → Tecnología;
TecProf → Técnica profesional; resto → Profesional universitario).

Uso: python auditoria/scripts/verificar_corpus.py
"""
import glob
import json
import os
import re
import sys
import unicodedata
import warnings
from collections import Counter, defaultdict

import openpyxl

warnings.filterwarnings('ignore')
RAW = 'data/raw/FORMATOS RA CICLO UNO RC'
OUT = 'auditoria/evidencia_corpus.json'
MAESTRIAS = {'MCE', 'MEDSTEM', 'MGEM', 'MGerTalentoHumano', 'MInnovacionEducativa'}


def ne(v):
    return v is not None and str(v).strip() != ''


def norm(t):
    t = unicodedata.normalize('NFKD', str(t)).encode('ascii', 'ignore').decode().lower()
    return re.sub(r'[\s\.;,:]+$', '', re.sub(r'\s+', ' ', t).strip())


def hoja(wb, p):
    return wb[[s for s in wb.sheetnames if s.startswith(p) and 'Backup' not in s][0]]


def nivel(prog):
    if prog.startswith('Esp'):
        return 'Especialización'
    if prog in MAESTRIAS:
        return 'Maestría'
    if prog.startswith('Tecnol'):
        return 'Tecnología'
    if prog.startswith('TecProf'):
        return 'Técnica profesional'
    return 'Profesional universitario'


def main():
    filas, firmas, comp_detalle = [], defaultdict(list), []
    nuc_global, asig_global = Counter(), Counter()
    for f in sorted(glob.glob(RAW + '/*.xlsx')):
        nombre = os.path.basename(f)
        prog, sede = re.match(r'FormatoRA_(.+?)_([A-Z]{4})\.xlsx', nombre).groups()
        wb = openpyxl.load_workbook(f, data_only=True)

        # Paso 2: competencias = filas con redacción (col F) desde la fila 3, sin textos de instrucción
        for i, r in enumerate(hoja(wb, 'Paso 2').iter_rows(min_row=3, values_only=True), start=3):
            f_ = r[5] if len(r) > 5 else None
            if not ne(f_) and not (len(r) > 0 and ne(r[0])):
                continue
            if not ne(f_):
                motivo = 'No: columna F vacía'
            elif str(f_).strip().startswith('['):
                motivo = 'No: texto de instrucción de la plantilla'
            elif norm(f_).replace(' ', '').startswith('redaccion'):
                motivo = 'No: encabezado repetido'
            else:
                motivo = 'Sí'
            comp_detalle.append({'archivo': nombre, 'nivel': nivel(prog), 'fila_excel': i,
                                 'col_A_numero': str(r[0]).strip() if ne(r[0]) else '',
                                 'col_F_competencia': str(f_).strip() if ne(f_) else '',
                                 'tipo': str(r[6]).strip() if len(r) > 6 and ne(r[6]) else '',
                                 'cuenta': motivo})
        comp = [norm(r[5]) for r in hoja(wb, 'Paso 2').iter_rows(min_row=3, values_only=True)
                if len(r) > 5 and ne(r[5]) and not str(r[5]).strip().startswith('[')
                and not norm(r[5]).replace(' ', '').startswith('redaccion')]  # encabezado repetido en la fila 3
        # Paso 3
        ra_f = [norm(r[8]) for r in hoja(wb, 'Paso 3 Redacción RA').iter_rows(min_row=3, values_only=True)
                if len(r) > 8 and ne(r[8])]
        # Paso 4
        p4 = list(hoja(wb, 'Paso 4').iter_rows(min_row=3, values_only=True))
        meso_f = sum(ne(r[0]) for r in p4)
        meso_e = sum(ne(r[0]) and ne(r[1]) for r in p4)
        # Paso 5
        ws = hoja(wb, 'Paso 5')
        enc = [norm(c.value) if ne(c.value) else '' for c in ws[2]]
        j_nuc = next((j for j, h in enumerate(enc) if h.startswith('nucleo')), None)
        asig, celdas_nuc, nuc_split = [], 0, []
        for r in ws.iter_rows(min_row=3, values_only=True):
            if ne(r[3]) and not re.fullmatch(r'[\d\.\s]+', str(r[3]).strip()):
                asig.append(re.sub(r'\s+', ' ', str(r[3]).strip()))
            if j_nuc is not None and j_nuc < len(r) and ne(r[j_nuc]):
                celdas_nuc += 1
                nuc_split += [p.strip() for p in re.split(r'[\n;•]+', str(r[j_nuc])) if len(p.strip()) > 2]
        asig_global.update({a.lower() for a in asig})
        nuc_global.update(norm(x) for x in nuc_split)

        firmas[prog].append((sede, tuple(sorted(set(ra_f))), tuple(sorted(a.lower() for a in asig))))
        filas.append({'archivo': nombre, 'programa': prog, 'sede': sede, 'nivel': nivel(prog),
                      'competencias': len(comp), 'competencias_unicas': len(set(comp)),
                      'ra_filas': len(ra_f), 'ra_unicos': len(set(ra_f)),
                      'meso_filas': meso_f, 'meso_estrategias': meso_e,
                      'asignaturas': len(asig), 'col_nucleos': enc[j_nuc] if j_nuc is not None else '(no encontrada)',
                      'celdas_nucleos': celdas_nuc, 'nucleos_menciones': len(nuc_split)})

    # Independencia: ¿las matrices multisede son idénticas?
    identicas = []
    for prog, lst in firmas.items():
        if len(lst) > 1:
            ra_ig = len({x[1] for x in lst}) == 1
            as_ig = len({x[2] for x in lst}) == 1
            identicas.append({'programa': prog, 'sedes': ', '.join(x[0] for x in lst),
                              'ra_identicos': ra_ig, 'asignaturas_identicas': as_ig})

    orden = ['Profesional universitario', 'Tecnología', 'Técnica profesional', 'Especialización', 'Maestría']
    campos = ['competencias', 'ra_filas', 'ra_unicos', 'meso_filas', 'meso_estrategias', 'asignaturas']
    tabla = {}
    for nv in orden:
        sub = [x for x in filas if x['nivel'] == nv]
        tabla[nv] = {'matrices': len(sub), **{c: sum(x[c] for x in sub) for c in campos}}
    for nombre_g, grupo in [('Pregrado', orden[:3]), ('Posgrado', orden[3:]), ('Total', orden)]:
        tabla[nombre_g] = {k: sum(tabla[n][k] for n in grupo) for k in tabla[orden[0]]}

    res = {
        'tabla_por_nivel': tabla,
        'denominaciones_asignatura_unicas': len(asig_global),
        'nucleos_celdas': sum(x['celdas_nucleos'] for x in filas),
        'nucleos_menciones_split': sum(x['nucleos_menciones'] for x in filas),
        'nucleos_unicos_split': len(nuc_global),
        'multisede': identicas,
        'modalidades': dict(Counter({'PBOG': 'Presencial', 'PMED': 'Presencial', 'VNAL': 'Virtual',
                                     'HMED': 'Híbrido', 'HBOG': 'Híbrido'}[x['sede']] for x in filas)),
        'ciudades': dict(Counter({'PBOG': 'Bogotá', 'HBOG': 'Bogotá', 'PMED': 'Medellín', 'HMED': 'Medellín',
                                  'VNAL': 'Nacional (virtual)'}[x['sede']] for x in filas)),
    }
    json.dump({'resumen': res, 'por_matriz': filas, 'competencias_detalle': comp_detalle}, open(OUT, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    sys.stdout.reconfigure(encoding='utf-8')
    print(json.dumps(res, ensure_ascii=False, indent=1))
    print(Counter(x['col_nucleos'] for x in filas))


if __name__ == '__main__':
    main()
