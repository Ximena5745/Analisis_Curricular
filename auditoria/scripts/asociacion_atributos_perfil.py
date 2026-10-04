"""
Asociación de los atributos del perfil de egreso (profesional y ocupacional) con el contenido de SU PROPIA matriz.

Cada matriz se procesa de forma independiente: no se usa información de otros programas, ni para comparar ni
para calibrar. Así el cálculo es el mismo al cargar un solo Excel en el dashboard.

Atributos: src/atributos_perfil.extraer_atributos_detalle (forma breve), Paso 1 columnas C y D.
Entidades de contenido de la matriz (sin la competencia genérica institucional ni electivas):
  Competencias  Paso 2: redacción y condición de contexto de cada competencia
  RA            Paso 3: RA únicos (D4)
  Asignaturas   Paso 5: nombre, cada núcleo temático (D7) e indicadores de cada asignatura
Puntaje atributo–entidad: máxima similitud coseno (embeddings multilingües, paraphrase-multilingual-MiniLM-L12-v2)
entre el atributo y las unidades de texto de la entidad.

Umbral por capa, calibrado DENTRO de la matriz con sus vínculos formales (índice de Youden):
  componente del perfil (Saber/SaberHacer/SaberSer del Paso 1) frente a
    Competencias  la competencia derivada de él (Objeto ← Saber, Finalidad ← SaberHacer, Condición ← SaberSer) = positivo;
                  las demás competencias = negativos
    RA            los RA cuyo SaberAsociado es ese componente = positivos; los demás RA = negativos
    Asignaturas   las asignaturas que declaran en el Paso 5 esos RA = positivos; las demás = negativos
Un atributo está asociado a una entidad si su puntaje ≥ umbral de la capa; la amplitud es el número de entidades.

Uso: python auditoria/scripts/asociacion_atributos_perfil.py [prefijo_programa ...]   (sin argumentos: todas)
Salida: auditoria/Asociacion_atributos_perfil.xlsx y auditoria/evidencia_asociacion_atributos.json
"""
import glob
import json
import logging
import os
import re
import sys
import unicodedata
import warnings
from difflib import SequenceMatcher

import numpy as np
import openpyxl
import pandas as pd

warnings.filterwarnings('ignore')
logging.disable(logging.CRITICAL)
os.environ.setdefault('HF_HUB_OFFLINE', '1')
sys.path.insert(0, '.')
sys.stdout.reconfigure(encoding='utf-8')
from sentence_transformers import SentenceTransformer  # noqa: E402

from src.atributos_perfil import extraer_atributos_detalle, lexico_de_matriz  # noqa: E402
from src.extractor import ExcelExtractor  # noqa: E402
from src.nucleos_cleaner import es_nucleo_valido, tokenizar_nucleo_celda  # noqa: E402

MODELO = 'paraphrase-multilingual-MiniLM-L12-v2'
CAPAS = ('Competencias', 'RA', 'Asignaturas')
TIPOS = ('Saber', 'SaberHacer', 'SaberSer')
PROGRAMAS = tuple(sys.argv[1:])
modelo = SentenceTransformer(MODELO)


def emb(textos):
    textos = list(textos)
    return modelo.encode(textos, batch_size=64, normalize_embeddings=True, show_progress_bar=False) if textos else np.zeros((0, 384))


def norm(t):
    t = unicodedata.normalize('NFKD', str(t)).encode('ascii', 'ignore').decode('ascii').lower()
    return re.sub(r'\s+', ' ', re.sub(r'[^a-z0-9 ]', ' ', t)).strip()


def ne(v):
    return v is not None and str(v).strip() not in ('', 'nan', 'None') and not str(v).strip().startswith('[')


def igual(a, b):
    a, b = norm(a), norm(b)
    return bool(a) and (a == b or SequenceMatcher(None, a, b, autojunk=False).ratio() >= 0.8)


def hoja(wb, pref):
    return next(wb[s] for s in wb.sheetnames if s.strip().startswith(pref) and 'backup' not in s.lower())


def col(df, prefijo):
    return next((c for c in df.columns if norm(c).startswith(prefijo)), None)


def leer(f):
    """Entidades, unidades de texto y vínculos formales de una matriz."""
    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    lex = lexico_de_matriz(wb)
    p1 = [list(r) + [None] * 12 for r in hoja(wb, 'Paso1').iter_rows(min_row=3, values_only=True)]
    perfiles = {'Perfil profesional': p1[0][2] if p1 else '', 'Perfil ocupacional': p1[0][3] if p1 else ''}
    componentes = [(t, str(r[j]).strip()) for r in p1 if ne(r[4]) and 'fenomenos contemporaneos' not in norm(r[4])
                   for j, t in ((4, 'Saber'), (5, 'SaberHacer'), (6, 'SaberSer')) if ne(r[j])]
    competencias = []
    for r in hoja(wb, 'Paso 2').iter_rows(min_row=3, values_only=True):
        if len(r) > 5 and ne(r[5]) and not norm(r[5]).startswith('redaccion') and 'fenomenos contemporaneos' not in norm(r[5]):
            competencias.append({'nombre': str(r[5]).strip(), 'campos': {'Saber': r[2], 'SaberHacer': r[3], 'SaberSer': r[4]},
                                 'unidades': [str(r[5]).strip()] + ([str(r[4]).strip()] if ne(r[4]) else [])})
    ra = {}
    for r in hoja(wb, 'Paso 3').iter_rows(min_row=3, values_only=True):
        if len(r) > 8 and ne(r[8]) and norm(r[8]) != 'resultados aprendizaje' and 'fenomenos contemporaneos' not in norm(r[0] or ''):
            d = ra.setdefault(str(r[8]).strip().lower(), {'nombre': str(r[8]).strip(), 'asociados': [], 'unidades': [str(r[8]).strip()]})
            if ne(r[3]):
                d['asociados'].append(str(r[3]).strip())
    wb.close()
    micro = ExcelExtractor(f).extract_estrategias_micro()
    c_as, c_nu, c_in, c_ra = (col(micro, 'nombre asignatura'), col(micro, 'nucleos tematicos'),
                              col(micro, 'indicadores de logro'), col(micro, 'resultado de aprendizaje'))
    asig = micro[c_as].ffill() if c_as else pd.Series('', index=micro.index)
    asignaturas = {}
    for i, row in micro.iterrows():
        a = str(asig[i]).strip()
        if not ne(a) or norm(a).startswith('electiva'):
            continue
        d = asignaturas.setdefault(a, {'nombre': a, 'unidades': [a], 'ra': set()})
        if c_nu and ne(row.get(c_nu)):
            d['unidades'] += [n for n in tokenizar_nucleo_celda(row[c_nu]) if es_nucleo_valido(n)[0]]
        if c_in and ne(row.get(c_in)):
            d['unidades'].append(str(row[c_in]).strip())
        if c_ra and ne(row.get(c_ra)):
            d['ra'].add(str(row[c_ra]).strip())
    for d in asignaturas.values():
        d['unidades'] = list(dict.fromkeys(d['unidades']))
    entidades = {'Competencias': competencias, 'RA': list(ra.values()), 'Asignaturas': list(asignaturas.values())}
    # Vínculos formales componente → entidad
    vinc = {c: np.zeros((len(componentes), len(entidades[c])), bool) for c in CAPAS}
    for i, (t, txt) in enumerate(componentes):
        for k, comp in enumerate(competencias):
            vinc['Competencias'][i, k] = igual(txt, comp['campos'][t] or '')
        for k, r in enumerate(entidades['RA']):
            vinc['RA'][i, k] = any(igual(txt, a) for a in r['asociados'])
        ra_vinc = [r['nombre'] for k, r in enumerate(entidades['RA']) if vinc['RA'][i, k]]
        for k, a in enumerate(entidades['Asignaturas']):
            vinc['Asignaturas'][i, k] = any(igual(x, y) for x in a['ra'] for y in ra_vinc)
    return lex, perfiles, componentes, entidades, vinc


def puntajes(E_items, entidad_emb):
    """Matriz ítems × entidades: máxima similitud con las unidades de cada entidad."""
    return np.array([[float((E @ e).max()) if len(E) else 0.0 for E in entidad_emb] for e in E_items]).reshape(len(E_items), len(entidad_emb))


filas, calibraciones = [], []
for f in sorted(glob.glob('data/raw/FORMATOS RA CICLO UNO RC/*.xlsx')):
    matriz = os.path.basename(f)[10:-5]
    programa, sede = matriz.rsplit('_', 1)
    if PROGRAMAS and not programa.startswith(PROGRAMAS):
        continue
    lex, perfiles, componentes, entidades, vinc = leer(f)
    ent_emb = {c: [emb(e['unidades']) for e in entidades[c]] for c in CAPAS}
    E_comp = emb(t for _, t in componentes)
    umbral = {}
    for c in CAPAS:
        S = puntajes(E_comp, ent_emb[c])
        pos, neg = S[vinc[c]], S[~vinc[c]]
        if len(pos) and len(neg):
            j, u = max(((pos >= u).mean() - (neg >= u).mean(), u) for u in np.round(np.arange(0.30, 0.95, 0.005), 3))
        else:
            j, u = float('nan'), float('nan')
        umbral[c] = float(u)
        calibraciones.append({'Matriz': matriz, 'Capa': c, 'Umbral': u, 'Youden': round(float(j), 3),
                              'Sensibilidad': round(float((pos >= u).mean()), 3) if len(pos) else None,
                              'Falsos positivos': round(float((neg >= u).mean()), 3) if len(neg) else None,
                              'Pares positivos': int(len(pos)), 'Pares negativos': int(len(neg))})
    for perfil, texto in perfiles.items():
        atributos = extraer_atributos_detalle(texto or '', lex, 'tercera' if perfil == 'Perfil profesional' else 'infinitivo')
        E_at = emb(a['breve'] for a in atributos)
        S = {c: puntajes(E_at, ent_emb[c]) for c in CAPAS}
        for i, a in enumerate(atributos):
            fila = {'Matriz': matriz, 'Programa': programa, 'Sede': sede, 'Perfil': perfil,
                    'ID': ('PP' if perfil == 'Perfil profesional' else 'PO') + str(i + 1), 'Tipo': a['tipo'],
                    'Atributo (forma breve)': a['breve'], 'Atributo completo': a['atributo']}
            for c in CAPAS:
                s = S[c][i] if S[c].size else np.array([])
                if not len(s) or np.isnan(umbral[c]):
                    fila[f'{c}: asociado'] = None
                    continue
                k = int(np.argmax(s))
                sel = [entidades[c][j]['nombre'] for j in np.where(s >= umbral[c])[0]]
                fila[f'{c}: asociado'] = bool(sel)
                fila[f'{c}: n.º'] = len(sel)
                fila[f'{c}: más afín'] = entidades[c][k]['nombre'][:200]
                fila[f'{c}: similitud'] = round(float(s[k]), 3)
                if c == 'Asignaturas':
                    fila['Asignaturas que lo desarrollan'] = '; '.join(sel)
            n = sum(bool(fila.get(f'{c}: asociado')) for c in CAPAS)
            fila['Capas asociadas'] = n
            fila['Estado'] = ('Asociado en las tres capas' if n == 3 else 'Sin asociación' if n == 0 else
                              'Parcial: ' + ', '.join(c for c in CAPAS if fila.get(f'{c}: asociado')))
            filas.append(fila)
    print('ok', matriz, {c: umbral[c] for c in CAPAS})

d = pd.DataFrame(filas)
cal = pd.DataFrame(calibraciones)


def resumen(g):
    r = {'Atributos': len(g)}
    for c in CAPAS:
        r[f'% asociado {c}'] = round(100 * g[f'{c}: asociado'].astype(float).mean(), 1)
    r['% en las tres capas'] = round(100 * (g['Capas asociadas'] == 3).mean(), 1)
    r['Sin asociación'] = int((g['Capas asociadas'] == 0).sum())
    r['Asignaturas por atributo (mediana)'] = float(g['Asignaturas: n.º'].median())
    return pd.Series(r)


res = d.groupby(['Programa', 'Sede', 'Perfil']).apply(resumen).reset_index()
with pd.ExcelWriter('auditoria/Asociacion_atributos_perfil.xlsx') as xw:
    res.to_excel(xw, sheet_name='Resumen', index=False)
    d.to_excel(xw, sheet_name='Atributos', index=False)
    cal.to_excel(xw, sheet_name='Calibracion', index=False)
    pd.DataFrame({'Elemento': ['Alcance', 'Atributo', 'Puntaje', 'Umbral', 'Amplitud', 'Limitación'],
                  'Definición': ['Cada matriz se compara solo con su propio contenido; no se usan otros programas.',
                                 'Forma breve del atributo extraído del perfil (src/atributos_perfil.py).',
                                 f'Máxima similitud coseno ({MODELO}) entre el atributo y las unidades de texto de cada competencia, RA o asignatura.',
                                 'Por capa y matriz, máximo de Youden entre los vínculos formales del perfil (positivos) y el resto de entidades de la misma matriz (negativos).',
                                 'Número de competencias, RA o asignaturas con puntaje ≥ umbral.',
                                 'Mide cercanía de significado, no suficiencia formativa; la entidad más afín permite la validación del comité curricular.']}
                 ).to_excel(xw, sheet_name='Definiciones', index=False)
json.dump({'modelo': MODELO, 'programas': list(PROGRAMAS), 'calibracion': cal.to_dict('records'), 'resumen': res.to_dict('records')},
          open('auditoria/evidencia_asociacion_atributos.json', 'w', encoding='utf-8'), ensure_ascii=False, indent=1, default=str)
print(cal.to_string(index=False))
print(res.to_string(index=False))
