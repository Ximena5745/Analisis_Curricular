"""
Revisión de la versión consolidada del manuscrito (V2) sobre una COPIA; el original no se modifica.

Marca en la copia, con comentarios de Word:
  - rosado:    dato o afirmación con error frente a la evidencia auditada (comentario con la corrección);
  - turquesa:  forma, coherencia interna o dato por verificar (❓);
  - amarillo:  citas y referencias agregadas en formato APA 7, para revisión.
No cambia la redacción de la autora salvo para insertar las citas faltantes (en amarillo).

Uso: python auditoria/scripts/revisar_v2.py <copia.docx>
"""
import copy
import re
import sys
import unicodedata

import docx
from docx.enum.text import WD_COLOR_INDEX as C
from docx.text.run import Run

RUTA = sys.argv[1]
D = docx.Document(RUTA)
AUTOR, INI = 'Revisión auditoría', 'RA'
ROSA, TURQ, AMAR = C.PINK, C.TURQUOISE, C.YELLOW
log = []


def parrafos():
    yield from D.paragraphs
    for t in D.tables:
        for f in t.rows:
            for c in f.cells:
                yield from c.paragraphs


def buscar(frase):
    for p in parrafos():
        if frase in ''.join(r.text for r in p.runs):
            return p
    raise ValueError(f'No encontrado: {frase[:60]}')


def partir(run, k):
    """Parte el run en la posición k; devuelve el run derecho."""
    t = run.text
    nuevo = copy.deepcopy(run._r)
    run._r.addnext(nuevo)
    run.text, der = t[:k], Run(nuevo, run._parent)
    der.text = t[k:]
    return der


def aislar(p, frase, n=0):
    texto = ''.join(r.text for r in p.runs)
    ini = [m.start() for m in re.finditer(re.escape(frase), texto)][n]
    fin = ini + len(frase)
    for limite in (fin, ini):  # partir primero al final para no mover los índices del inicio
        pos = 0
        for r in p.runs:
            if pos < limite < pos + len(r.text):
                partir(r, limite - pos)
                break
            pos += len(r.text)
    sel, pos = [], 0
    for r in p.runs:
        if ini <= pos and pos + len(r.text) <= fin and r.text:
            sel.append(r)
        pos += len(r.text)
    return sel


def marcar(frase, comentario, color=ROSA, n=0):
    p = buscar(frase)
    rs = aislar(p, frase, n)
    for r in rs:
        r.font.highlight_color = color
    D.add_comment(rs, text=comentario, author=AUTOR, initials=INI)
    log.append(('error' if color == ROSA else 'forma/verificar', frase[:70]))


def insertar(ancla, nuevo, comentario):
    """Inserta texto nuevo (amarillo) justo después de 'ancla'."""
    p = buscar(ancla)
    rs = aislar(p, ancla)
    r = rs[-1]
    nr = copy.deepcopy(r._r)
    r._r.addnext(nr)
    run = Run(nr, p)
    run.text = nuevo
    run.font.highlight_color = AMAR
    D.add_comment([run], text=comentario, author=AUTOR, initials=INI)
    log.append(('cita agregada', nuevo.strip()))


# 0. Leyenda --------------------------------------------------------------------------------------------
tit = D.paragraphs[0]
ley = copy.deepcopy(tit._p)
tit._p.addprevious(ley)
from docx.text.paragraph import Paragraph  # noqa: E402
pl = Paragraph(ley, tit._parent)
for r in pl.runs[1:]:
    r._r.getparent().remove(r._r)
pl.runs[0].text = ('COPIA DE REVISIÓN (no es la versión final). Rosado: error frente a la evidencia auditada; turquesa: forma, '
                   'coherencia o dato por verificar; amarillo: cita o referencia agregada en APA 7. Cada marca tiene un comentario.')
pl.runs[0].font.highlight_color = TURQ
pl.runs[0].font.size = docx.shared.Pt(9)
pl.runs[0].bold = False
pl.style = D.styles['Normal']

# 1. Errores de datos (rosado) ----------------------------------------------------------------------
E = [
    ('341 resultados de aprendizaje únicos',
     'Error: son 338 RA únicos (D23: la deduplicación ignora mayúsculas, tildes y puntuación). El resto del artículo ya usa 338.'),
    ('con 6.671 núcleos temáticos (2.815 únicos)',
     'Error: 6.684 núcleos temáticos (2.828 únicos), tras separar también los núcleos numerados después de dos o más espacios (D25).'),
    ('Se obtuvieron 6.671 núcleos (2.815 únicos)',
     'Error: 6.684 núcleos (2.828 únicos) (D25).'),
    ('umbrales de cobertura por campo del perfil, tasa de rechazo del filtrado de núcleos, puntaje orientativo de calidad temática y detector de valores atípicos',
     'Error: los umbrales de cobertura por campo y el detector de valores atípicos ya no forman parte del método (D9, D24). '
     'Redacción sugerida: «tasa de rechazo del filtrado de núcleos, control de calidad de núcleos con reglas explícitas y '
     'puntaje orientativo de calidad temática».'),
    ('contraste que orientó la fijación de umbrales y la depuración de reglas (P6)',
     'Error: solo la asociación del perfil (Etapa 4) se calibró con revisión experta (11 matrices; acuerdo 80,1 %; κ = 0,44). El '
     'umbral de 0,60 de homónimas y homologación es operativo y no se calibró. Sugerencia: «la asociación del perfil se calibró '
     'con revisión experta (Etapa 4); los umbrales de similitud de las Etapas 5 y 6 son operativos».'),
    ('y un detector de valores atípicos señala temas anómalos frente al corpus (Liu et al., 2008)',
     'Error: el detector (Isolation Forest) no se ejecuta en el flujo y se retiró (D24). Redacción sugerida: «Un control con reglas '
     'explícitas señala para revisión, sin excluirlos, los núcleos que agrupan varios temas, repiten el nombre de la asignatura '
     'o tienen una extensión atípica, superior a Q3 + 1,5 × RIC (Rahm & Do, 2000; Tukey, 1977): 264 núcleos (3,9 %)». Retirar '
     'Liu et al. (2008) de las referencias.'),
    ('Cada programa recibe un puntaje de 0 a 100 que pondera la exigencia de sus resultados de aprendizaje (40 %), el equilibrio entre tipos de saber (30 %) y la variedad de estrategias didácticas (30 %).',
     'Error: el puntaje 0–100 con pesos 40/30/30 se retiró por no tener línea base (D26) y contradice la Etapa 7, que menciona '
     '«la valoración por criterios de la Etapa 3». Redacción sugerida: «Etapa 3. Valoración por criterios. Cada matriz se valora '
     'contra la regla que exige la plantilla en cada tramo de la ruta (V1–V5): Cumple, Parcial o No cumple, sin puntaje '
     'compuesto (Glaser, 1963)».'),
    ('la distribución por tipo de saber, sobre registros',
     'Error: la distribución por tipo de saber se calcula sobre los 338 RA únicos (R2, Figura 3), no sobre registros: contar '
     'registros sobrerrepresenta SaberSer. Corregir a «la distribución por tipo de saber, sobre RA únicos».'),
    ('aportan evidencia directa (V5 = 0,8 %)',
     'Error: V5 = 95,4 % (estrategias con indicador N2 o superior, D29), como dice el párrafo anterior. El 0,8 % es la evidencia '
     'directa del logro (N3). Corregir a «aportan evidencia directa del logro (N3; 0,8 %)».'),
    ('1.466 elementos de perfil y 5.780 núcleos temáticos',
     'Error: cifras del método anterior. Corregir a «2.574 atributos del perfil y 6.684 núcleos temáticos».'),
    ('El 10,6 % de los elementos examinados generó alertas de posible falta de respaldo, concentradas principalmente en las áreas y poblaciones de actuación.',
     'Error: cifra del método anterior. Corregir a «El 3,8 % de los atributos del perfil (98 de 2.574) no tiene alineación con '
     'ninguna capa curricular; los vacíos se concentran en las poblaciones de actuación (10,9 %)».'),
]
for f, c in E:
    marcar(f, c)

# Tabla 2: definiciones operativas desactualizadas
T2 = [
    ('V1. Correspondencia perfil–resultados de aprendizaje',
     'Error: V1 se redefinió (D28): «Proporción de atributos del perfil de egreso con alineación en al menos una capa de la matriz '
     '(competencias, RA o asignaturas)». La definición actual (RA con competencia trazable) es una condición de la plantilla. '
     'Nombre sugerido: «V1. Correspondencia del perfil», como en R1.'),
    ('V2. Tipología y frecuencia de inconsistencias',
     'Error: V2 = «Proporción de matrices sin verbo repetido entre competencias específicas; la competencia genérica institucional '
     'no cuenta» (D14). Nombre en R2: «Coherencia horizontal».'),
    ('Proporción de RA con ruta completa perfil → competencia → RA → instrumento de evaluación',
     'Error: V4 = «Proporción de RA únicos vinculados a una estrategia mesocurricular con indicador e instrumento (Paso 4)» (D16), '
     'como se mide en R4.'),
    ('V5. Declaración de indicadores',
     'Error: V5 = «Proporción de estrategias cuyo indicador de mayor nivel supera la mera implementación (N2 o superior)» (D29), '
     'como en R5. Nombre sugerido: «V5. Nivel de evidencia de los indicadores».'),
]
for f, c in T2:
    marcar(f, c)

# 2. Forma, coherencia y datos por verificar (turquesa) ----------------------------------------------
F = [
    ('63% de la oferta institucional en renovación curricular',
     '❓ Falta la fuente y la fecha de corte de este 63 %; no se pudo verificar con los datos del proyecto.'),
    ('Tabla 5. Umbrales de cobertura por campo del perfil',
     'Forma: la tabla muestra atributos por componente, no umbrales. Título sugerido: «Atributos del perfil por componente».'),
    ('Una asignatura homónima se considera divergente cuando la similitud media entre sus versiones es inferior a 0,60.',
     'Sugerencia: añadir «umbral operativo no calibrado con revisión experta», para no contradecir la limitación del método.'),
    ('Etapa 7. Reportes y consulta. .',
     'Forma: hay un punto duplicado. ❓ El nombre «CurriculoPoli» no coincide con el título del aplicativo en el código '
     '(«Analisis Tematico Microcurricular»; P22).'),
    ('El código, las versiones de bibliotecas, semillas, parámetros, diccionarios, umbrales y reglas de preprocesamiento están documentados en el repositorio del proyecto.',
     'Forma: este párrafo tiene estilo Título 1 y aparecerá como encabezado en el índice. Pasar a estilo Normal. ❓ Las versiones '
     'de las bibliotecas aún no están fijadas (P17).'),
    ('R8. Resultados de los modelos analíticos y su valor por área de decisión',
     'Coherencia: el título promete un «valor por área de decisión» que la sección no presenta (la matriz de decisiones se '
     'retiró). Título sugerido: «R8. Contenido declarado de la oferta: asignaturas compartidas, tendencias y vacíos del perfil». '
     'La introducción dice «cuatro análisis» y «tres bloques»: unificar en tres bloques (R8.1–R8.3).'),
    ('Núcleos de asignaturas compartidas por área',
     'Forma: a la tabla le falta el rótulo «Tabla 14» y su título en cursiva (el texto la cita como Tabla 14).'),
    ('R8.2 Pertinencia y tendencias de la oferta',
     'Coherencia: la pertinencia exige un referente externo (sector productivo, egresados) que no se mide. Título sugerido: '
     '«R8.2 Tópicos y tendencias de la oferta».'),
    ('R8.3 Atributos del perfil sin respaldo curricular',
     'Coherencia: el texto y la Tabla 16 usan «alineación». Título sugerido: «R8.3 Atributos del perfil sin alineación curricular».'),
    ('El respaldo se clasifica como:',
     'Coherencia: Resultados usa «alineación». Sugerencia: «La alineación se clasifica como explícita, parcial o ausente».'),
    ('D2 (V3–V5): Indicadores (resultado/impacto) como núcleo de la toma de decisiones curriculares',
     'Coherencia: R5 clasifica los indicadores por nivel de evidencia (N1–N5), no como resultado o impacto. Título sugerido: «D2 '
     '(V3–V5): Nivel de evidencia de los indicadores y toma de decisiones curriculares».'),
    ('sitúa esta capacidad en Map y Measure (NIST, 2023)',
     'APA 7: si la sigla se usará en el texto, introducirla en la primera cita: «(National Institute of Standards and Technology '
     '[NIST], 2023)» en la Introducción, y después «(NIST, 2023)».'),
]
for f, c in F:
    marcar(f, c, TURQ)

marcar('También quedan sin desarrollo curricular escenarios como casinos y cruceros',
       'Sugerencia: se omitió la frase sobre los falsos negativos verificados en los núcleos temáticos. Recomendada como '
       'salvaguarda: «En otros casos la falta de alineación se debe al emparejamiento automático: Contaduría Pública incluye '
       'núcleos de dirección y Negocios Internacionales de regímenes aduaneros, por lo que estos resultados requieren validación '
       'experta».', TURQ)

# Estilos de título aplicados a texto corrido
n_est = 0
for p in D.paragraphs:
    est = p.style.name if p.style is not None else ''
    if est.startswith(('Heading 2', 'Heading 3', 'Título 2', 'Título 3')) and len(p.text) > 160 and p.runs:
        D.add_comment(p.runs, text='Forma: párrafo de texto con estilo de título; aparecerá en el índice y en el panel de '
                                   'navegación. Pasar a estilo Normal.', author=AUTOR, initials=INI)
        n_est += 1
log.append(('forma', f'{n_est} párrafos de texto con estilo de título'))

# 3. Citas faltantes en el texto (amarillo) -------------------------------------------------------
CITAS = [
    ('Bloom clásica', ' (Bloom, 1956)', 'Cita agregada: fuente de la taxonomía de Bloom.'),
    ('aportes de Krathwohl', ' (Krathwohl et al., 1964)', 'Cita agregada: dominio afectivo de Krathwohl, Bloom y Masia.'),
    ('cobertura léxica ponderada por IDF', ' (Spärck Jones, 1972)', 'Cita agregada: fuente original del IDF.'),
    ('paraphrase-multilingual-MiniLM-L12-v2', ' (Reimers & Gurevych, 2020)', 'Cita agregada: modelo multilingüe usado en la similitud semántica.'),
    ('un κ de Cohen', ' (Cohen, 1960)', 'Cita agregada: fuente del coeficiente kappa.'),
    ('coherencia (NPMI', '; Röder et al., 2015', 'Cita agregada: medida de coherencia NPMI de los tópicos.'),
    ('seis correlaciones de Spearman', ' (Spearman, 1904)', 'Cita agregada: coeficiente de correlación por rangos.'),
    ('La prueba de Kruskal-Wallis', ' (Kruskal & Wallis, 1952)', 'Cita agregada: prueba no paramétrica entre sedes.'),
]
for a, n_, c in CITAS:
    insertar(a, n_, c)

# 4. Referencias ---------------------------------------------------------------------------------
NUEVAS = [
    ('Biggs, J. (1996). Enhancing teaching through constructive alignment. *Higher Education, 32*(3), 347–364. '
     'https://doi.org/10.1007/BF00138871', 'Agregada: se cita en Resultados (Figura 1) y faltaba en la lista.'),
    ('Bloom, B. S. (Ed.). (1956). *Taxonomy of educational objectives: The classification of educational goals. Handbook I: '
     'Cognitive domain*. David McKay.', 'Agregada: taxonomía de Bloom (Etapa 3).'),
    ('Cohen, J. (1960). A coefficient of agreement for nominal scales. *Educational and Psychological Measurement, 20*(1), '
     '37–46. https://doi.org/10.1177/001316446002000104', 'Agregada: κ de Cohen (Etapa 4).'),
    ('Glaser, R. (1963). Instructional technology and the measurement of learning outcomes: Some questions. *American '
     'Psychologist, 18*(8), 519–521. https://doi.org/10.1037/h0049294',
     'Agregada: necesaria si se adopta la corrección de la Etapa 3 (valoración por criterios).'),
    ('Krathwohl, D. R., Bloom, B. S., & Masia, B. B. (1964). *Taxonomy of educational objectives: The classification of '
     'educational goals. Handbook II: Affective domain*. David McKay.', 'Agregada: dominio afectivo (Etapa 3).'),
    ('Kruskal, W. H., & Wallis, W. A. (1952). Use of ranks in one-criterion variance analysis. *Journal of the American '
     'Statistical Association, 47*(260), 583–621. https://doi.org/10.1080/01621459.1952.10483441', 'Agregada: prueba de R7.'),
    ('Rahm, E., & Do, H. H. (2000). Data cleaning: Problems and current approaches. *IEEE Data Engineering Bulletin, 23*(4), '
     '3–13.', 'Agregada: necesaria si se adopta la corrección de la Etapa 2 (control con reglas).'),
    ('Reimers, N., & Gurevych, I. (2020). Making monolingual sentence embeddings multilingual using knowledge distillation. En '
     '*Proceedings of the 2020 Conference on Empirical Methods in Natural Language Processing (EMNLP)* (pp. 4512–4525). '
     'Association for Computational Linguistics. https://doi.org/10.18653/v1/2020.emnlp-main.365', 'Agregada: modelo semántico (Etapa 4).'),
    ('Röder, M., Both, A., & Hinneburg, A. (2015). Exploring the space of topic coherence measures. En *Proceedings of the '
     'Eighth ACM International Conference on Web Search and Data Mining* (pp. 399–408). ACM. '
     'https://doi.org/10.1145/2684822.2685324', 'Agregada: coherencia NPMI (Etapa 6).'),
    ('Spärck Jones, K. (1972). A statistical interpretation of term specificity and its application in retrieval. *Journal of '
     'Documentation, 28*(1), 11–21. https://doi.org/10.1108/eb026526', 'Agregada: IDF (Etapa 4).'),
    ('Spearman, C. (1904). The proof and measurement of association between two things. *The American Journal of Psychology, '
     '15*(1), 72–101. https://doi.org/10.2307/1412159', 'Agregada: correlación de R6.'),
    ('Tukey, J. W. (1977). *Exploratory data analysis*. Addison-Wesley.',
     'Agregada: necesaria si se adopta la corrección de la Etapa 2 (criterio Q3 + 1,5 × RIC).'),
]
k = lambda t: re.sub(r'[^a-z0-9 ]', '', unicodedata.normalize('NFKD', t).encode('ascii', 'ignore').decode().lower())  # noqa: E731
P = D.paragraphs
i_ref = next(i for i, p in enumerate(P) if p.text.strip() == 'Referencias')
refs = [p for p in P[i_ref + 1:] if p.text.strip()]
for texto, com in NUEVAS:
    sig = next((p for p in refs if k(p.text) > k(texto.replace('*', ''))), None)
    modelo = sig if sig is not None else refs[-1]
    nuevo = copy.deepcopy(modelo._p)
    (sig._p.addprevious if sig is not None else refs[-1]._p.addnext)(nuevo)
    np_ = Paragraph(nuevo, modelo._parent)
    base = copy.deepcopy(np_.runs[0]._r)
    for r in np_.runs:
        r._r.getparent().remove(r._r)
    for h in list(nuevo):
        if h.tag.endswith('}hyperlink'):
            nuevo.remove(h)
    corridos = []
    for j, trozo in enumerate(texto.split('*')):
        if not trozo:
            continue
        rr = copy.deepcopy(base)
        nuevo.append(rr)
        run = Run(rr, np_)
        run.text = trozo
        run.italic = bool(j % 2)
        run.font.highlight_color = AMAR
        corridos.append(run)
    D.add_comment(corridos, text=com + ' Verificar formato APA 7 y metadatos.', author=AUTOR, initials=INI)
    refs = [p for p in D.paragraphs[i_ref + 1:] if p.text.strip()]
    log.append(('referencia agregada', texto[:45]))

R = [
    ('Liu, F. T., Ting, K. M., & Zhou, Z.-H. (2008)', 'Retirar si se adopta la corrección de la Etapa 2: el detector no se ejecuta (D24).', ROSA),
    ('Martínez, M. (2006)', 'Referencia sin cita en el texto: citarla o retirarla.', TURQ),
    ('Robertson, S., & Zaragoza, H. (2009)', 'Referencia sin cita en el texto: BM25 figura en la Tabla 6 sin cita. Citarla allí o retirarla.', TURQ),
    ('Milan Barman, Radha Barman, Nandi, A., & Das, T. (2025)',
     'APA 7: autores en formato Apellido, Inicial. Probablemente «Barman, M., Barman, R., Nandi, A., & Das, T. (2025)», y en el '
     'texto «Barman et al. (2025)». ❓ Verificar en la fuente.', TURQ),
    ('Chen, C. (2025). Intelligent implementation',
     '❓ El DOI 10.3233/ATDE… corresponde a la serie Advances in Transdisciplinary Engineering (IOS Press), no a una revista '
     'llamada «Medical Engineering and Education». Verificar la fuente y el formato (capítulo en actas).', TURQ),
    ('Blackford, D., & Maria Selvitella, A. (2026)',
     '❓ Verificar el orden de los apellidos («Selvitella, A. M.») y completar volumen y número de artículo.', TURQ),
]
for f, c, col in R:
    marcar(f, c, col)

D.save(RUTA)
from collections import Counter  # noqa: E402
print(Counter(t for t, _ in log))
for t, x in log:
    print(f'  [{t}] {x}')
