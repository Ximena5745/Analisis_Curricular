"""
Auditoría independiente - Etapa 5 (asignaturas compartidas) y resultado A10 (28,9 %).

Reconstruye la tabla de asignaturas homónimas (data/output/recalculo_2026.xlsx,
hoja Consistencia_asignaturas: 159 asignaturas, 46 con sim < 0,60), cuyo script
no está en el repositorio, y prueba variantes de método:

  Homónima: misma denominación normalizada (minúsculas, sin tildes ni puntuación)
  en 2 o más programas (D1: 39 denominaciones), sin electivas ni filas de totales.
  Contenido de cada versión (asignatura en una matriz): núcleos temáticos (D7)
  + indicadores de logro [+ texto de los RA del bloque, variante].
  sim = coseno TF-IDF medio entre todas las versiones [o solo entre versiones de
  programas distintos, sin pares de la misma denominación de programa].

Uso: python auditoria/scripts/verificar_asignaturas_compartidas.py
"""
import glob
import itertools
import json
import os
import re
import sys
import unicodedata
import warnings
from collections import defaultdict

import numpy as np
import openpyxl
import pandas as pd
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

warnings.filterwarnings('ignore')
sys.path.insert(0, '.')
from src.nucleos_cleaner import tokenizar_nucleo_celda  # noqa: E402


def ne(v):
    return v is not None and str(v).strip() != ''


def clave(t):
    t = unicodedata.normalize('NFKD', str(t)).encode('ascii', 'ignore').decode().lower()
    return re.sub(r'\s+', ' ', re.sub(r'[^a-z0-9 ]', ' ', t)).strip()


def leer_versiones():
    versiones = []
    for f in sorted(glob.glob('data/raw/FORMATOS RA CICLO UNO RC/*.xlsx')):
        prog = re.match(r'FormatoRA_(.+?)_[A-Z]{4}', os.path.basename(f)).group(1)
        wb = openpyxl.load_workbook(f, data_only=True)
        ws = wb[[s for s in wb.sheetnames if s.startswith('Paso 5')][0]]
        enc = [clave(c.value) if ne(c.value) else '' for c in ws[2]]
        j_nuc = next(i for i, h in enumerate(enc) if h.startswith('nucleo'))
        j_ind = next(i for i, h in enumerate(enc) if h.startswith('indicadores'))
        actual = None
        for r in ws.iter_rows(min_row=3, values_only=True):
            d = r[3]
            if ne(d) and re.fullmatch(r'[\d\.\s]+', str(d).strip()):
                actual = None
                continue
            if ne(d):
                actual = {'archivo': os.path.basename(f)[10:-5], 'programa': prog, 'nombre': str(d).strip(),
                          'clave': clave(d), 'nucleos': [], 'indicadores': [], 'ra': []}
                versiones.append(actual)
            if actual is None:
                continue
            if j_nuc < len(r) and ne(r[j_nuc]):
                actual['nucleos'] += tokenizar_nucleo_celda(r[j_nuc])
            if j_ind < len(r) and ne(r[j_ind]):
                actual['indicadores'].append(str(r[j_ind]))
            if ne(r[1]):
                actual['ra'].append(str(r[1]))
    return [v for v in versiones if not v['clave'].startswith('electiva')]


def evaluar(versiones, con_ra, solo_entre_programas):
    textos = [' '.join(v['nucleos'] + v['indicadores'] + (v['ra'] if con_ra else [])) for v in versiones]
    vec = TfidfVectorizer(max_df=0.85, ngram_range=(1, 2), sublinear_tf=True)
    mat = vec.fit_transform([clave(t) for t in textos])
    grupos = defaultdict(list)
    for i, v in enumerate(versiones):
        grupos[v['clave']].append(i)
    filas = []
    for k_, idx in grupos.items():
        progs = {versiones[i]['programa'] for i in idx}
        if len(progs) < 2:
            continue
        pares = [(a, b) for a, b in itertools.combinations(idx, 2)
                 if not solo_entre_programas or versiones[a]['programa'] != versiones[b]['programa']]
        sims = [float(cosine_similarity(mat[a], mat[b])[0, 0]) for a, b in pares]
        filas.append({'asignatura': versiones[idx[0]]['nombre'], 'programas': len(progs), 'versiones': len(idx),
                      'sim': round(float(np.mean(sims)), 3)})
    df = pd.DataFrame(filas)
    return df, int((df.sim < 0.60).sum())


def main():
    versiones = leer_versiones()
    rc = pd.read_excel('data/output/recalculo_2026.xlsx', 'Consistencia_asignaturas')
    rc['clave'] = rc.asig.map(clave)
    out = {'versiones_sin_electivas': len(versiones), 'proyecto': {'asignaturas': len(rc), 'divergentes': int((rc.sim < 0.6).sum()),
                                                                  'pct': round(100 * (rc.sim < 0.6).mean(), 1)}}
    for nombre, con_ra, entre in [('V1 núcleos+indicadores, todos los pares', False, False),
                                  ('V2 + texto RA, todos los pares', True, False),
                                  ('V3 núcleos+indicadores, solo entre programas distintos', False, True),
                                  ('V4 + texto RA, solo entre programas distintos', True, True)]:
        df, div = evaluar(versiones, con_ra, entre)
        m = df.merge(rc[['clave', 'sim', 'progs', 'n']], left_on=df.asignatura.map(clave), right_on='clave', how='inner',
                     suffixes=('', '_proyecto'))
        out[nombre] = {'asignaturas': len(df), 'divergentes': div, 'pct': round(100 * div / len(df), 1),
                       'coinciden_con_proyecto': len(m),
                       'correlacion_sim_con_proyecto': round(float(m.sim.corr(m.sim_proyecto)), 3) if len(m) > 2 else None}
        if nombre.startswith('V3'):
            df.sort_values('sim').to_json('auditoria/evidencia_asignaturas_compartidas.json', orient='records',
                                          force_ascii=False, indent=1)
    sys.stdout.reconfigure(encoding='utf-8')
    print(json.dumps(out, ensure_ascii=False, indent=1))


if __name__ == '__main__':
    main()
