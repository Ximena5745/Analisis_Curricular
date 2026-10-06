"""
Taxonomías de los RA: lo que declara cada matriz (Paso 3) frente a lo que la base de verbos asigna al verbo del RA.

Unidad: RA único por matriz (D4: clave = texto del RA sin mayúsculas, tildes ni puntuación), la misma que usan V3 y el
índice de exigencia. Cada RA declara en el Paso 3 su taxonomía (Bloom o BAK), dominio y nivel; la base de verbos
(assets/taxonomias/Taxonomias_MatrizBD.xlsx, hoja «Verbos») asigna cada verbo a una o varias subcategorías.

El contraste es descriptivo: un verbo puede admitir varias lecturas, así que «difiere» pide revisión, no declara error.

  Dominio: el dominio declarado debe estar entre los dominios que la base asigna al verbo.
  Nivel:   si el dominio coincide, el nivel declarado y el de la base se comparan en la escala común de 1 a 6 de
           indicadores_articulo.nivel_exigencia (cada nivel según su posición en la progresión de su dominio).
"""

import os
import re
import unicodedata
from typing import Dict, Iterable, List, Tuple

import openpyxl
import pandas as pd

from src.indicadores_articulo import _ne, _norm, nivel_exigencia, nombre_matriz

RUTA_BD = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'assets', 'taxonomias',
                       'Taxonomias_MatrizBD.xlsx')
DOMINIOS = ('Cognitivo', 'Procedimental', 'Actitudinal')
ESTADOS = ('Coincide', 'Nivel difiere', 'Dominio difiere', 'Verbo fuera de la base')


def _k(t) -> str:
    """Clave de comparación: minúsculas, sin tildes ni signos."""
    return re.sub(r'[^a-z]', '', unicodedata.normalize('NFKD', str(t)).encode('ascii', 'ignore').decode().lower())


def dominio_de(texto) -> str:
    """'CognitivoBAK', 'ProcedimentalB', 'Cognitivo()' → 'Cognitivo'; '' si no se reconoce."""
    k = _k(texto)
    return next((d for d in DOMINIOS if k.startswith(_k(d))), '')


def nivel_de(texto) -> str:
    """'AnálisisBAK', 'Comprensión B', 'Conocimientos' → 'analisis', 'comprension', 'conocimiento'."""
    k = re.sub(r'(bak|b)$', '', _k(texto))
    return 'conocimiento' if k.startswith('conocimiento') else k


def cargar_base(ruta: str = RUTA_BD) -> Dict[str, List[Tuple[str, str]]]:
    """{verbo normalizado: [(dominio, subcategoría normalizada), …]}; vacío si no existe la base."""
    if not os.path.exists(ruta):
        return {}
    d = pd.read_excel(ruta, sheet_name='Verbos', engine='openpyxl')
    d.columns = [_k(c) for c in d.columns]
    base: Dict[str, List[Tuple[str, str]]] = {}
    for verbo, sub, dom in zip(d['verbo'], d['nombresubcategoria'], d['nombredominio']):
        v, s, dm = _k(verbo), nivel_de(sub), dominio_de(dom)
        if len(v) >= 3 and s and dm and (dm, s) not in base.setdefault(v, []):
            base[v].append((dm, s))
    return base


def leer_ra(archivos: Iterable[Tuple[str, object]]) -> List[Dict]:
    """RA únicos por matriz con su verbo y lo que declara el Paso 3 (taxonomía, dominio y nivel)."""
    filas = []
    for nombre, fuente in archivos:
        matriz, programa, sede = nombre_matriz(nombre)
        if hasattr(fuente, 'seek'):
            fuente.seek(0)
        wb = openpyxl.load_workbook(fuente, read_only=True, data_only=True)
        try:
            ws = next((wb[s] for s in wb.sheetnames if s.strip().startswith('Paso 3') and 'backup' not in s.lower()), None)
            vistos = {}
            for r in (ws.iter_rows(min_row=3, values_only=True) if ws is not None else []):
                if len(r) > 8 and _ne(r[8]) and _norm(r[8]) != 'resultados aprendizaje':
                    vistos.setdefault(_norm(r[8]), r)
        finally:
            wb.close()
        for r in vistos.values():
            tax = 'BAK' if 'bak' in _k(r[4]) else ('Bloom' if _ne(r[4]) else '')
            dom, niv = dominio_de(r[5]), nivel_de(r[6]) if _ne(r[6]) else ''
            filas.append({'Matriz': matriz, 'Programa': programa, 'Sede': sede, 'RA': str(r[8]).strip(), 'clave': _norm(r[8]),
                          'Verbo': str(r[7]).strip() if _ne(r[7]) else '', 'Tipo de saber': str(r[2]).strip() if _ne(r[2]) else '',
                          'Taxonomía': tax, 'Dominio declarado': dom, 'Nivel declarado': niv,
                          'Exigencia declarada': nivel_exigencia(r[4], r[5], r[6])})
    return filas


def _escala(dominio: str, nivel: str):
    """Nivel 1–6 de un (dominio, subcategoría) de la base; None si no se reconoce."""
    return nivel_exigencia('bak' if dominio != 'Cognitivo' else 'bloom', dominio, nivel)


def contrastar(ra: List[Dict], base: Dict[str, List[Tuple[str, str]]]) -> pd.DataFrame:
    """Agrega a cada RA el dominio y nivel que la base asigna a su verbo y el resultado del contraste."""
    salida = []
    for x in ra:
        cands = base.get(_k(x['Verbo']), [])
        doms = sorted({d for d, _ in cands})
        mismos = [(d, s) for d, s in cands if d == x['Dominio declarado']]
        base_nivel = ''
        dif = None
        if not cands:
            estado = 'Verbo fuera de la base'
        elif not mismos:
            estado = 'Dominio difiere'
        else:
            decl = x['Exigencia declarada']
            esc = [(abs((_escala(d, s) or 0) - decl) if decl is not None and _escala(d, s) is not None else 99, d, s)
                   for d, s in mismos]
            dist, d, s = min(esc)
            base_nivel = s
            if decl is None or dist == 99:
                estado = 'Nivel difiere'
            else:
                dif = round(decl - _escala(d, s), 1)
                estado = 'Coincide' if abs(dif) < 0.5 else 'Nivel difiere'
        salida.append({**x, 'Dominio(s) según la base': ', '.join(doms), 'Nivel según la base': base_nivel,
                       'Diferencia de nivel (declarado − base)': dif, 'Contraste': estado})
    return pd.DataFrame(salida)


def resumen_contraste(df: pd.DataFrame) -> Dict:
    """Conteos y porcentajes del contraste sobre el conjunto de RA (para la pantalla y el PDF)."""
    n = len(df)
    out = {'n': n, 'estados': {e: int((df['Contraste'] == e).sum()) for e in ESTADOS} if n else {e: 0 for e in ESTADOS}}
    out['pct'] = {e: round(100 * v / n, 1) if n else None for e, v in out['estados'].items()}
    return out


def indice_exigencia(exigencia: Iterable) -> float:
    """Índice de exigencia = (nivel medio − 1) / 5 × 100, sobre los RA con nivel reconocido (None si no hay)."""
    v = [x for x in exigencia if x is not None and not pd.isna(x)]
    return round((sum(v) / len(v) - 1) / 5 * 100, 1) if v else None


def por_grupo(df: pd.DataFrame, columna: str) -> pd.DataFrame:
    """Por matriz o programa: RA únicos, % Bloom/BAK, índice de exigencia y % de cada resultado del contraste."""
    filas = []
    for g, d in df.groupby(columna, sort=True):
        n = len(d)
        fila = {columna: g, 'RA únicos': n, '% Bloom': round(100 * (d['Taxonomía'] == 'Bloom').mean(), 1),
                '% BAK': round(100 * (d['Taxonomía'] == 'BAK').mean(), 1),
                'Índice de exigencia': indice_exigencia(d['Exigencia declarada'])}
        for e in ESTADOS:
            fila[f'% {e}'] = round(100 * (d['Contraste'] == e).mean(), 1)
        filas.append(fila)
    return pd.DataFrame(filas)
