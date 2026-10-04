"""
Asociación del perfil de egreso con el contenido de SU PROPIA matriz.

Replica la lógica de la reevaluación experta (auditoría, Administración de Empresas PBOG):
  - Unidades del perfil: funciones del texto narrativo (src/atributos_perfil) y componentes esenciales
    enumerados por el programa en el Paso 1 (Áreas profesionales, Tareas profesionales, Poblaciones de
    actuación; una línea = un ítem).
  - Entidades: competencias específicas (Paso 2), RA de programa (Paso 3) y asignaturas no electivas
    (Paso 5). La competencia genérica institucional y su RA no cuentan.
  - Unidades de contenido finas: cada competencia en redacción, objeto, finalidad y condición; cada RA
    en sus cláusulas; cada asignatura en las oraciones de sus indicadores de logro y en sus núcleos.
  - Decisión: un ítem está respaldado por una entidad si su CONCEPTO DISTINTIVO aparece en alguna unidad:
      cobertura = Σ IDF(lemas del ítem presentes en la unidad) / Σ IDF(lemas del ítem), con IDF calculado
      sobre todas las unidades de la matriz (los términos comunes del programa pesan poco; los específicos,
      como «táctico» o «gubernamental», mucho) y equivalencias de dominio (SINONIMOS).
      explícito  cobertura ≥ PARAMS['explicita']
      parcial    cobertura ≥ PARAMS['parcial'] y similitud semántica (embeddings multilingües) ≥ PARAMS['semantica']
      núcleo     si PARAMS['nucleo'], el sustantivo principal del ítem y su adjetivo distintivo pesan PESO_NUCLEO
                 veces más en la cobertura (regla suave): «Consultoría empresarial» apenas se respalda con
                 «estrategias empresariales», porque falta el término que más pesa. Una regla obligatoria se probó y
                 se descartó: el etiquetado del modelo pequeño de spaCy y las paráfrasis le hacían perder respaldos reales.
  - Vocabulario de contexto: los lemas presentes en más de UMBRAL_CONTEXTO de las unidades de la matriz pesan
    PESO_CONTEXTO («empresarial» en gestión, «pedagógico» en educación, «diseño» en artes); se calcula por matriz.
  - Negación: «no gubernamental», «sin ánimo de lucro» forman un solo término.
  - Asignaturas: se marca el bloque institucional (Paso 5, «B.Institucional»), común a todos los programas, para
    distinguir el respaldo disciplinar del transversal. Se distingue MENCIÓN (el concepto está en el nombre o en un
    núcleo) de DESARROLLO (está en un indicador de logro, que es lo que se evalúa); PARAMS['desarrollo'] exige lo segundo.
    La similitud semántica sola no basta: dos textos del mismo tema son parecidos aunque el concepto falte
    (relación temática, que no cuenta).
Umbrales calibrados contra la referencia experta (auditoria/scripts/asociacion_perfil_calibrada.py).
"""

import os
import re
import unicodedata
from functools import lru_cache
from typing import Dict, List

import numpy as np
import openpyxl

from src.atributos_perfil import extraer_atributos_detalle, lexico_de_matriz

CAPAS = ('Competencia', 'RA', 'Asignatura')
MODELO = 'paraphrase-multilingual-MiniLM-L12-v2'
GENERICA = 'fenomenos contemporaneos'
# Calibrados contra la referencia experta de 11 programas, 429 ítems (auditoria/evidencia_calibracion_asociacion.json):
# acuerdo 80,0 %, kappa 0,45
PARAMS = {'explicita': 0.25, 'parcial': 0.05, 'semantica': 0.58, 'nucleo': False}
# Equivalencias de dominio (lema → lemas que cuentan como el mismo concepto); ampliable por la institución
SINONIMOS = {
    'starups': {'startup', 'emprendimiento', 'emprender', 'negocio'}, 'startup': {'emprendimiento', 'emprender', 'negocio'},
    'emprendimiento': {'emprender', 'intraemprendimiento'}, 'emprender': {'emprendimiento', 'intraemprendimiento'},
    'finanza': {'financiero', 'financiera', 'financiar'}, 'financiero': {'finanza'},
    'consultoria': {'consultor', 'asesoria', 'asesorar', 'consultorio'}, 'asesoria': {'asesorar', 'consultoria'},
    'gerenciar': {'gerencia', 'gestionar', 'gestion', 'administrar', 'administracion'},
    'gerencia': {'gerenciar', 'gestion', 'gestionar', 'administracion'},
    'dirigir': {'direccion', 'liderar', 'liderazgo'}, 'liderar': {'liderazgo', 'dirigir', 'direccion'},
    'pyme': {'pequeno', 'mediano'}, 'mipyme': {'pequeno', 'mediano', 'micro'},
    'legal': {'juridico', 'legislacion', 'normatividad', 'normativa', 'norma', 'ley', 'normativo'},
    'normatividad': {'normativa', 'norma', 'legal', 'legislacion', 'juridico', 'normativo'},
    'normativa': {'normatividad', 'norma', 'legal', 'legislacion', 'juridico', 'normativo'},
    'corporativo': {'empresarial', 'organizacional', 'empresa', 'organizacion'},
    'plan': {'planeacion', 'planificacion', 'planear', 'planificar'},
    # Educación
    'docencia': {'docente', 'ensenanza', 'ensenar', 'pedagogia', 'pedagogico'}, 'docente': {'docencia', 'ensenanza', 'maestro', 'profesor'},
    'ensenanza': {'ensenar', 'docencia', 'pedagogia', 'didactica'}, 'pedagogia': {'pedagogico', 'didactica', 'ensenanza'},
    'nino': {'infancia', 'infantil', 'nina', 'primera infancia'}, 'infancia': {'nino', 'infantil'}, 'infantil': {'infancia', 'nino'},
    'estudiante': {'alumno', 'aprendiz', 'educando'}, 'curriculo': {'curricular', 'plan de estudio'},
    # Artes y diseño
    'crear': {'creacion', 'creativo', 'disenar', 'diseno'}, 'disenar': {'diseno', 'disenador', 'crear', 'creacion'},
    'prototipo': {'prototipar', 'prototipado', 'modelo'}, 'indumentaria': {'prenda', 'vestuario', 'moda'},
    'audiovisual': {'video', 'cine', 'multimedia'}, 'grafico': {'visual', 'grafica'},
    # TIC e ingeniería
    'software': {'aplicacion', 'programa', 'programacion', 'desarrollo'}, 'programacion': {'programar', 'codigo', 'software'},
    'red': {'telecomunicacion', 'conectividad'}, 'dato': {'informacion', 'base de dato'},
    'infraestructura': {'arquitectura', 'plataforma'}, 'algoritmo': {'algoritmia'},
    # Ciencias sociales, comunicación y servicios
    'comunidad': {'poblacion', 'comunitario', 'territorio'}, 'poblacion': {'comunidad', 'grupo'},
    'periodismo': {'periodistico', 'noticia', 'informativo'}, 'medio': {'medios de comunicacion'},
    'turismo': {'turistico', 'turista'}, 'hotel': {'hotelero', 'hoteleria', 'alojamiento'},
    'publico': {'estatal', 'gubernamental', 'estado', 'entidad publica'},
    # Ampliación 2026-10-03 (auditoria/scripts/detectar_sinonimos_faltantes.py): formas que spaCy lematiza
    # distinto y equivalencias funcionales validadas en la revisión uno a uno
    'publica': {'publico', 'estatal', 'gubernamental', 'estado'}, 'startups': {'startup', 'emprendimiento', 'emprender'},
    'report': {'reporte', 'informe'}, 'gestor': {'gestion', 'gestionar'}, 'ejecutar': {'ejecucion'},
    'actuar': {'actuacion'}, 'aportar': {'aporte'},
    # Sector público: entidades → vocabulario de lo público
    **{k: {'publico', 'publica', 'estatal', 'estado', 'gobierno', 'territorial'} for k in
       ('alcaldia', 'gobernacion', 'ministerio', 'superintendencia', 'secretaria', 'descentralizada', 'gubernamental',
        'departamental')},
    # Tercer sector
    **{k: {'comunitario', 'solidario'} for k in ('fundacion', 'ongs', 'sin_animo_de_lucro', 'no_gubernamentales')},
    # Funciones de rol
    'coordinar': {'coordinacion', 'gestionar', 'gestion', 'organizar'},
    'supervisar': {'supervision', 'control', 'controlar', 'seguimiento'},
    'supervisor': {'supervision', 'supervisar', 'control', 'seguimiento'},
    'asesorar': {'asesoria', 'consultoria', 'orientar', 'recomendar', 'recomendacion'},
    'asesor': {'asesoria', 'asesorar', 'consultoria', 'orientar'},
    'recomendacion': {'recomendar', 'propuesta', 'proponer'},
    'manager': {'gerente', 'gerencia', 'gerenciar'}, 'firma': {'empresa'}, 'departamento': {'area', 'dependencia'},
    'rural': {'territorio', 'territorial', 'campo'}, 'equipo': {'grupo', 'colaborativo'},
    'abogado': {'juridico', 'derecho'}, 'cierre': {'terminacion', 'liquidacion'},
    # Servicios, mercadeo y publicidad (revisión de matrices con más de 15 % sin asociación)
    'bar': {'bebida', 'coctel'}, 'feria': {'evento', 'exposicion'}, 'resort': {'alojamiento', 'hotel', 'hotelero'},
    'agroturismo': {'turismo', 'rural'}, 'venta': {'comercial', 'vender'}, 'community': {'comunidad', 'red', 'social'},
    'tablero': {'indicador', 'dashboard', 'visualizacion'}, 'copy': {'redaccion', 'redactar', 'texto', 'copywriting'},
    'pauta': {'medio', 'publicitario', 'anuncio'},
    # Educación, STEM y derecho (revisión de matrices con más de 10 % sin asociación)
    **{k: {'stem'} for k in ('ciencia', 'tecnologia', 'ingenieria', 'matematica')},
    'multicultural': {'cultural', 'intercultural', 'diversidad'}, 'plurilingue': {'lengua', 'linguistico', 'diversidad'},
    'discapacidad': {'inclusion', 'inclusivo', 'diversidad'}, 'excepcional': {'inclusion', 'inclusivo'},
    'conciliar': {'conflicto', 'negociacion', 'mediacion'}, 'conciliacion': {'conflicto', 'negociacion', 'mediacion'},
    'conciliador': {'conciliacion', 'conflicto', 'negociacion', 'mediacion'},
    **{k: {'familia', 'familiar'} for k in ('adopcion', 'custodia', 'parental')},
    'tutoria': {'acompanamiento', 'acompanar', 'orientar'},
}
# Verbos que el modelo pequeño de spaCy no lematiza (forma conjugada → infinitivo)
IRREGULARES = {'promueve': 'promover', 'resuelve': 'resolver', 'demuestra': 'demostrar', 'sostiene': 'sostener',
               'mantiene': 'mantener', 'obtiene': 'obtener', 'atiende': 'atender', 'entiende': 'entender',
               'invierte': 'invertir', 'convierte': 'convertir', 'sugiere': 'sugerir', 'adquiere': 'adquirir',
               'sirve': 'servir', 'mide': 'medir', 'sigue': 'seguir', 'elige': 'elegir', 'construye': 'construir',
               'contribuye': 'contribuir', 'distribuye': 'distribuir', 'incluye': 'incluir', 'influye': 'influir',
               'conduce': 'conducir', 'produce': 'producir', 'ofrece': 'ofrecer', 'reconoce': 'reconocer'}
PESO_NUCLEO = 3.0
# Sustantivos de agente → actividad que debe estar presente (roles: «Coordinador pedagógico», «Consultor…»)
AGENTE_IRREGULAR = {'consultor': ['consultoria', 'asesorar', 'asesoria', 'consultar'], 'asesor': ['asesorar', 'asesoria'],
                    'auditor': ['auditar', 'auditoria'], 'gerente': ['gerencia', 'gerenciar'], 'director': ['dirigir', 'direccion'],
                    'docente': ['docencia', 'ensenar', 'ensenanza'], 'investigador': ['investigar', 'investigacion'],
                    'analista': ['analizar', 'analisis'], 'pedagogo': ['pedagogia', 'pedagogico'], 'lider': ['liderar', 'liderazgo'],
                    'emprendedor': ['emprender', 'emprendimiento'], 'contratista': ['contratar', 'contratacion'],
                    'periodista': ['periodismo', 'periodistico'], 'disenador': ['disenar', 'diseno'],
                    'tutor': ['tutoria', 'acompanar', 'acompanamiento']}
UMBRAL_CONTEXTO = 0.15   # proporción de unidades de la matriz a partir de la cual un lema es vocabulario de contexto
PESO_CONTEXTO = 0.25
FACTOR_ROL = 0.5   # cobertura que conserva un rol cuya actividad no aparece en la unidad (0 = regla estricta;
                   # 0.5 maximiza el acuerdo, ver auditoria/evidencia_factor_rol.json)
RAIZ = 6   # familias de palabras: «empresa» ~ «empresarial», «organización» ~ «organizacional»

# Configuración externa y versionada (config_asociacion_perfil.json). Los valores de arriba son el respaldo si el
# archivo no existe; la institución amplía sinónimos o ajusta parámetros sin tocar el código.
CONFIG_DEFECTO = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'config_asociacion_perfil.json')
VERSION_CONFIG = 'interna'
VALIDACION_CONFIG = {}   # metadato: con qué referencia y qué acuerdo se calibró (no entra en la huella)
_CLAVES_CONFIG = ('PARAMS', 'FACTOR_ROL', 'PESO_NUCLEO', 'UMBRAL_CONTEXTO', 'PESO_CONTEXTO', 'RAIZ', 'SINONIMOS',
                  'AGENTE_IRREGULAR')


def cargar_config(ruta: str = None) -> Dict:
    import json
    with open(ruta or CONFIG_DEFECTO, encoding='utf-8') as f:
        return json.load(f)


def aplicar_config(cfg: Dict) -> None:
    """Activa una configuración del método (dict con las claves de config_asociacion_perfil.json)."""
    global PARAMS, FACTOR_ROL, PESO_NUCLEO, UMBRAL_CONTEXTO, PESO_CONTEXTO, RAIZ, SINONIMOS, AGENTE_IRREGULAR, VERSION_CONFIG
    global VALIDACION_CONFIG
    PARAMS = dict(cfg.get('PARAMS', PARAMS))
    FACTOR_ROL = float(cfg.get('FACTOR_ROL', FACTOR_ROL))
    PESO_NUCLEO = float(cfg.get('PESO_NUCLEO', PESO_NUCLEO))
    UMBRAL_CONTEXTO = float(cfg.get('UMBRAL_CONTEXTO', UMBRAL_CONTEXTO))
    PESO_CONTEXTO = float(cfg.get('PESO_CONTEXTO', PESO_CONTEXTO))
    RAIZ = int(cfg.get('RAIZ', RAIZ))
    if 'SINONIMOS' in cfg:
        SINONIMOS = {k: set(v) for k, v in cfg['SINONIMOS'].items()}
    if 'AGENTE_IRREGULAR' in cfg:
        AGENTE_IRREGULAR = {k: list(v) for k, v in cfg['AGENTE_IRREGULAR'].items()}
    VERSION_CONFIG = str(cfg.get('VERSION', 'sin versión'))
    VALIDACION_CONFIG = dict(cfg.get('VALIDACION') or {})


def config_actual() -> Dict:
    """Configuración activa, con su versión y una huella SHA-256 que identifica exactamente el método aplicado."""
    import hashlib
    import json
    cfg = {'PARAMS': PARAMS, 'FACTOR_ROL': FACTOR_ROL, 'PESO_NUCLEO': PESO_NUCLEO, 'UMBRAL_CONTEXTO': UMBRAL_CONTEXTO,
           'PESO_CONTEXTO': PESO_CONTEXTO, 'RAIZ': RAIZ, 'MODELO_SEMANTICO': MODELO,
           'SINONIMOS': {k: sorted(v) for k, v in sorted(SINONIMOS.items())},
           'AGENTE_IRREGULAR': {k: list(v) for k, v in sorted(AGENTE_IRREGULAR.items())}}
    huella = hashlib.sha256(json.dumps(cfg, sort_keys=True, ensure_ascii=False).encode()).hexdigest()
    return {'VERSION': VERSION_CONFIG, 'HUELLA': huella, 'VALIDACION': VALIDACION_CONFIG, **cfg}


try:
    aplicar_config(cargar_config())
except (OSError, ValueError):
    pass


def _raiz(l: str) -> str:
    return l[:RAIZ] if len(l) >= RAIZ else l
LEMAS_GENERICOS = {'cualquier', 'tipo', 'especialmente', 'distinto', 'diferente', 'clave', 'forma', 'manera', 'nivel',
                   'ambito', 'capaz', 'capacidad', 'grande', 'nuevo', 'propio', 'general', 'campo', 'existente',
                   'emergente', 'junior', 'ser', 'estar', 'tener', 'hacer', 'poder', 'permitir', 'contar', 'realizar', 'llevar'}


def _norm(t) -> str:
    t = unicodedata.normalize('NFKD', str(t)).encode('ascii', 'ignore').decode('ascii').lower()
    return re.sub(r'\s+', ' ', re.sub(r'[^a-z0-9 ]', ' ', t)).strip()


def _ne(v) -> bool:
    return v is not None and str(v).strip() not in ('', 'nan', 'None') and not str(v).strip().startswith('[')


@lru_cache(maxsize=1)
def _nlp():
    import spacy
    return spacy.load('es_core_news_sm', disable=['ner', 'parser'])


@lru_cache(maxsize=1)
def _modelo():
    os.environ.setdefault('HF_HUB_OFFLINE', '1')
    from sentence_transformers import SentenceTransformer
    return SentenceTransformer(MODELO)


NEGACION = re.compile(r'\b(no|sin)\s+(animo de lucro|[a-z]+)', re.I)


def _tokens(texto: str):
    """(lema, pos) de las palabras de contenido; la negación se une a la palabra siguiente."""
    t = _norm(texto)
    t = NEGACION.sub(lambda m: m.group(1) + '_' + m.group(2).replace(' ', '_'), t)
    out = []
    for tok in _nlp()(t):
        lema = _norm(tok.lemma_) if '_' not in tok.text else tok.text
        lema = IRREGULARES.get(tok.text, lema).replace(' ', '_')
        if (tok.pos_ in ('NOUN', 'PROPN', 'ADJ', 'VERB') or '_' in tok.text) and not tok.is_stop and len(tok.text) > 2 \
                and lema not in LEMAS_GENERICOS:
            out.append((lema, tok.pos_))
    return out


def lemas(texto: str) -> frozenset:
    """Lemas de contenido (sustantivos, adjetivos, verbos, nombres propios), sin tildes."""
    return frozenset(l for l, _ in _tokens(texto))


SUSTANTIVOS_LIGEROS = {'gestion', 'sector', 'area', 'organizacion', 'empresa', 'proceso', 'actividad', 'campo',
                       'funcion', 'entidad', 'grupo', 'conjunto', 'parte', 'ambito', 'contexto', 'escenario', 'aspecto'}


def sustantivo_principal(texto: str, contexto: frozenset = frozenset()) -> str:
    """Concepto que debe estar presente: primer sustantivo del ítem que no sea «ligero» (gestión, sector, área…)
    ni vocabulario de contexto de la matriz; si no lo hay, el primer adjetivo distintivo; si no, el primer sustantivo."""
    return nucleo_item(texto, contexto)[0]


def actividad_de_rol(texto: str) -> List[str]:
    """Si el ítem empieza con un sustantivo de agente (rol), lemas de la actividad que lo define:
    «Coordinador pedagógico» → coordinar, coordinacion; «Creador de contenidos» → crear, creacion."""
    toks = _tokens(texto)
    if not toks:
        return []
    l = toks[0][0]
    l = re.sub(r'a$', '', l) if re.search(r'(dor|tor|sor)a$', l) else l
    if l in AGENTE_IRREGULAR:
        return AGENTE_IRREGULAR[l]
    m = re.match(r'^([a-z]+?)(ador|edor|idor)$', l)
    if m:
        v = m.group(1) + {'ador': 'ar', 'edor': 'er', 'idor': 'ir'}[m.group(2)]
        return [v, m.group(1) + 'acion' if m.group(2) == 'ador' else m.group(1) + 'cion']
    return []


def nucleo_item(texto: str, contexto: frozenset = frozenset()) -> List[str]:
    """Términos obligatorios: el sustantivo principal y, si lo sigue de inmediato un adjetivo que no es vocabulario
    de contexto, ese adjetivo («planes tácticos» → plan + táctico; «planes estratégicos» → plan, si «estratégico»
    es contexto de la matriz)."""
    toks = _tokens(texto)
    libre = lambda l: l not in SUSTANTIVOS_LIGEROS and _raiz(l) not in contexto
    for pos in (('NOUN', 'PROPN'), ('ADJ',)):
        j = next((j for j, (l, p) in enumerate(toks) if p in pos and libre(l)), None)
        if j is not None:
            out = [toks[j][0]]
            if pos[0] == 'NOUN' and j + 1 < len(toks) and toks[j + 1][1] == 'ADJ' and libre(toks[j + 1][0]):
                out.append(toks[j + 1][0])
            return out
    j = next((l for l, p in toks if p in ('NOUN', 'PROPN')), toks[0][0] if toks else '')
    return [j] if j else []


def _embed(textos: List[str]) -> np.ndarray:
    if not textos:
        return np.zeros((0, 384))
    return _modelo().encode(textos, batch_size=64, normalize_embeddings=True, show_progress_bar=False)


def _hoja(wb, pref):
    """Hoja cuyo nombre empieza por `pref` (sin hojas de respaldo), o None si la matriz no la tiene."""
    return next((wb[s] for s in wb.sheetnames if s.strip().startswith(pref) and 'backup' not in s.lower()), None)


TAXONOMIA_RESPALDO = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'data', 'raw',
                                  'Taxonomias.xlsx')


@lru_cache(maxsize=1)
def _lexico_respaldo() -> frozenset:
    """Verbos en infinitivo del archivo institucional de taxonomías, para matrices sin hojas de taxonomía."""
    try:
        wb = openpyxl.load_workbook(TAXONOMIA_RESPALDO, read_only=True, data_only=True)
    except OSError:
        return frozenset()
    try:
        return frozenset(x.strip().lower() for ws in wb.worksheets for r in ws.iter_rows(values_only=True) for x in r
                         if isinstance(x, str) and re.fullmatch(r'[A-Za-zÁÉÍÓÚáéíóúñ]+(?:ar|er|ir)(?:se)?', x.strip()))
    finally:
        wb.close()


def _filas(wb, pref, avisos, capa):
    h = _hoja(wb, pref)
    if h is None:
        avisos.append(f'Falta la hoja «{pref}…»: {capa} no se evalúa.')
        return []
    return h.iter_rows(min_row=3, values_only=True)


def _expandir(ls: frozenset) -> List[tuple]:
    """Cada lema del ítem con sus formas equivalentes expresadas como raíces: (lema, {raíces})."""
    return [(l, {_raiz(x) for x in {l} | SINONIMOS.get(l, set())}) for l in sorted(ls)]


def _clausulas(texto: str) -> List[str]:
    partes = re.split(r',?\s+(?:para|que|considerando|seg[uú]n|as[ií] como|mediante|teniendo en cuenta)\s+|;\s*', str(texto))
    return [p.strip(' ,.') for p in [str(texto)] + partes if len(p.split()) >= 2]


ETIQUETA = re.compile(r'^\s*(?:\d+[.)]\s*)?([^:/\n]{3,90}):\s+(?=\S)')


def _items_componente(celda) -> List[str]:
    """Ítems de un componente del Paso 1. Cada línea puede empezar con una etiqueta de categoría
    («1. Educación inicial y preescolar: Educador infantil/Pedagogo…»), que se descarta; el resto se divide
    por «/», «;» o viñeta."""
    out = []
    for linea in re.split(r'\n+', str(celda)):
        linea = re.sub(r'\by\s*/\s*o\b', 'y o', ETIQUETA.sub('', linea.strip()))   # «y/o» no separa ítems
        linea = re.sub(r'(\w)\s*/\s*(as?|os?)\b', r'\1', linea)   # marca de género «Diseñador/a» no separa ítems
        for p in re.split(r'/|;|•', linea):
            p = re.sub(r'^\s*(?:\d+[.)]|[-–—*])\s*', '', p).strip(' .,:-\t')
            if p.count('(') != p.count(')'):          # paréntesis partido entre líneas: «(Gestión…» / «…milla)»
                p = p.replace('(', '').replace(')', '').strip(' .,:-\t')
            if len(p) > 3 and re.search(r'[A-Za-zÁÉÍÓÚáéíóúñ]{3}', p):
                out.append(p)
    return out


def _oraciones(texto: str) -> List[str]:
    return [o.strip() for o in re.split(r'(?<=[.;])\s+|\n+', str(texto)) if len(o.split()) >= 2]


def leer_matriz(fuente) -> Dict:
    """Perfil (texto y componentes) y entidades con sus unidades de texto, desde el Excel original."""
    from src.nucleos_cleaner import es_nucleo_valido, tokenizar_nucleo_celda
    if hasattr(fuente, 'seek'):
        fuente.seek(0)
    wb = openpyxl.load_workbook(fuente, read_only=True, data_only=True)
    avisos = []
    try:
        lex = lexico_de_matriz(wb)
        if not lex:
            lex = set(_lexico_respaldo())
            avisos.append('La matriz no tiene hojas de taxonomía: se usa el léxico de verbos institucional (Taxonomias.xlsx).')
        p1 = [list(r) + [None] * 12 for r in _filas(wb, 'Paso1', avisos, 'el perfil')]
        if not p1 or not (_ne(p1[0][2]) or _ne(p1[0][3]) or any(_ne(r[j]) for r in p1 for j in (7, 8, 9))):
            avisos.append('El Paso 1 no tiene perfil profesional, perfil ocupacional ni componentes: no hay atributos que evaluar.')
        perfil = {'Perfil profesional': p1[0][2] if p1 else '', 'Perfil ocupacional': p1[0][3] if p1 else ''}
        componentes = []
        for j, campo in ((7, 'Áreas profesionales'), (8, 'Tareas profesionales'), (9, 'Poblaciones actuación')):
            for r in p1:
                if _ne(r[j]):
                    componentes += [(campo, it) for it in _items_componente(r[j])]
        competencias = []
        for r in _filas(wb, 'Paso 2', avisos, 'la capa de competencias'):
            r = list(r) + [None] * 8
            if _ne(r[5]) and not _norm(r[5]).startswith('redaccion') and GENERICA not in _norm(r[5]):
                partes = [str(r[5]).strip()] + [str(x).strip() for x in (r[2], r[3], r[4]) if _ne(x)]
                if _ne(r[1]) and _ne(r[2]):
                    partes.append(f'{r[1]} {r[2]}')
                competencias.append({'id': f'C{len(competencias) + 1}', 'nombre': str(r[5]).strip(), 'unidades': partes})
        ras, vistos = [], set()
        for r in _filas(wb, 'Paso 3', avisos, 'la capa de RA'):
            r = list(r) + [None] * 10
            if _ne(r[8]) and _norm(r[8]) != 'resultados aprendizaje' and GENERICA not in _norm(r[0] or '') \
                    and _norm(r[8]) not in vistos:  # D4 (D23): sin mayúsculas, tildes ni puntuación
                vistos.add(_norm(r[8]))
                ras.append({'id': f'RA{len(ras) + 1}', 'nombre': str(r[8]).strip(), 'unidades': _clausulas(r[8])})
        asignaturas, actual = {}, None
        h5 = _hoja(wb, 'Paso 5')
        filas5, col_nuc, col_inst = [], 13, None
        if h5 is None:
            avisos.append('Falta la hoja «Paso 5…»: la capa de asignaturas no se evalúa.')
        else:
            # Columnas por encabezado: la plantilla de posgrado tiene una columna de bloque menos que la de pregrado
            enc = [_norm(v) for v in next(h5.iter_rows(min_row=2, max_row=2, values_only=True), ())]
            col_nuc = next((j for j, v in enumerate(enc) if v.startswith('nucleos')), None)
            col_inst = next((j for j, v in enumerate(enc) if 'institucional' in v), None)
            if col_nuc is None:
                col_nuc = 13
                avisos.append('El Paso 5 no tiene encabezado «Núcleos temáticos»: se usa la columna 13 de la plantilla.')
            filas5 = h5.iter_rows(min_row=3, values_only=True)
        for r in filas5:
            r = list(r) + [None] * 18
            if _ne(r[3]):
                actual = str(r[3]).strip()
            if not actual or _norm(actual).startswith('electiva') or re.fullmatch(r'[\d\s.]+', actual):
                continue
            d = asignaturas.setdefault(actual, {'id': actual, 'nombre': actual, 'unidades': [actual], 'tipos': ['nombre'],
                                                'institucional': False})
            if col_inst is not None and _ne(r[3]) and _ne(r[col_inst]):
                d['institucional'] = True
            nuevas = []
            if _ne(r[4]):
                nuevas += [(o, 'indicador') for o in _oraciones(r[4])]
            if _ne(r[col_nuc]):
                nuevas += [(n, 'nucleo') for n in tokenizar_nucleo_celda(r[col_nuc]) if es_nucleo_valido(n)[0]]
            for u, t in nuevas:
                if u not in d['unidades']:
                    d['unidades'].append(u)
                    d['tipos'].append(t)
    finally:
        wb.close()
    for capa, ents in (('competencias', competencias), ('RA', ras), ('asignaturas', asignaturas)):
        if not ents and not any(capa in a for a in avisos):
            avisos.append(f'No se encontraron {capa} con contenido: esa capa queda sin respaldo para todos los atributos.')
    return {'lexico': lex, 'perfil': perfil, 'componentes': componentes, 'avisos': avisos,
            'entidades': {'Competencia': competencias, 'RA': ras, 'Asignatura': list(asignaturas.values())}}


# Ítems que no son atributos del perfil: cifras de encuestas de empleo («Empleados del estado (86)»), situación
# laboral y etiquetas de nivel. Se informan como «No evaluable» y no entran en los indicadores.
CIFRA_ENCUESTA = re.compile(r'\(\s*\d+\s*\)\s*$')
NO_ATRIBUTO = re.compile(r'^(operativo|tactico|estrategico)$')


def evaluable(item: str) -> bool:
    return not (CIFRA_ENCUESTA.search(str(item)) or NO_ATRIBUTO.search(_norm(item)))


def items_perfil(m: Dict) -> List[Dict]:
    """Ítems a evaluar: funciones del texto narrativo y componentes esenciales."""
    items = []
    for perfil, texto in m['perfil'].items():
        persona = 'tercera' if perfil == 'Perfil profesional' else 'infinitivo'
        for a in extraer_atributos_detalle(texto or '', m['lexico'], persona):
            items.append({'perfil': perfil, 'origen': 'Texto narrativo', 'tipo': a['tipo'], 'item': a['breve']})
    vistos = {_norm(t) for _, t in m['componentes']}
    items = [i for i in items if not (i['perfil'] == 'Perfil ocupacional' and _norm(i['item']) in vistos)]   # sin duplicar componentes
    for campo, texto in m['componentes']:
        items.append({'perfil': 'Perfil ocupacional', 'origen': campo, 'tipo': 'Componente', 'item': texto})
    return items


def puntajes(m: Dict, items: List[Dict]) -> Dict:
    """Por capa: matrices ítem × entidad de cobertura ponderada por IDF y similitud semántica, y unidad de evidencia."""
    unidades = {capa: [(k, u) for k, e in enumerate(m['entidades'][capa]) for u in e['unidades']] for capa in CAPAS}
    # tipo de unidad: en asignaturas, «indicador» = desarrollo evaluado; nombre y núcleo = mención
    tipo_u = {capa: [t for e in m['entidades'][capa] for t in e.get('tipos', ['indicador'] * len(e['unidades']))] for capa in CAPAS}
    L_u = {capa: [frozenset(_raiz(x) for x in lemas(u)) for _, u in unidades[capa]] for capa in CAPAS}
    todas = [l for capa in CAPAS for l in L_u[capa]]
    n = max(len(todas), 1)
    df = {}
    for ls in todas:
        for l in ls:
            df[l] = df.get(l, 0) + 1
    contexto = {l for l, c in df.items() if c / n > UMBRAL_CONTEXTO}
    idf = lambda l: float(np.log((n + 1) / (df.get(_raiz(l), 0) + 1)) + 1) * (PESO_CONTEXTO if _raiz(l) in contexto else 1.0)
    E_items = _embed([i['item'] for i in items])
    out = {}
    for capa in CAPAS:
        ents = m['entidades'][capa]
        U = _embed([u for _, u in unidades[capa]])
        S = E_items @ U.T if len(U) else np.zeros((len(items), 0))
        cob = np.zeros((len(items), len(ents)))
        cob_n = np.zeros((len(items), len(ents)))   # cobertura con el núcleo ponderado
        cob_d = np.zeros((len(items), len(ents)))   # solo unidades de desarrollo (indicadores)
        cob_nd = np.zeros((len(items), len(ents)))
        sem = np.zeros((len(items), len(ents)))
        evid = [[None] * len(ents) for _ in items]
        for i, it in enumerate(items):
            actividad = actividad_de_rol(it['item']) if it.get('origen') != 'Texto narrativo' else []
            # la actividad del rol y sus sinónimos: «conciliador» → conciliar, conciliación, conflicto, negociación
            rol = {_raiz(x) for a in actividad for x in {a} | SINONIMOS.get(a, set())}
            # el sustantivo de agente cuenta como su actividad: «Educador» ~ «educación», «Formador» ~ «formación»
            grupos = [(l, g | rol if re.search(r'(dor|tor|sor)$', l) or l in AGENTE_IRREGULAR else g)
                      for l, g in _expandir(lemas(it['item']))]
            total = sum(idf(l) for l, _ in grupos) or 1.0
            nucleo = [{_raiz(x) for x in {t} | SINONIMOS.get(t, set())} for t in nucleo_item(it['item'], frozenset(contexto))]
            raices_n = {r for g in nucleo for r in g}
            w_n = {l: idf(l) * (PESO_NUCLEO if g & raices_n else 1.0) for l, g in grupos}
            total_n = sum(w_n.values()) or 1.0
            mejor, mejor_n = {}, {}
            d_c, d_n = {}, {}
            for j, ((k, u), lu) in enumerate(zip(unidades[capa], L_u[capa])):
                c = sum(idf(l) for l, g in grupos if g & lu) / total
                if rol and not (rol & lu):          # regla estructural: un rol solo se respalda si su actividad está presente
                    c *= FACTOR_ROL
                s_ = float(S[i, j])
                if (c, s_) > mejor.get(k, (-1, -1, None))[:2]:
                    mejor[k] = (c, s_, u)
                c_n = sum(w_n[l] for l, g in grupos if g & lu) / total_n
                if rol and not (rol & lu):
                    c_n *= FACTOR_ROL
                if (c_n, s_) > mejor_n.get(k, (-1, -1, None))[:2]:
                    mejor_n[k] = (c_n, s_, u)
                if tipo_u[capa][j] == 'indicador':
                    d_c[k], d_n[k] = max(d_c.get(k, 0.0), c), max(d_n.get(k, 0.0), c_n)
                sem[i, k] = max(sem[i, k], s_)
            for k, (c, _, u) in mejor.items():
                cob[i, k], evid[i][k] = c, u
            for k, (c, _, u) in mejor_n.items():
                cob_n[i, k] = c
            for k in d_c:
                cob_d[i, k], cob_nd[i, k] = d_c[k], d_n[k]
        out[capa] = {'cob': cob, 'cob_n': cob_n, 'cob_d': cob_d, 'cob_nd': cob_nd, 'sem': sem, 'evidencia': evid}
    return out


def decidir(cob: float, sem: float, params: Dict = None) -> str:
    """cob: cobertura (con núcleo exigido si params['nucleo'])."""
    p = params or PARAMS
    if cob >= p['explicita']:
        return 'Explícita'
    if cob >= p['parcial'] and sem >= p['semantica']:
        return 'Parcial'
    return 'No asociada'


def asociar(fuente, params: Dict = None, avisos: List[str] = None) -> List[Dict]:
    """Filas en formato largo: ítem × capa × entidad respaldada (o NINGUNA), con evidencia y nivel.
    Si se pasa `avisos` (lista), se le agregan las hojas o capas que la matriz no permitió evaluar."""
    m = leer_matriz(fuente)
    if avisos is not None:
        avisos.extend(m.get('avisos', []))
    items = items_perfil(m)
    P = puntajes(m, items)
    filas = []
    for i, it in enumerate(items):
        for capa in CAPAS:
            ents = m['entidades'][capa]
            prm = params or PARAMS
            clave = ('cob_n' if prm.get('nucleo') else 'cob')
            if capa == 'Asignatura' and prm.get('desarrollo'):
                clave = 'cob_nd' if prm.get('nucleo') else 'cob_d'
            C = P[capa][clave]
            hits = [(k, decidir(C[i, k], P[capa]['sem'][i, k], prm)) for k in range(len(ents))]
            hits = [(k, nv) for k, nv in hits if nv != 'No asociada']
            if not hits:
                k = int(np.argmax(P[capa]['cob'][i] + P[capa]['sem'][i])) if len(ents) else None
                filas.append({**it, 'capa': capa, 'entidad': 'NINGUNA', 'nivel': 'No asociada', 'evidencia': '',
                              'mas_cercana': ents[k]['nombre'][:150] if k is not None else '',
                              'cobertura': round(float(P[capa]['cob'][i, k]), 3) if k is not None else None,
                              'similitud': round(float(P[capa]['sem'][i, k]), 3) if k is not None else None})
            for k, nivel in hits:
                filas.append({**it, 'capa': capa, 'entidad': ents[k]['id'] if capa != 'Asignatura' else ents[k]['nombre'],
                              'nivel': nivel, 'evidencia': P[capa]['evidencia'][i][k], 'mas_cercana': '',
                              'cobertura': round(float(P[capa]['cob'][i, k]), 3), 'similitud': round(float(P[capa]['sem'][i, k]), 3)})
    return filas


def resumen_items(filas: List[Dict]) -> List[Dict]:
    """Una fila por ítem del perfil: capas en que tiene respaldo y estado global (sin asociación en ninguna capa)."""
    por_item = {}
    for f in filas:
        clave = (f['perfil'], f['origen'], f['item'])
        d = por_item.setdefault(clave, {'perfil': f['perfil'], 'origen': f['origen'], 'tipo': f['tipo'], 'item': f['item'],
                                        **{c: 0 for c in CAPAS}})
        if f['entidad'] != 'NINGUNA':
            d[f['capa']] = 1
    out = []
    for d in por_item.values():
        n = sum(d[c] for c in CAPAS)
        d['capas_con_respaldo'] = n
        d['evaluable'] = evaluable(d['item'])
        d['estado'] = 'No evaluable' if not d['evaluable'] else 'Sin asociación' if n == 0 else ('Asociado en las tres capas' if n == len(CAPAS) else 'Asociación parcial')
        out.append(d)
    return out
