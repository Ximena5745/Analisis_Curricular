"""
Indicadores V1–V5 del artículo, calculados con las definiciones auditadas (D4, D13–D17).

Fuente: hojas Paso 2, Paso 3 y Paso 4 de cada matriz. Unidad de RA: RA único por matriz (D4,
clave = texto del RA de la columna I del Paso 3 sin mayúsculas, tildes ni puntuación; ajuste C03). Cada estrategia del Paso 4 es un bloque
de celdas combinadas (B, C y E); sus indicadores (D) e instrumentos (F) ocupan filas propias.

  V1 Correspondencia del perfil: % de atributos del perfil alineados en al menos una capa (competencias, RA o
                            asignaturas; src/asociacion_perfil.py; se recibe ya calculado, por matriz).
     RA con competencia     % de RA únicos que citan una competencia redactada en el Paso 2. Condición del
                            instrumento (la plantilla exige la referencia); antes figuraba como V1 (D13).
  V2 Coherencia (D14)       Matriz sin verbo repetido entre competencias específicas; la competencia
                            genérica institucional («Analizar fenómenos contemporáneos») no cuenta.
  V3 Evaluabilidad (D15)    RA con verbo observable (en SaberSer se acepta el verbo afectivo) y
                            finalidad de desempeño o producto concreto. Condición del instrumento.
  V4 Trazabilidad (D16)     % de RA únicos con estrategia meso que declara instrumento, emparejados
                            por similitud ≥ 0,80 o mejor coincidencia mutua ≥ 0,50.
  V5 Evidencia directa (D17) % de estrategias con al menos un indicador N3 o N4.

Los scripts de auditoria/scripts (verificar_v3, v4_alternativas, verificar_v5) son la evidencia del
artículo; este módulo reproduce sus reglas para el dashboard (tests: test_indicadores_articulo).
"""

import os
import re
import unicodedata
from collections import Counter
from difflib import SequenceMatcher
from typing import Dict, Iterable, List, Tuple

import openpyxl

SEDES = ('PBOG', 'PMED', 'VNAL', 'HMED', 'HBOG', 'HVAL')
GENERICA = 'fenomenos contemporaneos'

# --- V3 (D15) -------------------------------------------------------------------------------
NO_OBSERVABLES = {'conocer', 'comprender', 'entender', 'reflexionar', 'valorar', 'apreciar', 'interiorizar',
                  'sensibilizar', 'concientizar', 'percibir', 'aceptar', 'respetar', 'saber', 'creer',
                  'pensar', 'sentir', 'asumir', 'apropiar', 'apropiarse', 'familiarizarse', 'considerar'}
FINALIDAD = re.compile(r'\b(?:para|mediante|a traves de|con el fin de|con el proposito de)\s+(?:(?:la|el|los|las|su|sus)\s+)?'
                       r'([a-z]+(?:ar|er|ir)(?:se|lo|la|los|las)?)\b')
PRODUCTO = re.compile(r'\b(?:un|una|unos|unas)\s+(?:estudio|propuesta|proyecto|informe|plan|prototipo|modelo|diagnostico|'
                      r'producto|documento|ensayo|articulo|programa|estrategia|solucion|analisis|disen[oa]|reporte|portafolio)\b')

# --- V5 (D17): niveles de evidencia -------------------------------------------------------------
IMPACTO = re.compile(r'egresados? (vinculad|emplead|contratad|ubicad)|emplea|insercion|vinculad[oa]s? laboral|laboralmente|contratad|ubicacion laboral|'
                     r'emprendimientos? creados?|impacto en (la|el|las|los) (comunidad|entorno|sector|region)|'
                     r'beneficiari|convenios? (firmad|suscrit)|transferencia')
LOGRO = re.compile(r'puntaje|promedio|aprueb|aprobaci|desempe|nivel de (logro|dominio|desarrollo|competencia)|'
                   r'calificaci|\bnota\b|alcanzan|superan|igual o superior|superior a|por encima|mejora en|rubrica|'
                   r'resultados? (de|en) (las? )?prueba|saber (pro|tyt)')
PERCEPCION = re.compile(r'valoraci|\bvalora\b|satisfacci|percepci|opini|apreciaci|utilidad percibida|nivel de contribuci')
APRENDIZAJE = re.compile(LOGRO.pattern + r'|sustentaci|jurad|evaluad[oa]s? con (rubrica|criterios)|arbitrad|indexad')
TRANSFERENCIA = re.compile(r'(evaluaci|desempe|calificaci|valoracion del (tutor|jefe|empleador))[^.]*'
                           r'(practica (empresarial|profesional|laboral)|pasanti|contexto real|escenario real|en la empresa|usuarios reales)')
NIVELES = ('N1 Implementación', 'N2 Percepción y reacción', 'N3 Aprendizaje demostrado', 'N4 Transferencia', 'N5 Efecto externo')


def _norm(t, puntuacion=False) -> str:
    t = unicodedata.normalize('NFKD', str(t)).encode('ascii', 'ignore').decode('ascii').lower()
    if not puntuacion:
        t = re.sub(r'[^a-z0-9 ]', ' ', t)
    return re.sub(r'\s+', ' ', t).strip()


def _ne(v) -> bool:
    return v is not None and str(v).strip() not in ('', 'nan', 'None')


def _sim(a: str, b: str) -> float:
    return SequenceMatcher(None, a, b, autojunk=False).ratio()


def _hoja(wb, prefijo):
    return next(wb[s] for s in wb.sheetnames if s.strip().startswith(prefijo) and 'backup' not in s.lower())


def nivel_indicador(texto: str) -> str:
    t = _norm(texto, puntuacion=True)
    if IMPACTO.search(t):
        return NIVELES[4]
    if TRANSFERENCIA.search(t) and 'simulad' not in t:
        return NIVELES[3]
    if APRENDIZAJE.search(t):
        return NIVELES[2]
    if PERCEPCION.search(t):
        return NIVELES[1]
    return NIVELES[0]


def ra_evaluable(verbo: str, texto: str, actitudinal: bool) -> Tuple[bool, str]:
    v, t = _norm(verbo), _norm(texto)
    obs = v not in NO_OBSERVABLES or actitudinal
    m = FINALIDAD.search(t)
    fin = m.group(1) if m else None
    if fin and len(fin) > 6:
        fin = re.sub(r'(se|lo|la|los|las)$', '', fin)
    prod = (fin is not None and fin not in NO_OBSERVABLES) or PRODUCTO.search(t) is not None
    motivo = [] if obs else [f'verbo no observable ({v})']
    if not prod:
        motivo.append('sin finalidad o producto declarado' if not m else f'finalidad no observable (para {fin})')
    return obs and prod, '; '.join(motivo)


def nombre_matriz(nombre_archivo: str) -> Tuple[str, str, str]:
    """('Programa_SEDE', programa, sede) desde 'FormatoRA_Programa_SEDE.xlsx'."""
    base = re.sub(r'\.xlsx?$', '', os.path.basename(nombre_archivo))
    base = re.sub(r'^FormatoRA[_-]', '', base)
    programa, _, sede = base.rpartition('_')
    return (base, programa, sede) if sede in SEDES else (base, base, 'ND')


def _leer_matriz(nombre: str, fuente) -> Dict:
    matriz, programa, sede = nombre_matriz(nombre)
    if hasattr(fuente, 'seek'):
        fuente.seek(0)
    wb = openpyxl.load_workbook(fuente, read_only=True, data_only=True)
    try:
        comp = []
        for r in _hoja(wb, 'Paso 2').iter_rows(min_row=3, values_only=True):
            if len(r) > 5 and _ne(r[5]) and not str(r[5]).strip().startswith('['):
                if _norm(r[5]).startswith('redaccion') or _norm(r[1] or '') == 'verbo competencia':
                    continue
                comp.append({'verbo': _norm(r[1]) if _ne(r[1]) else '', 'texto': _norm(r[5]),
                             'generica': GENERICA in _norm(r[5])})
        ra = {}
        for r in _hoja(wb, 'Paso 3').iter_rows(min_row=3, values_only=True):
            if len(r) > 8 and _ne(r[8]) and _norm(r[8]) != 'resultados aprendizaje':
                clave = _norm(r[8])  # D4: sin mayúsculas, tildes ni puntuación (C03)
                d = ra.setdefault(clave, {'texto': str(r[8]).strip(), 'norm': _norm(r[8]), 'verbo': r[7] or '',
                                          'competencias': set(), 'tipos': set(), 'actitudinal': False})
                if _ne(r[0]):
                    d['competencias'].add(_norm(r[0]))
                if _ne(r[2]):
                    d['tipos'].add(str(r[2]).strip())
                d['actitudinal'] |= 'ser' in _norm(r[2] or '') or 'actitudinal' in _norm(r[5] or '')
        estrategias, meso, b = [], [], None
        for r in _hoja(wb, 'Paso 4').iter_rows(min_row=3, values_only=True):
            if not r or len(r) < 6:
                continue
            if _ne(r[1]):
                b = {'estrategia': str(r[1]).strip(), 'indicadores': [], 'ins': False}
                estrategias.append(b)
            if b is None:
                continue
            if _ne(r[3]):
                b['indicadores'].append(str(r[3]).strip())
            b['ins'] |= _ne(r[5])
            if _ne(r[0]):
                meso.append((_norm(r[0]), b))
    finally:
        wb.close()
    return {'matriz': matriz, 'programa': programa, 'sede': sede, 'competencias': comp, 'ra': ra,
            'estrategias': estrategias, 'meso': meso}


def _evaluar_matriz(m: Dict) -> Dict:
    comp_textos = [c['texto'] for c in m['competencias']]
    especificas = Counter(c['verbo'] for c in m['competencias'] if c['verbo'] and not c['generica'])
    repetidos = sorted(v for v, n in especificas.items() if n > 1)

    p4 = sorted({x for x, _ in m['meso']})
    p3 = [d['norm'] for d in m['ra'].values()]
    mejor_p3 = {x: max(p3, key=lambda t: _sim(t, x)) for x in p4} if p3 else {}
    mutuo = set()
    for t in p3:
        if p4:
            x = max(p4, key=lambda x: _sim(t, x))
            if mejor_p3.get(x) == t and _sim(t, x) >= 0.5:
                mutuo.add((t, x))

    ras = []
    for d in m['ra'].values():
        t = d['norm']
        v1 = any(c == ct or _sim(c, ct) >= 0.8 for c in d['competencias'] for ct in comp_textos)
        ok3, motivo3 = ra_evaluable(d['verbo'], d['texto'], d['actitudinal'])
        vinc = [e for x, e in m['meso'] if t == x or _sim(t, x) >= 0.8 or (t, x) in mutuo]
        ras.append({'matriz': m['matriz'], 'programa': m['programa'], 'sede': m['sede'], 'ra': d['texto'],
                    'generica': GENERICA in t, 'RA_competencia': v1, 'V3': ok3, 'motivo_V3': motivo3,
                    'V4': any(e['ins'] for e in vinc), 'estrategias': len({id(e) for e in vinc})})
    est = []
    for e in m['estrategias']:
        niveles = [nivel_indicador(i) for i in e['indicadores']]
        est.append({'matriz': m['matriz'], 'programa': m['programa'], 'sede': m['sede'], 'estrategia': e['estrategia'],
                    'indicadores': len(e['indicadores']), 'niveles': niveles, 'instrumento': e['ins'],
                    'nivel_max': max(niveles) if niveles else None,
                    'V5': any(n[:2] in ('N3', 'N4') for n in niveles)})
    return {'matriz': m['matriz'], 'programa': m['programa'], 'sede': m['sede'],
            'competencias': len(m['competencias']), 'verbos_repetidos': repetidos, 'V2': not repetidos,
            'ras': ras, 'estrategias': est,
            'indicadores': [(e['estrategia'], i, nivel_indicador(i)) for e in m['estrategias'] for i in e['indicadores']]}


def calcular_indicadores(archivos: Iterable[Tuple[str, object]], v1: Dict[str, Tuple[int, int]] = None) -> Dict:
    """
    Calcula V1–V5 sobre un conjunto de matrices.

    Args:
        archivos: pares (nombre_archivo, ruta o archivo abierto) de las matrices FormatoRA.
        v1: por matriz ('Programa_SEDE'), (atributos con alineación en al menos una capa, atributos evaluables), de
            src.asociacion_perfil. Sin este dato, V1 queda vacío.

    Returns:
        dict con 'por_matriz', 'ra', 'estrategias', 'indicadores', 'globales' y 'errores'.
    """
    v1 = v1 or {}
    matrices, errores = [], []
    for nombre, fuente in archivos:
        try:
            matrices.append(_evaluar_matriz(_leer_matriz(nombre, fuente)))
        except Exception as exc:  # hoja ausente o formato distinto
            errores.append({'archivo': nombre, 'causa': str(exc)[:120]})
    ras = [r for m in matrices for r in m['ras']]
    est = [e for m in matrices for e in m['estrategias']]
    inds = [{'matriz': m['matriz'], 'sede': m['sede'], 'estrategia': s, 'indicador': i, 'nivel': n}
            for m in matrices for s, i, n in m['indicadores']]
    pct = lambda xs: round(100 * sum(xs) / len(xs), 1) if xs else None
    por_matriz = []
    for m in matrices:
        r, e = m['ras'], m['estrategias']
        por_matriz.append({'Matriz': m['matriz'], 'Programa': m['programa'], 'Sede': m['sede'],
                           'Competencias': m['competencias'], 'RA únicos': len(r), 'Estrategias': len(e),
                           'V1 %': (round(100 * v1[m['matriz']][0] / v1[m['matriz']][1], 1)
                                    if v1.get(m['matriz'], (0, 0))[1] else None),
                           'RA con competencia %': pct([x['RA_competencia'] for x in r]),
                           'V2': 'Sí' if m['V2'] else 'No',
                           'Verbos repetidos': ', '.join(m['verbos_repetidos']),
                           'V3 %': pct([x['V3'] for x in r]), 'V4 %': pct([x['V4'] for x in r]),
                           'RA sin estrategia': sum(not x['V4'] for x in r),
                           'V5 %': pct([x['V5'] for x in e]),
                           'Estrategias con indicador e instrumento': sum(bool(x['indicadores']) and x['instrumento'] for x in e),
                           'Indicadores por nivel': dict(Counter(n[:2] for x in e for n in x['niveles']))})
    globales = {
        'matrices': len(matrices), 'programas': len({m['programa'] for m in matrices}),
        'competencias': sum(m['competencias'] for m in matrices),
        'ra_unicos': len(ras), 'estrategias': len(est), 'indicadores': len(inds),
        'V1': (round(100 * sum(v1[m['matriz']][0] for m in matrices if m['matriz'] in v1) /
                     sum(v1[m['matriz']][1] for m in matrices if m['matriz'] in v1), 1)
               if any(v1.get(m['matriz'], (0, 0))[1] for m in matrices) else None),
        'atributos_perfil': sum(v1[m['matriz']][1] for m in matrices if m['matriz'] in v1),
        'RA_competencia': pct([r['RA_competencia'] for r in ras]), 'V2': pct([m['V2'] for m in matrices]),
        'V3': pct([r['V3'] for r in ras]), 'V4': pct([r['V4'] for r in ras]),
        'V4_programa': pct([r['V4'] for r in ras if not r['generica']]),
        'ra_genericos': sum(r['generica'] for r in ras),
        'genericos_con_estrategia': sum(r['V4'] for r in ras if r['generica']),
        'ra_sin_estrategia': sum(not r['V4'] for r in ras),
        'ra_sin_estrategia_programa': sum(not r['V4'] and not r['generica'] for r in ras),
        'V5': pct([e['V5'] for e in est]), 'estrategias_directa': sum(e['V5'] for e in est),
        'niveles': dict(Counter(i['nivel'] for i in inds)),
        'matrices_con_repeticion': sum(not m['V2'] for m in matrices),
        'programas_con_repeticion': len({m['programa'] for m in matrices if not m['V2']}),
    }
    return {'globales': globales, 'por_matriz': por_matriz, 'ra': ras, 'estrategias': est,
            'indicadores': inds, 'errores': errores}


def contrastes(por_matriz: List[Dict], min_n: int = 10) -> Dict:
    """
    R6 (Spearman entre V1, V2, V4 y V5 por matriz) y R7 (Kruskal-Wallis entre sedes), solo cuando hay datos
    suficientes. Variables constantes, grupos unitarios (n = 1) o menos de `min_n` matrices se informan como
    motivo en lugar de un coeficiente.
    """
    import itertools

    import numpy as np
    from scipy.stats import kruskal, spearmanr

    filas = [{'Matriz': m['Matriz'], 'Sede': m['Sede'], 'V1': m.get('V1 %'), 'V2': 100.0 if m['V2'] == 'Sí' else 0.0,
              'V4': m.get('V4 %'), 'V5': m.get('V5 %')} for m in por_matriz]
    variables = [v for v in ('V1', 'V2', 'V4', 'V5') if all(f[v] is not None for f in filas)]
    out = {'spearman': [], 'kruskal': [], 'motivos': []}
    if len(filas) < min_n:
        out['motivos'].append(f'Hay {len(filas)} matrices; los contrastes requieren al menos {min_n}.')
        return out
    valores = {v: np.array([f[v] for f in filas], dtype=float) for v in variables}
    constantes = [v for v in variables if np.ptp(valores[v]) == 0]
    if constantes:
        out['motivos'].append(f'Variables constantes en el conjunto cargado (sin varianza): {", ".join(constantes)}.')
    activas = [v for v in variables if v not in constantes]
    for a, b in itertools.combinations(activas, 2):
        rho, p = spearmanr(valores[a], valores[b])
        out['spearman'].append({'Par': f'{a}–{b}', 'rho': round(float(rho), 2), 'p': round(float(p), 3), 'n': len(filas)})
    sedes = {}
    for f in filas:
        sedes.setdefault(f['Sede'], []).append(f)
    unitarias = [s for s, g in sedes.items() if len(g) < 2]
    if unitarias:
        out['motivos'].append(f'Sedes con una sola matriz, excluidas de Kruskal-Wallis: {", ".join(unitarias)}.')
    grupos = {s: g for s, g in sedes.items() if len(g) >= 2}
    if len(grupos) < 2:
        out['motivos'].append('Se necesitan al menos dos sedes con dos o más matrices para comparar sedes.')
        return out
    n, k = sum(len(g) for g in grupos.values()), len(grupos)
    for v in activas:
        muestras = [[f[v] for f in g] for g in grupos.values()]
        if len({x for m in muestras for x in m}) < 2:
            continue
        h, p = kruskal(*muestras)
        out['kruskal'].append({'Variable': v, 'H': round(float(h), 2), 'gl': k - 1, 'p': round(float(p), 3),
                               'épsilon²': round(max(0.0, (float(h) - k + 1) / (n - k)), 3), 'n': n,
                               'Sedes': ', '.join(f'{s} ({len(g)})' for s, g in grupos.items())})
    return out


# --- Exigencia de los RA (Etapa 3) -----------------------------------------------------------
# Progresión de cada dominio, de menor a mayor exigencia. Cada nivel se traduce a una escala común de 1 a 6
# según su posición en su propio dominio: nivel = 1 + (posición − 1) × 5 / (número de niveles − 1).
PROGRESIONES = {
    ('bloom', None): ['conocimiento', 'comprension', 'aplicacion', 'analisis', 'sintesis', 'evaluacion'],
    ('bak', 'cognitivo'): ['conocimiento', 'comprension', 'aplicacion', 'analisis', 'sintesis', 'evaluacion'],
    ('bak', 'procedimental'): ['imitacion', 'manipulacion', 'precision', 'control'],
    ('bak', 'actitudinal'): ['percepcion', 'responder', 'valorar', 'organizar', 'caracterizar'],
}


def _k(t) -> str:
    return re.sub(r'[^a-z]', '', unicodedata.normalize('NFKD', str(t)).encode('ascii', 'ignore').decode().lower())


def nivel_exigencia(taxonomia, dominio, nivel):
    """Nivel 1–6 de un RA según su taxonomía (Bloom o BAK), dominio y nivel declarados; None si no se reconoce."""
    tax = 'bak' if 'bak' in _k(taxonomia) else 'bloom'
    dom = next((d for d in ('cognitivo', 'procedimental', 'actitudinal') if d in _k(dominio)), None)
    prog = PROGRESIONES.get((tax, None if tax == 'bloom' else dom))
    if prog is None:
        return None
    base = re.sub(r'(bak|b)$', '', _k(nivel))
    base = 'conocimiento' if base.startswith('conocimiento') else base
    if base not in prog:
        return None
    return round(1 + (prog.index(base)) * 5 / (len(prog) - 1), 2)


def calcular_exigencia(archivos: Iterable[Tuple[str, object]]) -> List[Dict]:
    """Por matriz: RA únicos, proporción Bloom/BAK e índice de exigencia = (nivel medio − 1) / 5 × 100."""
    out = []
    for nombre, fuente in archivos:
        matriz, programa, sede = nombre_matriz(nombre)
        if hasattr(fuente, 'seek'):
            fuente.seek(0)
        wb = openpyxl.load_workbook(fuente, read_only=True, data_only=True)
        try:
            ws = next((wb[s] for s in wb.sheetnames if s.strip().startswith('Paso 3') and 'backup' not in s.lower()), None)
            vistos = {}
            for r in (ws.iter_rows(min_row=3, values_only=True) if ws is not None else []):
                if len(r) > 8 and _ne(r[8]) and _norm(r[8]) != 'resultados aprendizaje':
                    vistos.setdefault(_norm(r[8]), (r[4], r[5], r[6]))  # D4 (C03)
        finally:
            wb.close()
        niveles = [nivel_exigencia(t, d, n) for t, d, n in vistos.values()]
        validos = [n for n in niveles if n is not None]
        bak = sum('bak' in _k(t) for t, _, _ in vistos.values())
        out.append({'Matriz': matriz, 'Programa': programa, 'Sede': sede, 'RA únicos': len(vistos),
                    '% BAK': round(100 * bak / len(vistos), 1) if vistos else None,
                    '% Bloom': round(100 * (len(vistos) - bak) / len(vistos), 1) if vistos else None,
                    'RA sin nivel reconocido': len(niveles) - len(validos),
                    'Índice de exigencia': round((sum(validos) / len(validos) - 1) / 5 * 100, 1) if validos else None})
    return out


# --- Valoración por criterios (sustituye al puntaje de calidad 0–100; decisión D26) ----------------------------
# Cada tramo se valora contra la regla que la propia plantilla institucional exige (criterio absoluto, sin pesos
# ni línea base): Cumple = la regla se cumple en todos los casos; Parcial = en algunos; No cumple = en ninguno.
# La mediana del conjunto cargado se informa solo como referencia descriptiva; no interviene en el estado.
CRITERIOS = [
    ('V1', 'Perfil alineado', 'Cada atributo del perfil alineado en al menos una capa (competencias, RA o asignaturas)'),
    ('V2', 'Coherencia de competencias', 'Ningún verbo repetido entre competencias específicas'),
    ('V3', 'Evaluabilidad de los RA', 'Todo RA con verbo observable y finalidad de desempeño o producto'),
    ('V4', 'Trazabilidad de los RA', 'Todo RA vinculado a una estrategia mesocurricular con instrumento'),
    ('V5', 'Indicadores de las estrategias', 'Cada estrategia declara al menos un indicador y un instrumento (Paso 4)'),
]
# El nivel de evidencia de los indicadores (N1–N5) no es una regla de la plantilla: se describe y se recomienda,
# no se califica (un «No cumple» por no tener N3–N4 castigaría algo que el instrumento no pide).
ESTADOS = ('Cumple', 'Parcial', 'No cumple')


def estado_criterio(pct) -> str:
    if pct is None:
        return 'Sin dato'
    return 'Cumple' if pct >= 100 else ('No cumple' if pct <= 0 else 'Parcial')


def referencias_conjunto(por_matriz: List[Dict], exigencia: Dict[str, float] = None) -> Dict[str, float]:
    """Medianas del conjunto cargado (solo para situar cada matriz; no definen estados)."""
    import statistics
    ref = {}
    for v in ('V1 %', 'V3 %', 'V4 %', 'V5 %'):
        xs = [m[v] for m in por_matriz if m.get(v) is not None]
        ref[v] = round(statistics.median(xs), 1) if xs else None
    ref['V2'] = round(100 * sum(m['V2'] == 'Sí' for m in por_matriz) / len(por_matriz), 1) if por_matriz else None
    xs = [x for x in (exigencia or {}).values() if x is not None]
    ref['exigencia'] = round(statistics.median(xs), 1) if xs else None
    ref['n'] = len(por_matriz)
    return ref


def valorar_matriz(fila: Dict, sin_respaldo_pct: float = None, exigencia: float = None,
                   referencia: Dict = None, tres_capas_pct: float = None) -> List[Dict]:
    """
    Valoración por criterios de una matriz (fila de calcular_indicadores()['por_matriz']).

    Returns:
        lista de dicts: Variable, Tramo, Criterio, Resultado, Estado y Referencia (mediana del conjunto).
    """
    ref = referencia or {}
    f = lambda x: '—' if x is None else f'{x:.1f} %'.replace('.', ',')
    n_ra, n_est = fila.get('RA únicos') or 0, fila.get('Estrategias') or 0
    v5 = fila.get('V5 %')
    directas = round((v5 or 0) * n_est / 100)
    resultado = {
        'V1': (fila.get('V1 %'), f"{f(fila.get('V1 %'))} con alguna alineación"
               + (f"; {f(sin_respaldo_pct)} sin alineación" if sin_respaldo_pct is not None else '')
               + (f"; {f(tres_capas_pct)} con las tres capas" if tres_capas_pct is not None else '')),
        'V2': (100.0 if fila.get('V2') == 'Sí' else 0.0,
               'Sin verbos repetidos' if fila.get('V2') == 'Sí' else f"Verbo repetido: {fila.get('Verbos repetidos')}"),
        'V3': (fila.get('V3 %'), f"{f(fila.get('V3 %'))} de {n_ra} RA"),
        'V4': (fila.get('V4 %'), f"{f(fila.get('V4 %'))} de {n_ra} RA ({fila.get('RA sin estrategia', 0)} sin estrategia)"),
        'V5': ((100 * fila.get('Estrategias con indicador e instrumento', 0) / n_est) if n_est else None,
               f"{fila.get('Estrategias con indicador e instrumento', 0)} de {n_est} estrategias"),
    }
    refs = {'V1': f(ref.get('V1 %')), 'V2': (f"{f(ref.get('V2'))} de las matrices cumple" if ref.get('V2') is not None else '—'),
            'V3': f(ref.get('V3 %')), 'V4': f(ref.get('V4 %')), 'V5': '—'}
    filas = [{'Variable': v, 'Tramo': e, 'Criterio': c, 'Resultado': resultado[v][1],
              'Estado': estado_criterio(resultado[v][0]), 'Referencia (mediana)': refs[v]} for v, e, c in CRITERIOS]
    niv = fila.get('Indicadores por nivel') or {}
    n_ind = sum(niv.values())
    filas.append({'Variable': '—', 'Tramo': 'Nivel de evidencia de los indicadores (descriptivo)',
                  'Criterio': 'Sin criterio en la plantilla. Recomendación: al menos un indicador de aprendizaje '
                              'demostrado (N3) por estrategia',
                  'Resultado': (f"{n_ind} indicadores: N1 implementación {niv.get('N1', 0)}, N2 percepción "
                                f"{niv.get('N2', 0)}, N3–N4 aprendizaje {niv.get('N3', 0) + niv.get('N4', 0)}, "
                                f"N5 efecto externo {niv.get('N5', 0)}; {directas} de {n_est} estrategias con N3–N4")
                                if n_ind else '—',
                  'Estado': 'Descriptiva', 'Referencia (mediana)': f(ref.get('V5 %')) + ' de estrategias con N3–N4'
                  if ref.get('V5 %') is not None else '—'})
    filas.append({'Variable': '—', 'Tramo': 'Exigencia de los RA (descriptiva)',
                  'Criterio': 'Sin criterio: la exigencia adecuada depende del nivel de formación',
                  'Resultado': '—' if exigencia is None else f'índice {exigencia:.1f}'.replace('.', ','),
                  'Estado': 'Descriptiva',
                  'Referencia (mediana)': '—' if ref.get('exigencia') is None else f"{ref['exigencia']:.1f}".replace('.', ',')})
    return filas
